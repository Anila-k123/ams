"""Merge phase 02: drafting authenticates as the AMS advocate and records who did what
by plain advocate id, in its own database."""

import tempfile
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from core.testing import ALL_PERMISSIONS, auth, make_advocate

from .models import Client, DraftFile, DraftSession, Playbook, Project, Sample, Template


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
@mock.patch('drafting.views._dispatch')      # no parsing / LLM calls in tests
class DraftingAuthTest(TestCase):
    databases = {'default'}

    def setUp(self):
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        self.junior = make_advocate(permissions=ALL_PERMISSIONS, parent_advocate_id=self.senior.id)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.client_row = Client.objects.create(name='Sharma Traders')
        self.project = Project.objects.create(client=self.client_row, name='OS/12/2026')

    def upload(self, advocate, name='NDA.pdf'):
        return self.client.post('/api/drafting/samples/', {
            'file': SimpleUploadedFile(name, b'%PDF-1.4 test'), 'name': name,
            'client': self.client_row.id, 'project': self.project.id}, **auth(advocate))

    def test_upload_records_the_ams_advocate(self, dispatch):
        resp = self.upload(self.junior)
        self.assertEqual(resp.status_code, 201, resp.content)
        sample = Sample.objects.get(id=resp.json()['id'])
        self.assertEqual(sample.uploaded_by_id, self.junior.id)
        dispatch.assert_called_once()

    def test_documents_shared_within_the_practice_only(self, dispatch):
        mine = Sample.objects.get(id=self.upload(self.junior).json()['id'])
        Sample.objects.create(name='Other firm.pdf', file='x.pdf', uploaded_by_id=self.outsider.id)
        for who in (self.senior, self.junior):
            ids = [s['id'] for s in self.client.get('/api/drafting/samples/', **auth(who)).json()['results']]
            self.assertEqual(ids, [mine.id])
        self.assertEqual(self.client.get('/api/drafting/samples/', **auth(self.outsider)).json()['count'], 1)
        self.assertEqual(self.client.get(f'/api/drafting/samples/{mine.id}/', **auth(self.outsider)).status_code, 404)

    def test_draft_sessions_are_the_creators(self, dispatch):
        template = Template.objects.create(name='Mutual NDA', document_type='nda')
        mine = DraftSession.objects.create(template=template, created_by_id=self.senior.id, facts={}, status='ready')
        DraftSession.objects.create(template=template, created_by_id=self.junior.id, facts={}, status='ready')
        body = self.client.get('/api/drafting/draft-sessions/', **auth(self.senior)).json()
        self.assertEqual([s['id'] for s in body['results']], [mine.id])
        self.assertEqual(set(body), {'count', 'next', 'previous', 'results'})   # InstaDraft's paging shape

    def test_login_required_and_clients_refused(self, dispatch):
        self.assertEqual(self.client.get('/api/drafting/samples/').status_code, 401)
        from clientaccess.models import ClientUser
        from core.testing import make_client
        client_login = make_advocate()
        ClientUser.objects.create(advocate_id=client_login.id, client_id=make_client(self.senior).id)
        self.assertEqual(self.client.get('/api/drafting/samples/', **auth(client_login)).status_code, 403)


# Merge phase 03: drafting uses AMS's permission codes (seed_drafting_permissions).
INTERN = ('DRAFT_VIEW', 'DRAFT_CREATE')
JUNIOR = INTERN + ('DRAFT_EXPORT',)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
class DraftingPermissionTest(TestCase):
    databases = {'default'}

    def setUp(self):
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        self.intern = make_advocate(permissions=INTERN, parent_advocate_id=self.senior.id)
        self.junior = make_advocate(permissions=JUNIOR, parent_advocate_id=self.senior.id)
        self.nobody = make_advocate(permissions=('CASE_VIEW',), parent_advocate_id=self.senior.id)
        self.template = Template.objects.create(name='Mutual NDA', document_type='nda',
                                                created_by_id=self.senior.id)

    def test_no_drafting_codes_means_no_drafting(self):
        for url in ('/api/drafting/samples/', '/api/drafting/templates/', '/api/drafting/draft-sessions/',
                    '/api/drafting/playbooks/', '/api/drafting/playbook-clauses/'):
            self.assertEqual(self.client.get(url, **auth(self.nobody)).status_code, 403, url)

    def test_view_code_reads_but_shared_setup_needs_manage(self):
        self.assertEqual(self.client.get('/api/drafting/templates/', **auth(self.intern)).status_code, 200)
        self.assertEqual(self.client.get('/api/drafting/playbooks/', **auth(self.intern)).status_code, 200)
        for who in (self.intern, self.junior):
            self.assertEqual(self.client.delete(f'/api/drafting/templates/{self.template.id}/',
                                                **auth(who)).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/drafting/templates/{self.template.id}/',
                                            **auth(self.senior)).status_code, 204)

    def test_playbooks_no_longer_open(self):
        # InstaDraft had permission_classes = [] on these; inside AMS that would be public.
        self.assertEqual(self.client.get('/api/drafting/playbooks/').status_code, 401)
        self.assertEqual(self.client.get('/api/drafting/playbook-clauses/').status_code, 401)

    def test_export_needs_export_code(self):
        from django.urls import reverse
        for who, expected in ((self.intern, 403), (self.junior, 200)):
            session = DraftSession.objects.create(template=self.template, created_by_id=who.id,
                                                  facts={}, status='ready')
            url = f'/api/drafting/drafts/{session.id}/export/docx/'
            resp = self.client.get(url, **auth(who))
            self.assertEqual(resp.status_code, expected, (url, resp.content[:200]))
        self.assertIn('attachment', resp['Content-Disposition'])


    def test_client_project_member_endpoints_are_gone(self):
        # Merge phase 08: drafting clients/projects are mirrors of AMS clients/cases
        # (link-case, uploads with case_id); they had no practice scope as endpoints.
        for url in ('/api/drafting/clients/', '/api/drafting/projects/', '/api/drafting/members/'):
            self.assertEqual(self.client.get(url, **auth(self.senior)).status_code, 404, url)


# Merge phase 04: AMS cases/documents in-process, and files behind the login.
@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'),
                   DOCUMENT_UPLOAD_DIR=tempfile.mkdtemp(prefix='drafting-test-uploads-'))
@mock.patch('drafting.views._dispatch')
class DraftingAmsInProcessTest(TestCase):
    databases = {'default'}

    def setUp(self):
        from core.testing import make_case, make_client as make_ams_client
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        self.junior = make_advocate(permissions=ALL_PERMISSIONS, parent_advocate_id=self.senior.id)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.ams_client = make_ams_client(self.senior, name='Sharma Traders')
        self.case = make_case(self.senior, self.ams_client, case_title='Sharma v. State')
        self.foreign_case = make_case(self.outsider, make_ams_client(self.outsider))

    def ams_document(self, owner, name='Plaint.pdf', body=b'%PDF-1.4 plaint'):
        import datetime, os
        from core.models import Document
        from django.conf import settings
        # Uploaded AMS files live in DOCUMENT_UPLOAD_DIR/documents (documents/storage.py).
        folder = os.path.join(settings.DOCUMENT_UPLOAD_DIR, 'documents')
        os.makedirs(folder, exist_ok=True)
        stored = f'ams-{owner.id}-{name}'
        path = os.path.join(folder, stored)
        with open(path, 'wb') as fh:
            fh.write(body)
        return Document.objects.create(document_name=name.rsplit('.', 1)[0], original_name=name, stored_name=stored,
                                       file_path=path, version=1, upload_date=datetime.datetime.now(),
                                       advocate_id=owner.id, case=self.case if owner == self.senior else None)

    def test_cases_are_the_practices(self, dispatch):
        body = self.client.get('/api/drafting/ams-cases/', **auth(self.junior)).json()
        self.assertEqual([c['id'] for c in body['content']], [self.case.id])

    def test_link_case_creates_project_once(self, dispatch):
        from workspace.models import CaseParty, CaseTask
        CaseParty.objects.create(advocate_id=self.senior.id, case_id=self.case.id, name='State of TN', is_opponent=True)
        task = CaseTask.objects.create(advocate_id=self.senior.id, case_id=self.case.id, title='Draft plaint')
        first = self.client.post('/api/drafting/link-case/', {'caseId': self.case.id, 'taskId': task.id},
                                 content_type='application/json', **auth(self.junior))
        self.assertEqual(first.status_code, 200, first.content)
        data = first.json()
        self.assertEqual(data['prefill']['party_a_name'], 'Sharma Traders')
        self.assertEqual(data['prefill']['party_b_name'], 'State of TN')
        self.assertEqual(data['task']['id'], task.id)
        again = self.client.post('/api/drafting/link-case/', {'caseId': self.case.id},
                                 content_type='application/json', **auth(self.senior)).json()
        self.assertEqual(again['projectId'], data['projectId'])
        self.assertEqual(Project.objects.get(id=data['projectId']).case_id, self.case.id)
        self.assertEqual(self.client.post('/api/drafting/link-case/', {'caseId': self.foreign_case.id},
                                          content_type='application/json', **auth(self.senior)).status_code, 404)

    def test_documents_list_and_import(self, dispatch):
        mine = self.ams_document(self.senior)
        self.ams_document(self.outsider, 'Theirs.pdf')
        self.ams_document(self.senior, 'photo.jpg')                  # not draftable
        rows = self.client.get('/api/drafting/ams-documents/', **auth(self.junior)).json()['content']
        self.assertEqual([r['id'] for r in rows], [mine.id])
        self.assertIsNone(rows[0]['sample'])
        resp = self.client.post(f'/api/drafting/ams-documents/{mine.id}/import/', **auth(self.junior))
        self.assertEqual(resp.status_code, 200, resp.content)
        sample = Sample.objects.get(id=resp.json()['id'])
        self.assertEqual((sample.ams_document_id, sample.uploaded_by_id), (mine.id, self.junior.id))
        Sample.objects.filter(id=sample.id).update(status='ready')     # processing finished
        sample.refresh_from_db()
        again = self.client.post(f'/api/drafting/ams-documents/{mine.id}/import/', **auth(self.senior)).json()
        self.assertEqual(again['id'], sample.id)                      # shared across the practice
        self.assertEqual(dispatch.call_count, 1)
        row = self.client.get('/api/drafting/ams-documents/', **auth(self.senior)).json()['content'][0]
        self.assertEqual(row['sample'], {'id': sample.id, 'status': sample.status, 'current': True})

    def test_stuck_pending_import_is_retried(self, dispatch):
        doc = self.ams_document(self.senior)
        first = self.client.post(f'/api/drafting/ams-documents/{doc.id}/import/', **auth(self.senior)).json()
        # Its job never ran (still pending): importing again re-dispatches it.
        again = self.client.post(f'/api/drafting/ams-documents/{doc.id}/import/', **auth(self.senior)).json()
        self.assertEqual(again['id'], first['id'])
        self.assertEqual(dispatch.call_count, 2)
        Sample.objects.filter(id=first['id']).update(status='processing')
        self.client.post(f'/api/drafting/ams-documents/{doc.id}/import/', **auth(self.senior))
        self.assertEqual(dispatch.call_count, 2)                      # in flight: left alone

    def test_files_need_login_and_practice(self, dispatch):
        resp = self.client.post('/api/drafting/samples/', {
            'file': SimpleUploadedFile('NDA.pdf', b'%PDF-1.4 secret'), 'name': 'NDA.pdf'}, **auth(self.junior))
        url = resp.json()['file']
        self.assertIn('/api/drafting/samples/', url)
        self.assertTrue(url.endswith('.pdf'))
        path = url.split('testserver', 1)[1]
        self.assertEqual(self.client.get(path).status_code, 401)
        self.assertEqual(self.client.get(path, **auth(self.outsider)).status_code, 404)
        ok = self.client.get(path, **auth(self.senior))
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(b''.join(ok.streaming_content), b'%PDF-1.4 secret')


# Merge phase 04: uploads pick an AMS case; the drafting client/project follow from it.
@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
@mock.patch('drafting.views._dispatch')
class DraftingUploadCaseTest(TestCase):
    databases = {'default'}

    def setUp(self):
        from core.testing import make_case, make_client as make_ams_client
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.case = make_case(self.senior, make_ams_client(self.senior, name='Sharma Traders'))
        self.foreign = make_case(self.outsider, make_ams_client(self.outsider))

    def upload(self, **extra):
        return self.client.post('/api/drafting/samples/', {
            'file': SimpleUploadedFile('Plaint.pdf', b'%PDF-1.4 x'), 'name': 'Plaint.pdf', **extra},
            **auth(self.senior))

    def test_case_sets_client_and_project(self, dispatch):
        resp = self.upload(case_id=self.case.id)
        self.assertEqual(resp.status_code, 201, resp.content)
        sample = Sample.objects.get(id=resp.json()['id'])
        self.assertEqual(sample.project.case_id, self.case.id)
        self.assertEqual(sample.client.name, 'Sharma Traders')

    def test_client_and_project_cannot_be_set_directly(self, dispatch):
        from drafting.ams_cases import link_case
        foreign_client, foreign_project = link_case(self.foreign)
        resp = self.upload(client=foreign_client.id, project=foreign_project.id)
        self.assertEqual(resp.status_code, 201, resp.content)
        sample = Sample.objects.get(id=resp.json()['id'])
        self.assertEqual((sample.client_id, sample.project_id), (None, None))

    def test_no_case_is_fine(self, dispatch):
        resp = self.upload()
        self.assertEqual(resp.status_code, 201, resp.content)
        self.assertIsNone(Sample.objects.get(id=resp.json()['id']).project_id)

    def test_other_firms_case_refused(self, dispatch):
        self.assertEqual(self.upload(case_id=self.foreign.id).status_code, 400)
        self.assertFalse(Sample.objects.exists())


# Merge phase 05: creating a draft session (the new-draft wizard's last call).
@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
@mock.patch('drafting.views._dispatch')
@mock.patch('drafting.serializers._dispatch', create=True)
class DraftSessionCreateTest(TestCase):
    databases = {'default'}

    def setUp(self):
        from core.testing import make_case, make_client as make_ams_client
        from drafting.ams_cases import link_case
        from workspace.models import CaseTask
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.template = Template.objects.create(name='Mutual NDA', document_type='nda', status='ready',
                                                created_by_id=self.senior.id)
        self.case = make_case(self.senior, make_ams_client(self.senior))
        self.client_row, self.project = link_case(self.case)
        self.task = CaseTask.objects.create(advocate_id=self.senior.id, case_id=self.case.id, title='Draft NDA')
        foreign_case = make_case(self.outsider, make_ams_client(self.outsider))
        _, self.foreign_project = link_case(foreign_case)

    def create(self, **body):
        data = {'template': self.template.id, 'samples': [], 'facts': {}, 'llm': 'gemini', 'mode': 'library', **body}
        return self.client.post('/api/drafting/draft-sessions/', data, content_type='application/json',
                                **auth(self.senior))

    def test_from_scratch_needs_no_member_or_case(self, *mocks):
        resp = self.create()
        self.assertEqual(resp.status_code, 201, resp.content)
        self.assertEqual(DraftSession.objects.get(id=resp.json()['id']).created_by_id, self.senior.id)

    def test_case_and_task_link(self, *mocks):
        resp = self.create(client=self.client_row.id, project=self.project.id, ams_task_id=self.task.id)
        self.assertEqual(resp.status_code, 201, resp.content)
        session = DraftSession.objects.get(id=resp.json()['id'])
        self.assertEqual((session.project.case_id, session.ams_task_id), (self.case.id, self.task.id))

    def test_other_firms_project_refused(self, *mocks):
        self.assertEqual(self.create(project=self.foreign_project.id).status_code, 400)

    def test_task_from_another_case_refused(self, *mocks):
        from workspace.models import CaseTask
        other = CaseTask.objects.create(advocate_id=self.senior.id, case_id=None, title='Unrelated')
        self.assertEqual(self.create(project=self.project.id, ams_task_id=other.id).status_code, 400)


# drafting/jobs.py: where background jobs run.
class JobDispatchTest(TestCase):
    databases = {'default'}

    def test_celery_when_not_eager(self):
        from drafting import jobs
        task = mock.Mock(name='task')
        with override_settings(CELERY_TASK_ALWAYS_EAGER=False):
            jobs.dispatch(task, 5)
        task.delay.assert_called_once_with(5)

    def test_worker_process_by_default(self):
        from drafting import jobs
        task = mock.Mock()
        task.name = 'drafting.tasks.process_sample'
        with override_settings(CELERY_TASK_ALWAYS_EAGER=True, DRAFTING_JOB_RUNNER='process'), \
                mock.patch.object(jobs, '_get_pool') as pool:
            jobs.dispatch(task, 7)
        pool.return_value.submit.assert_called_once_with(jobs._run_in_worker, 'drafting.tasks.process_sample', [7])
        task.apply.assert_not_called()                      # nothing ran in the web process

    def test_dead_worker_is_replaced_once(self):
        from concurrent.futures.process import BrokenProcessPool
        from drafting import jobs
        task = mock.Mock()
        task.name = 'drafting.tasks.process_sample'
        broken, healthy = mock.Mock(), mock.Mock()
        broken.submit.side_effect = BrokenProcessPool()
        with override_settings(CELERY_TASK_ALWAYS_EAGER=True, DRAFTING_JOB_RUNNER='process'), \
                mock.patch.object(jobs, '_get_pool', side_effect=[broken, healthy]), \
                mock.patch.object(jobs, '_reset_pool') as reset:
            jobs.dispatch(task, 9)
        reset.assert_called_once()
        healthy.submit.assert_called_once()


# Senior review of a junior's draft in the editor (drafting/access.py).
@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
@mock.patch('drafting.views._dispatch')
class DraftReviewAccessTest(TestCase):
    databases = {'default'}

    def setUp(self):
        from core.testing import make_case
        from drafting.ams_cases import link_case
        from drafting.models import DraftBlock
        from workspace.models import CaseTask
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        drafting = ('CASE_VIEW', 'DOCUMENT_VIEW', 'TASK_VIEW', 'DRAFT_VIEW', 'DRAFT_CREATE', 'DRAFT_EXPORT')
        self.junior = make_advocate(permissions=drafting, parent_advocate_id=self.senior.id)
        self.other_junior = make_advocate(permissions=drafting, parent_advocate_id=self.senior.id)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.case = make_case(self.senior)
        client, project = link_case(self.case)
        self.task = CaseTask.objects.create(advocate_id=self.senior.id, case_id=self.case.id, title='Draft plaint',
                                            assigned_to_id=self.junior.id, assigned_by_id=self.senior.id)
        template = Template.objects.create(name='Plaint', document_type='plaint')
        self.session = DraftSession.objects.create(template=template, client=client, project=project,
                                                   created_by_id=self.junior.id, ams_task_id=self.task.id,
                                                   facts={}, status='ready')
        self.block = DraftBlock.objects.create(session=self.session, position=0, block_type='clause',
                                               heading='Plaint', text='The plaintiff states.')
        self.url = f'/api/drafting/draft-sessions/{self.session.id}/'

    def set_status(self, value):
        self.task.review_status = value
        self.task.save(update_fields=['review_status'])

    def save_blocks(self, who):
        return self.client.post(self.url + 'save-blocks/', [{'id': self.block.id, 'text': 'Edited.'}],
                                content_type='application/json', **auth(who))

    def test_not_visible_before_submission(self, dispatch):
        self.assertEqual(self.client.get(self.url, **auth(self.senior)).status_code, 404)

    def test_reviewer_reads_exports_and_sees_review_block(self, dispatch):
        self.set_status('SUBMITTED')
        body = self.client.get(self.url, **auth(self.senior)).json()
        self.assertEqual(body['review']['taskId'], self.task.id)
        self.assertTrue(body['review']['canReview'])
        self.assertTrue(body['review']['canEdit'])
        self.assertFalse(body['review']['isOwner'])
        self.assertEqual(self.client.get(f'/api/drafting/drafts/{self.session.id}/export/docx/',
                                         **auth(self.senior)).status_code, 200)
        self.assertEqual(self.client.get(f'/api/drafting/drafts/{self.session.id}/ams-task/',
                                         **auth(self.senior)).json()['task']['id'], self.task.id)
        # The junior's own view: owner, cannot review their own work.
        own = self.client.get(self.url, **auth(self.junior)).json()['review']
        self.assertEqual((own['isOwner'], own['canReview']), (True, False))

    def test_others_cannot_see_it(self, dispatch):
        self.set_status('SUBMITTED')
        for who in (self.other_junior, self.outsider):
            self.assertEqual(self.client.get(self.url, **auth(who)).status_code, 404)
        # It is not in the reviewer's own draft list either (review is opened from the task).
        self.assertEqual(self.client.get('/api/drafting/draft-sessions/', **auth(self.senior)).json()['count'], 0)

    def test_reviewer_edits_only_while_submitted(self, dispatch):
        self.set_status('SUBMITTED')
        self.assertEqual(self.save_blocks(self.senior).status_code, 200)
        self.block.refresh_from_db()
        self.assertEqual(self.block.text, 'Edited.')
        self.set_status('APPROVED')
        self.assertEqual(self.save_blocks(self.senior).status_code, 403)
        self.assertEqual(self.save_blocks(self.junior).status_code, 200)       # the author still can

    def test_reviewer_never_files_redrafts_or_deletes(self, dispatch):
        self.set_status('SUBMITTED')
        h = auth(self.senior)
        self.assertEqual(self.client.post(f'/api/drafting/drafts/{self.session.id}/send-to-ams/', {},
                                          content_type='application/json', **h).status_code, 404)
        self.assertEqual(self.client.post(self.url + 'regenerate/', **h).status_code, 404)
        self.assertEqual(self.client.delete(self.url, **h).status_code, 404)

    def test_task_payload_links_the_draft(self, dispatch):
        tasks = self.client.get(f'/api/workspace/cases/{self.case.id}/tasks', **auth(self.senior)).json()
        self.assertEqual([t['draftSessionId'] for t in tasks if t['id'] == self.task.id], [self.session.id])


@mock.patch('drafting.views._dispatch')
class DraftVersionTest(TestCase):
    """Saved versions: frozen copies of a draft that redlines compare against."""
    databases = {'default'}

    def setUp(self):
        from .models import DraftBlock
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(session=self.session, position=0, block_type='clause',
                                               heading='Rent', text='Rent is Rs. 5,00,000.', source='generated')
        self.url = f'/api/drafting/draft-sessions/{self.session.id}/versions/'

    def save(self, who, label=''):
        return self.client.post(self.url, {'label': label}, content_type='application/json', **auth(who))

    def test_save_and_list_newest_first(self, dispatch):
        self.assertEqual(self.save(self.owner, 'Sent to client').status_code, 201)
        resp = self.save(self.owner)
        self.assertEqual([(v['number'], v['label']) for v in resp.json()], [(2, ''), (1, 'Sent to client')])
        self.assertEqual(self.client.get(self.url, **auth(self.owner)).json()[0]['number'], 2)

    def test_version_does_not_change_after_later_edits(self, dispatch):
        self.save(self.owner)
        self.block.text = 'Rent is Rs. 7,50,000.'
        self.block.save()
        version = self.session.versions.get()
        self.assertEqual(version.blocks[0]['text'], 'Rent is Rs. 5,00,000.')
        self.assertEqual(version.blocks[0]['block_id'], self.block.id)

    def test_empty_draft_refused(self, dispatch):
        self.block.delete()
        self.assertEqual(self.save(self.owner).status_code, 400)

    def test_other_advocates_cannot_see_or_save(self, dispatch):
        self.assertEqual(self.client.get(self.url, **auth(self.outsider)).status_code, 404)
        self.assertEqual(self.save(self.outsider).status_code, 404)


def _block(block_id, heading, html):
    return {'block_id': block_id, 'position': 0, 'block_type': 'clause', 'heading': heading,
            'text': '', 'content_html': html, 'style_json': {}}


def _redline_xml(before, after):
    import io
    import zipfile
    from .export.redline import render_redline_docx
    data, stats = render_redline_docx(before, after, author='A. Advocate')
    return zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml').decode(), stats


class RedlineDocxTest(TestCase):
    """Word tracked changes: <w:ins> / <w:del> around exactly the words that changed."""

    def test_changed_words_only(self):
        xml, stats = _redline_xml([_block(1, 'Rent', '<p>Monthly rent is Rs. 5,00,000 payable in advance.</p>')],
                                  [_block(1, 'Rent', '<p>Monthly rent is Rs. 7,50,000 payable in advance.</p>')])
        self.assertIn('<w:delText xml:space="preserve">5</w:delText>', xml)
        self.assertIn('w:author="A. Advocate"', xml)
        self.assertRegex(xml, r'<w:ins [^>]*>.*?<w:t xml:space="preserve">7</w:t>')
        self.assertIn('payable in advance', xml)                 # unchanged text stays plain
        self.assertEqual(stats, {'inserted': 2, 'deleted': 2})   # "5","00,000" → "7","50,000" as tokens

    def test_added_and_removed_clauses(self):
        xml, stats = _redline_xml(
            [_block(1, 'Rent', '<p>Rent.</p>'), _block(2, 'Arbitration', '<p>Disputes go to arbitration.</p>')],
            [_block(1, 'Rent', '<p>Rent.</p>'), _block(3, 'Jurisdiction', '<p>Courts at Chennai.</p>')])
        self.assertIn('<w:delText xml:space="preserve">Disputes go to arbitration.</w:delText>', xml)
        self.assertRegex(xml, r'<w:ins [^>]*><w:r><w:t xml:space="preserve">Courts at Chennai.</w:t>')
        self.assertGreater(stats['deleted'], 0)

    def test_redrafted_clause_paired_by_heading(self):
        # A re-draft gives the clause a new block id; the heading still pairs it, so only
        # the changed word is marked, not the whole clause.
        _, stats = _redline_xml([_block(1, 'Term', '<p>The term is three years.</p>')],
                                [_block(9, 'Term', '<p>The term is five years.</p>')])
        self.assertEqual(stats, {'inserted': 1, 'deleted': 1})

    def test_identical_versions_have_no_marks(self):
        blocks = [_block(1, 'Rent', '<p>Rent is <strong>due</strong> monthly.</p><ul><li><p>One</p></li></ul>')]
        xml, stats = _redline_xml(blocks, blocks)
        self.assertNotIn('<w:ins', xml)
        self.assertNotIn('<w:del ', xml)
        self.assertEqual(stats, {'inserted': 0, 'deleted': 0})


@mock.patch('drafting.views._dispatch')
class RedlineExportViewTest(TestCase):
    databases = {'default'}

    def setUp(self):
        from .models import DraftBlock
        from .versions import save_version
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(session=self.session, position=0, block_type='clause',
                                               heading='Rent', text='Rent is Rs. 5,00,000.', source='generated')
        self.url = f'/api/drafting/drafts/{self.session.id}/export/redline/'
        self.save_version = save_version

    def test_needs_a_version(self, dispatch):
        self.assertEqual(self.client.get(self.url, **auth(self.owner)).status_code, 400)

    def test_compares_last_version_with_current(self, dispatch):
        v1 = self.save_version(self.session)
        self.block.text = 'Rent is Rs. 7,50,000.'
        self.block.save()
        resp = self.client.get(self.url, **auth(self.owner))
        self.assertEqual(resp.status_code, 200)
        self.assertIn(f'redline_v{v1.number}-current', resp['Content-Disposition'])
        self.assertEqual((resp['X-Redline-Inserted'], resp['X-Redline-Deleted']), ('2', '2'))

    def test_other_sessions_versions_refused(self, dispatch):
        other = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        from .models import DraftBlock
        DraftBlock.objects.create(session=other, position=0, block_type='clause', text='x', source='generated')
        foreign = self.save_version(other)
        self.assertEqual(self.client.get(self.url + f'?from={foreign.id}', **auth(self.owner)).status_code, 404)

    def test_outsider_cannot_export(self, dispatch):
        self.save_version(self.session)
        outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.assertEqual(self.client.get(self.url, **auth(outsider)).status_code, 404)


@mock.patch('drafting.views._dispatch')
class PdfExportTest(TestCase):
    """?output=pdf converts the Word export with LibreOffice; 503 when it is missing."""
    databases = {'default'}

    def setUp(self):
        from .models import DraftBlock
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        DraftBlock.objects.create(session=self.session, position=0, block_type='clause',
                                  heading='Rent', text='Rent is due.', source='generated')
        self.base = f'/api/drafting/drafts/{self.session.id}/export/'

    @override_settings(LIBREOFFICE_PATH=r'C:\nowhere\soffice.exe')
    def test_503_without_libreoffice(self, dispatch):
        resp = self.client.get(self.base + 'docx/?output=pdf', **auth(self.owner))
        self.assertEqual(resp.status_code, 503)
        # The Word export is unaffected.
        self.assertEqual(self.client.get(self.base + 'docx/', **auth(self.owner)).status_code, 200)

    @mock.patch('drafting.export.pdf.docx_to_pdf', return_value=b'%PDF-1.7 converted')
    def test_pdf_of_draft_and_redline(self, convert, dispatch):
        from .versions import save_version
        resp = self.client.get(self.base + 'docx/?output=pdf', **auth(self.owner))
        self.assertEqual((resp.status_code, resp['Content-Type']), (200, 'application/pdf'))
        self.assertTrue(resp['Content-Disposition'].endswith('.pdf"'))
        self.assertEqual(resp.content, b'%PDF-1.7 converted')
        self.assertTrue(convert.call_args[0][0].startswith(b'PK'))   # it converted a .docx
        v = save_version(self.session)
        resp = self.client.get(self.base + f'redline/?from={v.id}&output=pdf', **auth(self.owner))
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertIn('X-Redline-Inserted', resp)


class CompareEngineTest(TestCase):
    """export/compare.py: the data behind the on-screen compare view and the redline."""

    def test_changes_numbered_and_marked(self):
        from .export.compare import compare_blocks
        before = [_block(1, 'Rent', '<p>Rent is Rs. 40,000.</p><p>Paid monthly.</p>'),
                  _block(2, 'Sub-letting', '<p>The Lessee may sub-let.</p>')]
        after = [_block(1, 'Rent', '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>'),
                 _block(3, 'Lock-in', '<p>Six months lock-in.</p>')]
        data = compare_blocks(before, after).as_json()
        self.assertEqual([c['status'] for c in data['clauses']], ['changed', 'removed', 'added'])
        rent = data['clauses'][0]['body']
        self.assertEqual(rent[0]['kind'], 'changed')
        self.assertEqual([(s['op'], s['text']) for s in rent[0]['segs'] if s['op'] != 'same'],
                         [('del', '40'), ('ins', '45')])
        self.assertEqual(rent[1]['kind'], 'same')
        self.assertIsNone(rent[1]['change_id'])
        # Every changed paragraph is numbered in order: rent, removed heading + body, added heading + body.
        ids = [p['change_id'] for c in data['clauses'] for p in c['heading'] + c['body'] if p['change_id']]
        self.assertEqual(ids, [1, 2, 3, 4, 5])
        self.assertEqual(data['changes'], 5)

    def test_screen_and_word_count_the_same(self):
        from .export.compare import compare_blocks
        before = [_block(1, 'Term', '<p>The term is three years from today.</p>')]
        after = [_block(1, 'Term', '<p>The term is five years from the start date.</p>')]
        data = compare_blocks(before, after).as_json()
        _, stats = _redline_xml(before, after)
        self.assertEqual((data['inserted'], data['deleted']), (stats['inserted'], stats['deleted']))


@mock.patch('drafting.views._dispatch')
class CompareViewTest(TestCase):
    databases = {'default'}

    def setUp(self):
        from .models import DraftBlock
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(session=self.session, position=0, block_type='clause',
                                               heading='Rent', text='Rent is Rs. 40,000.', source='generated')
        self.url = f'/api/drafting/draft-sessions/{self.session.id}/compare/'

    def test_needs_a_version(self, dispatch):
        self.assertEqual(self.client.get(self.url, **auth(self.owner)).status_code, 400)

    def test_version_against_current(self, dispatch):
        from .versions import save_version
        v1 = save_version(self.session)
        self.block.text = 'Rent is Rs. 45,000.'
        self.block.save()
        body = self.client.get(self.url, **auth(self.owner)).json()
        self.assertEqual(body['from']['id'], v1.id)
        self.assertIsNone(body['to'])
        self.assertEqual((body['inserted'], body['deleted'], body['changes']), (1, 1, 1))   # 40 -> 45; ',000' unchanged

    def test_outsider_refused(self, dispatch):
        from .versions import save_version
        save_version(self.session)
        outsider = make_advocate(permissions=ALL_PERMISSIONS)
        self.assertEqual(self.client.get(self.url, **auth(outsider)).status_code, 404)


DRAFTER = ('DRAFT_VIEW', 'DRAFT_CREATE', 'DRAFT_EXPORT')


class AuthorityTest(TestCase):
    """drafting/authority.py: who is senior over whose drafts."""

    def test_rule(self):
        from .authority import is_senior_over
        head = make_advocate(permissions=DRAFTER)
        junior = make_advocate(permissions=DRAFTER, parent_advocate_id=head.id)
        peer = make_advocate(permissions=DRAFTER, parent_advocate_id=head.id)
        assigner = make_advocate(permissions=DRAFTER + ('TASK_ASSIGN',), parent_advocate_id=head.id)
        outsider = make_advocate(permissions=DRAFTER + ('TASK_ASSIGN',))
        self.assertTrue(is_senior_over(head, junior))          # team head
        self.assertTrue(is_senior_over(assigner, junior))      # can assign tasks in the team
        self.assertFalse(is_senior_over(peer, junior))         # a peer
        self.assertFalse(is_senior_over(junior, head))         # nobody is over the team head
        self.assertFalse(is_senior_over(assigner, head))
        self.assertFalse(is_senior_over(outsider, junior))     # another team
        self.assertFalse(is_senior_over(head, head))


class HtmlWriteTest(TestCase):
    """export/htmlwrite.py turns paragraphs back into the HTML the editor reads."""

    def test_round_trip(self):
        from .export.docx import html_to_blocks
        from .export.htmlwrite import paras_to_html
        src = ('<h2>Rent</h2><p style="text-align: justify">Rent is <strong>Rs. 40,000</strong>, <em>monthly</em>.</p>'
               '<ul><li><p>One</p><ul><li><p>One a</p></li></ul></li><li><p>Two</p></li></ul>'
               '<ol><li><p>First</p></li></ol><p>Line<br>break &amp; more</p>')
        paras = html_to_blocks(src)
        again = html_to_blocks(paras_to_html(paras))
        self.assertEqual([(p.style, p.align, p.runs) for p in again], [(p.style, p.align, p.runs) for p in paras])


@mock.patch('drafting.views._dispatch')
class ReviewRoundTest(TestCase):
    """drafting/review.py through the API: suggest, decide, finish (docs/DRAFT_REVIEW.md)."""
    databases = {'default'}

    def setUp(self):
        from .models import DraftBlock, DraftReviewRequest
        self.head = make_advocate(permissions=DRAFTER)
        self.junior = make_advocate(permissions=DRAFTER, parent_advocate_id=self.head.id)
        self.outsider = make_advocate(permissions=DRAFTER)
        # The senior's own draft; the junior is asked to proofread it (suggestions only).
        self.session = DraftSession.objects.create(created_by_id=self.head.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(
            session=self.session, position=0, block_type='clause', heading='Rent', text='',
            content_html='<p>Rent is Rs. 40,000.</p><p>Paid monthly.</p>', source='generated')
        self.request = DraftReviewRequest.objects.create(
            session=self.session, requested_by_id=self.head.id, reviewer_id=self.junior.id,
            authority=DraftReviewRequest.Authority.SUGGEST)
        self.base = '/api/drafting/'

    def suggest(self, who, html):
        return self.client.post(f'{self.base}draft-sessions/{self.session.id}/suggest/',
                                {'blocks': [{'id': self.block.id, 'heading': 'Rent', 'content_html': html, 'text': ''}],
                                 'note': 'Please check'}, content_type='application/json', **auth(who))

    def decide(self, rid, who, change, decision, reason=''):
        return self.client.post(f'{self.base}rounds/{rid}/decide/',
                                {'change': change, 'decision': decision, 'reason': reason},
                                content_type='application/json', **auth(who))

    def finish(self, rid, who):
        return self.client.post(f'{self.base}rounds/{rid}/finish/', **auth(who))

    def test_requested_reviewer_sees_but_cannot_edit_directly(self, dispatch):
        url = f'{self.base}draft-sessions/{self.session.id}/'
        self.assertEqual(self.client.get(url, **auth(self.junior)).status_code, 200)
        resp = self.client.post(url + 'save-blocks/', [{'id': self.block.id, 'text': 'x'}],
                                content_type='application/json', **auth(self.junior))
        self.assertEqual(resp.status_code, 403)                     # suggest-only
        self.assertEqual(self.client.get(url, **auth(self.outsider)).status_code, 404)

    def test_binding_reviewer_may_edit(self, dispatch):
        from .models import DraftReviewRequest
        junior_draft = DraftSession.objects.create(created_by_id=self.junior.id, facts={}, status='ready')
        DraftReviewRequest.objects.create(session=junior_draft, requested_by_id=self.junior.id,
                                          reviewer_id=self.head.id, authority=DraftReviewRequest.Authority.BINDING)
        resp = self.client.post(f'{self.base}draft-sessions/{junior_draft.id}/save-blocks/', [],
                                content_type='application/json', **auth(self.head))
        self.assertEqual(resp.status_code, 200)

    def test_mixed_decisions_in_one_clause(self, dispatch):
        resp = self.suggest(self.junior, '<p>Rent is Rs. 45,000.</p><p>Paid monthly in advance.</p>')
        self.assertEqual(resp.status_code, 201, resp.content)
        rnd = resp.json()
        self.assertEqual((rnd['kind'], rnd['pending'], rnd['can_decide']), ('suggestions', 2, False))
        self.block.refresh_from_db()
        self.assertIn('40,000', self.block.content_html)            # nothing applied yet
        rid = rnd['id']
        self.assertEqual(self.decide(rid, self.junior, 1, 'accepted').status_code, 400)   # not the decider
        self.assertEqual(self.decide(rid, self.head, 2, 'declined').status_code, 400)     # reason needed
        self.assertEqual(self.decide(rid, self.head, 1, 'rejected').status_code, 400)     # wrong kind
        self.assertEqual(self.decide(rid, self.head, 1, 'accepted').status_code, 200)
        self.assertEqual(self.finish(rid, self.head).status_code, 400)                    # one still pending
        self.assertEqual(self.decide(rid, self.head, 2, 'declined', 'Advance is in 2.2').status_code, 200)
        resp = self.finish(rid, self.head)
        self.assertEqual(resp.status_code, 200, resp.content)
        self.block.refresh_from_db()
        self.assertIn('45,000', self.block.content_html)            # accepted
        self.assertNotIn('in advance', self.block.content_html)     # declined
        self.assertEqual(self.block.text, 'Rent is Rs. 45,000.\nPaid monthly.')
        version = self.session.versions.order_by('-number').first()
        self.assertEqual((version.kind, version.label), ('review', 'Review: 1 accepted, 1 declined'))

    def test_accept_all_keeps_the_html_exactly(self, dispatch):
        html = '<p>Rent is Rs. 45,000.</p><p>Paid <strong>monthly</strong>.</p>'
        rid = self.suggest(self.junior, html).json()['id']
        self.assertEqual(self.decide(rid, self.head, 'all', 'accepted').status_code, 200)
        self.finish(rid, self.head)
        self.block.refresh_from_db()
        self.assertEqual(self.block.content_html, html)

    def test_outdated_suggestion_is_not_applied(self, dispatch):
        rid = self.suggest(self.junior, '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>').json()['id']
        self.block.content_html = '<p>Rent is Rs. 50,000.</p><p>Paid monthly.</p>'   # the owner moved on
        self.block.save()
        body = self.client.get(f'{self.base}rounds/{rid}/', **auth(self.head)).json()
        self.assertTrue(body['clauses'][0]['body'][0]['outdated'])
        resp = self.finish(rid, self.head)                         # outdated changes need no decision
        self.assertEqual(resp.json()['result']['outdated'], [1])   # change 1: its paragraph changed
        self.block.refresh_from_db()
        self.assertIn('50,000', self.block.content_html)

    def test_edit_elsewhere_in_the_clause_keeps_suggestion_valid(self, dispatch):
        # A whole deed often sits in one clause: an edit to another paragraph must not outdate it.
        rid = self.suggest(self.junior, '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>').json()['id']
        self.block.content_html = '<p>Rent is Rs. 40,000.</p><p>Paid quarterly.</p>'   # owner edits paragraph 2
        self.block.save()
        body = self.client.get(f'{self.base}rounds/{rid}/', **auth(self.head)).json()
        self.assertFalse(body['clauses'][0]['body'][0]['outdated'])
        self.decide(rid, self.head, 1, 'accepted')
        self.assertEqual(self.finish(rid, self.head).status_code, 200)
        self.block.refresh_from_db()
        self.assertEqual(self.block.text, 'Rent is Rs. 45,000.\nPaid quarterly.')   # both edits kept

    def test_only_formatting_is_refused(self, dispatch):
        resp = self.suggest(self.junior, '<p>Rent is <strong>Rs. 40,000.</strong></p><p>Paid monthly.</p>')
        self.assertEqual(resp.status_code, 400)

    def test_owner_cannot_suggest_and_outsider_cannot_see(self, dispatch):
        self.assertEqual(self.suggest(self.head, '<p>x</p>').status_code, 403)
        rid = self.suggest(self.junior, '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>').json()['id']
        self.assertEqual(self.client.get(f'{self.base}rounds/{rid}/', **auth(self.outsider)).status_code, 404)

    def test_author_withdraws(self, dispatch):
        rid = self.suggest(self.junior, '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>').json()['id']
        self.assertEqual(self.client.post(f'{self.base}rounds/{rid}/cancel/', **auth(self.head)).status_code, 400)
        resp = self.client.post(f'{self.base}rounds/{rid}/cancel/', **auth(self.junior))
        self.assertEqual(resp.json()['status'], 'cancelled')

    def test_changes_round_reject_puts_old_text_back(self, dispatch):
        from . import review
        from .versions import snapshot_blocks
        base = snapshot_blocks(self.session)
        self.block.content_html = '<p>Rent is Rs. 45,000.</p><p>Paid weekly.</p>'
        self.block.save()
        rnd = review.create_changes_round(self.session, base, author_id=self.junior.id, decider_id=self.head.id)
        review.decide(rnd, self.head, 1, 'accepted')
        review.decide(rnd, self.head, 2, 'rejected')
        review.finish(rnd, self.head)
        self.block.refresh_from_db()
        self.assertEqual(self.block.text, 'Rent is Rs. 45,000.\nPaid monthly.')

    def test_binding_round_is_only_acknowledged(self, dispatch):
        from . import review
        from .versions import snapshot_blocks
        base = snapshot_blocks(self.session)
        self.block.content_html = '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>'
        self.block.save()
        rnd = review.create_changes_round(self.session, base, author_id=self.head.id,
                                          decider_id=self.junior.id, binding=True)
        with self.assertRaises(review.ReviewError):
            review.decide(rnd, self.junior, 1, 'rejected')
        review.decide(rnd, self.junior, 1, 'queried', 'Client agreed 40,000?')
        review.finish(rnd, self.junior)
        self.block.refresh_from_db()
        self.assertIn('45,000', self.block.content_html)          # a query never undoes a correction


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
@mock.patch('drafting.views._dispatch')
class TaskReviewLoopTest(TestCase):
    """drafting/task_review.py: the whole task loop with review rounds (docs/DRAFT_REVIEW.md, situation 1)."""
    databases = {'default'}

    def setUp(self):
        from core.testing import make_case
        from drafting.ams_cases import link_case
        from drafting.models import DraftBlock
        from workspace.models import CaseTask
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        drafting = ('CASE_VIEW', 'DOCUMENT_VIEW', 'TASK_VIEW', 'DRAFT_VIEW', 'DRAFT_CREATE', 'DRAFT_EXPORT')
        self.junior = make_advocate(permissions=drafting, parent_advocate_id=self.senior.id)
        self.case = make_case(self.senior)
        client, project = link_case(self.case)
        self.task = CaseTask.objects.create(advocate_id=self.senior.id, case_id=self.case.id, title='Draft lease',
                                            assigned_to_id=self.junior.id, assigned_by_id=self.senior.id)
        self.session = DraftSession.objects.create(client=client, project=project, created_by_id=self.junior.id,
                                                   ams_task_id=self.task.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(
            session=self.session, position=0, block_type='clause', heading='Rent', text='',
            content_html='<p>Rent is Rs. 40,000.</p><p>Notice is one month.</p><p>Paid monthly.</p>', source='generated')
        self.s = f'/api/drafting/draft-sessions/{self.session.id}/'

    def post(self, path, who, body=None):
        return self.client.post(path, body or {}, content_type='application/json', **auth(who))

    def save(self, who, html):
        return self.post(self.s + 'save-blocks/', who, [{'id': self.block.id, 'heading': 'Rent', 'content_html': html, 'text': ''}])

    def submit(self, note=''):
        return self.post(f'/api/drafting/drafts/{self.session.id}/send-to-ams/', self.junior, {'note': note} if note else {})

    def review(self, action, note=''):
        return self.post(f'/api/workspace/tasks/{self.task.id}/review', self.senior, {'action': action, 'note': note})

    def rounds(self, who):
        return self.client.get(self.s + 'rounds/', **auth(who)).json()

    def test_full_loop(self, dispatch):
        from . import review
        from .models import DraftComment, DraftReviewRound
        self.assertEqual(self.submit().status_code, 200)

        # Senior: no suggestions in a task review (docs/DRAFT_REVIEW.md, "Revision") ...
        resp = self.post(self.s + 'suggest/', self.senior, {'blocks': [{'id': self.block.id, 'heading': 'Rent',
            'content_html': '<p>x</p>', 'text': ''}]})
        self.assertEqual(resp.status_code, 403)
        # ... he corrects directly, and asks a question as a comment.
        self.assertEqual(self.save(self.senior,
                                   '<p>Rent is Rs. 45,000.</p><p>Notice is one month.</p><p>Paid monthly.</p>').status_code, 200)
        resp = self.post(self.s + 'comments/', self.senior, {'block_id': self.block.id, 'quote': 'Paid monthly.',
                                                             'body': 'Did the client agree to rent in advance?'})
        question = resp.json()['threads'][0]['id']
        self.assertEqual(self.review('approve').status_code, 400)          # an open question
        self.assertEqual(self.review('request_changes', 'Fix the notice period.').status_code, 200)

        # Junior: the correction arrives as a binding round (OK / Query only).
        binding = DraftReviewRound.objects.get(session=self.session, binding=True)
        self.assertEqual((binding.decider_id, binding.author_id), (self.junior.id, self.senior.id))
        with self.assertRaises(review.ReviewError):
            review.decide(binding, self.junior, 1, 'rejected')
        # Her query opens a comment thread on the corrected words, for the senior.
        review.decide(binding, self.junior, 1, 'queried', 'Client agreed 40,000?')
        query = DraftComment.objects.get(session=self.session, author_id=self.junior.id, parent=None)
        self.assertEqual((query.body, query.block_id), ('Client agreed 40,000?', self.block.id))
        self.assertIn('45,000', query.quote)
        # She answers his question, fixes what was asked and resubmits (comments don't block a resubmit).
        self.post(f'/api/drafting/comments/{question}/reply/', self.junior, {'body': 'Arrears, as agreed.'})
        self.block.refresh_from_db()
        self.assertEqual(self.save(self.junior, self.block.content_html.replace('one month', 'two months')).status_code, 200)
        self.assertEqual(self.submit('Fixed notice').status_code, 200)
        binding.refresh_from_db()
        self.assertEqual(binding.status, 'finished')                       # closed, query kept

        # Senior: a Keep / Reject round of her change since "Sent back", and two open threads.
        changes = DraftReviewRound.objects.get(session=self.session, kind='changes', binding=False, status='open')
        self.assertEqual((changes.decider_id, review.comparison(changes).changes), (self.senior.id, 1))
        threads = self.client.get(self.s + 'comments/', **auth(self.senior)).json()
        self.assertEqual((threads['open'], threads['answered']), (2, 1))
        review.decide(changes, self.senior, 'all', 'accepted')
        review.finish(changes, self.senior)
        self.assertEqual(self.review('approve').status_code, 400)          # threads still open
        for t in threads['threads']:
            self.assertEqual(self.post(f'/api/drafting/comments/{t["id"]}/resolve/', self.senior).status_code, 200)
        self.assertEqual(self.review('approve').status_code, 200)
        self.task.refresh_from_db()
        self.assertEqual(self.task.review_status, 'APPROVED')
        kinds = list(self.session.versions.order_by('number').values_list('kind', flat=True))
        self.assertIn('returned', kinds)


@mock.patch('drafting.views._dispatch')
class ReviewRequestTest(TestCase):
    """drafting/review_requests.py: Request review on drafts without a task (situations 2-4)."""
    databases = {'default'}

    def setUp(self):
        from .models import DraftBlock
        self.head = make_advocate(permissions=DRAFTER)
        self.junior = make_advocate(permissions=DRAFTER, parent_advocate_id=self.head.id)
        self.peer = make_advocate(permissions=DRAFTER, parent_advocate_id=self.head.id)
        self.outsider = make_advocate(permissions=DRAFTER)
        self.session = DraftSession.objects.create(created_by_id=self.junior.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(
            session=self.session, position=0, block_type='clause', heading='Rent', text='',
            content_html='<p>Rent is Rs. 40,000.</p><p>Paid monthly.</p>', source='generated')
        self.s = f'/api/drafting/draft-sessions/{self.session.id}/'

    def post(self, path, who, body=None):
        return self.client.post(path, body or {}, content_type='application/json', **auth(who))

    def ask(self, reviewer, note='Please check the rent'):
        return self.post(self.s + 'review-requests/', self.junior, {'reviewer': reviewer.id, 'note': note})

    def test_reviewers_and_their_power(self, dispatch):
        rows = {r['id']: r['authority'] for r in self.client.get(self.s + 'reviewers/', **auth(self.junior)).json()}
        self.assertEqual(rows, {self.head.id: 'binding', self.peer.id: 'suggest'})   # never the outsider

    def test_request_rules(self, dispatch):
        self.assertEqual(self.ask(self.outsider).status_code, 400)                    # not in the team
        self.assertEqual(self.post(self.s + 'review-requests/', self.peer, {'reviewer': self.head.id}).status_code, 404)
        resp = self.ask(self.head)
        self.assertEqual((resp.status_code, resp.json()['authority']), (201, 'binding'))
        self.assertEqual(self.ask(self.head).status_code, 400)                        # already reviewing
        self.assertEqual(self.session.versions.get().label, f'Sent for review to {self.head.full_name}')
        from core.models import NotificationQueue
        self.assertTrue(NotificationQueue.objects.filter(type='DRAFT_REVIEW_REQUESTED').exists())
        task_draft = DraftSession.objects.create(created_by_id=self.junior.id, facts={}, status='ready', ams_task_id=99)
        resp = self.post(f'/api/drafting/draft-sessions/{task_draft.id}/review-requests/', self.junior, {'reviewer': self.head.id})
        self.assertEqual(resp.status_code, 400)                                        # tasks use their own review

    def test_senior_corrects_then_owner_acknowledges(self, dispatch):
        rid = self.ask(self.head).json()['id']
        self.assertEqual([r['id'] for r in self.client.get('/api/drafting/drafts/for-review/', **auth(self.head)).json()], [rid])
        session = self.client.get(self.s, **auth(self.head)).json()
        self.assertEqual((session['access']['canWrite'], session['access']['canSuggest']), (True, True))
        resp = self.post(self.s + 'save-blocks/', self.head, [{'id': self.block.id, 'heading': 'Rent',
                         'content_html': '<p>Rent is Rs. 45,000.</p><p>Paid monthly.</p>', 'text': ''}])
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(self.post(f'/api/drafting/review-requests/{rid}/done/', self.head).status_code, 200)
        self.assertEqual(self.client.get('/api/drafting/drafts/for-review/', **auth(self.head)).json(), [])
        rounds = self.client.get(self.s + 'rounds/', **auth(self.junior)).json()
        self.assertEqual([(r['binding'], r['changes'], r['can_decide']) for r in rounds], [(True, 1, True)])
        # After Done the senior can still read the draft but no longer edit it.
        self.assertEqual(self.client.get(self.s, **auth(self.head)).status_code, 200)
        self.assertEqual(self.post(self.s + 'save-blocks/', self.head, []).status_code, 403)

    def test_peer_suggests_owner_decides(self, dispatch):
        rid = self.ask(self.peer).json()['id']
        session = self.client.get(self.s, **auth(self.peer)).json()
        self.assertEqual((session['access']['canWrite'], session['access']['canSuggest']), (False, True))
        resp = self.post(self.s + 'suggest/', self.peer, {'blocks': [{'id': self.block.id, 'heading': 'Rent',
                         'content_html': '<p>Rent is Rs. 40,000.</p><p>Paid monthly in advance.</p>', 'text': ''}]})
        self.assertEqual(resp.status_code, 201)
        self.post(f'/api/drafting/review-requests/{rid}/done/', self.peer)
        rounds = self.client.get(self.s + 'rounds/', **auth(self.junior)).json()
        self.assertEqual([(r['kind'], r['binding'], r['can_decide']) for r in rounds], [('suggestions', False, True)])

    def test_owner_cancels(self, dispatch):
        rid = self.ask(self.peer).json()['id']
        sug = self.post(self.s + 'suggest/', self.peer, {'blocks': [{'id': self.block.id, 'heading': 'Rent',
                        'content_html': '<p>Rent is Rs. 41,000.</p><p>Paid monthly.</p>', 'text': ''}]}).json()['id']
        self.assertEqual(self.post(f'/api/drafting/review-requests/{rid}/cancel/', self.peer).status_code, 404)
        self.assertEqual(self.post(f'/api/drafting/review-requests/{rid}/cancel/', self.junior).json()['status'], 'cancelled')
        self.assertEqual(self.client.get(self.s, **auth(self.peer)).status_code, 404)      # access gone
        from .models import DraftReviewRound
        self.assertEqual(DraftReviewRound.objects.get(id=sug).status, 'cancelled')


    def test_progress_shown_to_both_sides(self, dispatch):
        from . import review
        from core.models import NotificationQueue
        from .models import DraftReviewRound
        status = lambda: self.client.get('/api/drafting/drafts/review-status/', **auth(self.junior)).json()
        progress = lambda: self.client.get('/api/drafting/drafts/for-review/', **auth(self.peer)).json()[0]['progress']
        self.assertEqual(status(), {})                                                   # never reviewed
        rid = self.ask(self.peer).json()['id']
        self.assertEqual(status()[str(self.session.id)]['state'], 'with_reviewer')
        self.post(self.s + 'suggest/', self.peer, {'blocks': [{'id': self.block.id, 'heading': 'Rent',
                  'content_html': '<p>Rent is Rs. 40,000.</p><p>Paid monthly in advance.</p>', 'text': ''}]})
        self.assertEqual(status()[str(self.session.id)], {'state': 'to_decide', 'who': self.peer.full_name})
        self.assertEqual((progress()['sent'], progress()['waiting']), (1, 1))
        rnd = DraftReviewRound.objects.get(session=self.session)
        review.decide(rnd, self.junior, 'all', 'accepted')
        review.finish(rnd, self.junior)
        self.assertEqual((progress()['waiting'], progress()['accepted']), (0, 1))
        told = NotificationQueue.objects.filter(type='DRAFT_REVIEW_DONE', advocate_id=self.peer.id)
        self.assertTrue(told.exists())
        self.assertEqual(status()[str(self.session.id)]['state'], 'with_reviewer')
        self.post(f'/api/drafting/review-requests/{rid}/done/', self.peer)
        self.assertEqual(status()[str(self.session.id)]['state'], 'reviewed')


class DecideAllTest(TestCase):
    """'all' only fills in undecided changes."""

    def test_all_keeps_earlier_decisions(self):
        from . import review
        from .models import DraftBlock
        from .versions import snapshot_blocks
        owner = make_advocate(permissions=DRAFTER)
        senior = make_advocate(permissions=DRAFTER)
        session = DraftSession.objects.create(created_by_id=owner.id, facts={}, status='ready')
        block = DraftBlock.objects.create(session=session, position=0, block_type='clause', heading='Rent', text='',
                                          content_html='<p>Rent 40.</p><p>Notice 1.</p>', source='generated')
        base = snapshot_blocks(session)
        block.content_html = '<p>Rent 45.</p><p>Notice 2.</p>'
        block.save()
        rnd = review.create_changes_round(session, base, author_id=senior.id, decider_id=owner.id, binding=True)
        review.decide(rnd, owner, 1, 'queried', 'Why 45?')
        review.decide(rnd, owner, 'all', 'acknowledged')
        self.assertEqual(dict(rnd.decisions.values_list('change_id', 'decision')), {1: 'queried', 2: 'acknowledged'})


class CompareStorageNoiseTest(TestCase):
    """The AI draft (plain text) and an edited draft (editor HTML) store the same things differently;
    comparing them must show only real edits (draft #121: 11 reported changes, 2 real)."""

    def compare(self, before, after, title='Residential Lease Deed'):
        from .export.compare import compare_blocks
        return compare_blocks(before, after, title)

    def block(self, heading, text='', html=''):
        return {'block_id': 1, 'position': 0, 'block_type': 'clause', 'heading': heading,
                'text': text, 'content_html': html, 'style_json': {}}

    def test_empty_placeholder_and_title_line_are_not_changes(self):
        ai = self.block('', text='Residential Lease Deed\nCommencing from [[Lease Start Date]].')
        edited = self.block('', html='<p>Commencing from <span data-placeholder="Lease Start Date" '
                                     'data-value="">[Lease Start Date]</span>.</p>')
        self.assertEqual(self.compare([ai], [edited]).changes, 0)

    def test_filled_placeholder_is_one_change(self):
        ai = self.block('', text='Executed on [[Date of Execution]].')
        edited = self.block('', html='<p>Executed on <span data-placeholder="Date of Execution" '
                                     'data-value="22/03/2026">22/03/2026</span>.</p>')
        comp = self.compare([ai], [edited])
        self.assertEqual(comp.changes, 1)
        segs = [(s.op, s.text) for s in comp.clauses[0].body[0].segs if s.op != 'same']
        self.assertIn(('ins', '22/03/2026'), segs)

    def test_repeated_clause_heading_line_is_not_a_change(self):
        ai = self.block('SCHEDULE OF PROPERTY', text='SCHEDULE OF PROPERTY\nAll that flat.')
        edited = self.block('SCHEDULE OF PROPERTY', html='<p>All that flat.</p>')
        self.assertEqual(self.compare([ai], [edited]).changes, 0)

    def test_double_escaped_ampersand_in_placeholder_name(self):
        ai = self.block('', text='Witness: [[Witness 2 Name & Address]]')
        edited = self.block('', html='<p>Witness: <span data-placeholder="Witness 2 Name &amp;amp; Address" '
                                     'data-value="">[Witness 2 Name &amp;amp; Address]</span></p>')
        self.assertEqual(self.compare([ai], [edited]).changes, 0)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
@mock.patch('drafting.views._dispatch')
class DraftCommentTest(TestCase):
    """drafting/comments.py: questions on a passage, replies, resolving; Approve waits for them."""
    databases = {'default'}

    def setUp(self):
        from core.testing import make_case
        from drafting.ams_cases import link_case
        from drafting.models import DraftBlock
        from workspace.models import CaseTask
        self.senior = make_advocate(permissions=ALL_PERMISSIONS)
        drafting = ('CASE_VIEW', 'DOCUMENT_VIEW', 'TASK_VIEW', 'DRAFT_VIEW', 'DRAFT_CREATE', 'DRAFT_EXPORT')
        self.junior = make_advocate(permissions=drafting, parent_advocate_id=self.senior.id)
        self.peer = make_advocate(permissions=drafting, parent_advocate_id=self.senior.id)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        case = make_case(self.senior)
        client, project = link_case(case)
        self.task = CaseTask.objects.create(advocate_id=self.senior.id, case_id=case.id, title='Draft lease',
                                            assigned_to_id=self.junior.id, assigned_by_id=self.senior.id)
        self.session = DraftSession.objects.create(client=client, project=project, created_by_id=self.junior.id,
                                                   ams_task_id=self.task.id, facts={}, status='ready')
        self.block = DraftBlock.objects.create(session=self.session, position=0, block_type='clause', heading='Rent',
                                               text='', content_html='<p>Rent is paid monthly.</p>', source='generated')
        self.s = f'/api/drafting/draft-sessions/{self.session.id}/'

    def post(self, path, who, body=None):
        return self.client.post(path, body or {}, content_type='application/json', **auth(who))

    def comment(self, who, body='Did the client agree to rent in advance?'):
        return self.post(self.s + 'comments/', who, {'block_id': self.block.id, 'quote': 'paid monthly', 'body': body})

    def test_task_loop_with_a_question(self, dispatch):
        self.assertEqual(self.post(f'/api/drafting/drafts/{self.session.id}/send-to-ams/', self.junior).status_code, 200)
        resp = self.comment(self.senior)
        self.assertEqual(resp.status_code, 201)
        thread = resp.json()['threads'][0]
        self.assertEqual((thread['quote'], thread['author_name'], thread['resolved']), ('paid monthly', self.senior.full_name, False))
        from core.models import NotificationQueue
        self.assertTrue(NotificationQueue.objects.filter(type='DRAFT_COMMENT', advocate_id=self.junior.id).exists())
        # The junior answers; she can't close the senior's question herself.
        body = self.post(f'/api/drafting/comments/{thread["id"]}/reply/', self.junior,
                         {'body': 'Client agreed arrears, see meeting note.'}).json()
        self.assertEqual((body['open'], body['answered']), (1, 1))
        self.assertEqual(self.post(f'/api/drafting/comments/{thread["id"]}/resolve/', self.junior).status_code, 403)
        # Approve waits for the question to be resolved.
        review = lambda: self.post(f'/api/workspace/tasks/{self.task.id}/review', self.senior, {'action': 'approve'})  # noqa: E731
        resp = review()
        self.assertEqual(resp.status_code, 400)
        self.assertIn('comment', resp.json()['error'])
        self.assertEqual(self.post(f'/api/drafting/comments/{thread["id"]}/resolve/', self.senior).json()['open'], 0)
        self.assertEqual(review().status_code, 200)

    def test_who_resolves(self, dispatch):
        from .models import DraftReviewRequest
        # A peer's question on the junior's draft: the junior (author) may resolve it.
        DraftReviewRequest.objects.create(session=self.session, requested_by_id=self.junior.id,
                                          reviewer_id=self.peer.id, authority='suggest')
        tid = self.comment(self.peer, 'Typo in clause 2?').json()['threads'][0]['id']
        rows = self.client.get(self.s + 'comments/', **auth(self.junior)).json()['threads']
        self.assertTrue(rows[0]['can_resolve'])
        self.assertEqual(self.post(f'/api/drafting/comments/{tid}/resolve/', self.junior).status_code, 200)
        self.assertEqual(self.post(f'/api/drafting/comments/{tid}/reopen/', self.peer).json()['open'], 1)

    def test_delete_and_outsider(self, dispatch):
        tid = self.comment(self.junior, 'Note to self').json()['threads'][0]['id']
        self.assertEqual(self.client.get(self.s + 'comments/', **auth(self.outsider)).status_code, 404)
        self.assertEqual(self.client.delete(f'/api/drafting/comments/{tid}/', **auth(self.outsider)).status_code, 404)
        self.assertEqual(self.client.delete(f'/api/drafting/comments/{tid}/', **auth(self.junior)).json()['threads'], [])
        tid = self.comment(self.junior, 'Check dates').json()['threads'][0]['id']
        self.post(f'/api/drafting/comments/{tid}/reply/', self.junior, {'body': 'Done'})
        self.assertEqual(self.client.delete(f'/api/drafting/comments/{tid}/', **auth(self.junior)).status_code, 400)

    def test_empty_and_foreign_passage_refused(self, dispatch):
        self.assertEqual(self.comment(self.junior, '   ').status_code, 400)
        resp = self.post(self.s + 'comments/', self.junior, {'block_id': 999999, 'body': 'x'})
        self.assertEqual(resp.status_code, 400)


class CaseFileTest(TestCase):
    """drafting/casefile.py: the linked case's summary and documents, never beyond the viewer's rights."""

    def setUp(self):
        from django.utils import timezone
        from core.models import Case, Client as AmsClient, Document
        self.owner = make_advocate(permissions=DRAFTER + ('CASE_VIEW', 'DOCUMENT_VIEW'))
        self.no_docs = make_advocate(permissions=DRAFTER + ('CASE_VIEW',), parent_advocate_id=self.owner.id)
        self.outsider = make_advocate(permissions=ALL_PERMISSIONS)
        client = AmsClient.objects.create(name='K. Kannan', address='12 Anna Salai', advocate=self.owner)
        self.case = Case.objects.create(case_number='RC 12/2026', case_title='Kannan v Arun', advocate=self.owner,
                                        client=client)
        now = timezone.now()
        for name, case in (('Lease deed.pdf', self.case), ('Aadhaar.pdf', None)):
            Document.objects.create(document_name=name, original_name=name, stored_name=name, file_path=name,
                                    upload_date=now, advocate=self.owner, case=case, client=client)
        project = Project.objects.create(name='p', client=Client.objects.create(name='c'), case_id=self.case.id)
        self.linked = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready', project=project)
        self.unlinked = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')

    def get(self, session, who):
        return self.client.get(f'/api/drafting/draft-sessions/{session.id}/case-file/', **auth(who))

    def test_linked_case(self):
        data = self.get(self.linked, self.owner).json()
        self.assertEqual((data['case']['caseNumber'], data['case']['client']['address']), ('RC 12/2026', '12 Anna Salai'))
        self.assertEqual(sorted(d['name'] for d in data['documents']), ['Aadhaar.pdf', 'Lease deed.pdf'])

    def test_no_case_and_rights(self):
        self.assertIsNone(self.get(self.unlinked, self.owner).json()['case'])
        self.assertEqual(self.get(self.linked, self.outsider).status_code, 404)     # not their draft
        self.linked.created_by_id = self.no_docs.id
        self.linked.save()
        data = self.get(self.linked, self.no_docs).json()
        self.assertEqual((data['canSeeDocuments'], data['documents']), (False, []))


class SaveConflictAndRestoreTest(TestCase):
    """Two people saving the same draft; putting a draft back to a saved version."""

    def setUp(self):
        from .models import DraftBlock
        self.owner = make_advocate(permissions=DRAFTER)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.a = DraftBlock.objects.create(session=self.session, position=0, block_type='clause', heading='Rent',
                                           text='', content_html='<p>Rent 40.</p>', source='generated')
        self.b = DraftBlock.objects.create(session=self.session, position=1, block_type='clause', heading='Term',
                                           text='', content_html='<p>11 months.</p>', source='generated')
        self.url = f'/api/drafting/draft-sessions/{self.session.id}/'

    def loaded(self):
        data = self.client.get(self.url, **auth(self.owner)).json()
        return {str(b['id']): b['rev'] for b in data['blocks']}

    def save(self, base, rent, term):
        blocks = [{'id': self.a.id, 'heading': 'Rent', 'content_html': rent, 'text': ''},
                  {'id': self.b.id, 'heading': 'Term', 'content_html': term, 'text': ''}]
        return self.client.post(self.url + 'save-blocks/', {'blocks': blocks, 'base': base},
                                content_type='application/json', **auth(self.owner))

    def test_different_clauses_merge_same_clause_refused(self):
        first, second = self.loaded(), self.loaded()
        self.assertEqual(self.save(first, '<p>Rent 45.</p>', '<p>11 months.</p>').status_code, 200)
        # The second person changed only the term: both changes stay, and they're told to reload.
        resp = self.save(second, '<p>Rent 40.</p>', '<p>12 months.</p>')
        self.assertEqual((resp.status_code, resp.json()['merged']), (200, True))
        self.a.refresh_from_db(); self.b.refresh_from_db()
        self.assertEqual((self.a.content_html, self.b.content_html), ('<p>Rent 45.</p>', '<p>12 months.</p>'))
        # A third, stale copy changing the rent again is refused and nothing is saved.
        resp = self.save(second, '<p>Rent 50.</p>', '<p>12 months.</p>')
        self.assertEqual((resp.status_code, resp.json()['conflicts']), (409, ['Rent']))
        self.a.refresh_from_db()
        self.assertEqual(self.a.content_html, '<p>Rent 45.</p>')

    def test_restore_version(self):
        from .versions import save_version
        v1 = save_version(self.session, label='first', user_id=self.owner.id)
        self.save(self.loaded(), '<p>Rent 99.</p>', '<p>11 months.</p>')
        self.b.delete()
        resp = self.client.post(self.url + f'versions/{v1.id}/restore/', **auth(self.owner))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual([b['content_html'] for b in resp.json()['blocks']], ['<p>Rent 40.</p>', '<p>11 months.</p>'])
        self.assertEqual(resp.json()['blocks'][0]['id'], self.a.id)                 # same clause, anchors kept
        labels = list(self.session.versions.order_by('number').values_list('label', flat=True))
        self.assertEqual(labels, ['first', 'Before restoring v1', 'Restored v1'])
        other = make_advocate(permissions=DRAFTER)
        self.assertEqual(self.client.post(self.url + f'versions/{v1.id}/restore/', **auth(other)).status_code, 404)


class ImportChangesTest(TestCase):
    """drafting/incoming.py: a Word file from outside becomes a suggestions round + comment threads."""

    def setUp(self):
        from .models import DraftBlock
        self.owner = make_advocate(permissions=DRAFTER)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={'document_title': 'Notice'},
                                                   status='ready')
        DraftBlock.objects.create(session=self.session, position=0, block_type='clause', heading='Rent', text='',
                                  content_html='<p>Rent is Rs. <span data-placeholder="Rent" data-value="">[Rent]</span> a month.</p>'
                                               '<p>Paid monthly.</p>', source='generated')
        DraftBlock.objects.create(session=self.session, position=1, block_type='clause', heading='Term', text='',
                                  content_html='<p>Eleven months.</p>', source='generated')

    def upload(self, data, name='back.docx', **extra):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(f'/api/drafting/draft-sessions/{self.session.id}/import-changes/',
                                {'file': SimpleUploadedFile(name, data), **extra}, **auth(self.owner))

    def test_our_own_export_has_nothing(self):
        from .export.docx import render_session_docx
        resp = self.upload(render_session_docx(self.session))
        self.assertEqual(resp.status_code, 400)                       # reads the same as the draft

    def test_tracked_changes_and_comments(self):
        import io
        import zipfile
        from .export.redline import render_redline_docx
        from .models import DraftComment, DraftReviewRound
        from .versions import snapshot_blocks
        base = snapshot_blocks(self.session)
        theirs = [dict(b) for b in base]
        theirs[1]['content_html'] = '<p>Twelve months.</p>'
        out = render_redline_docx(base, theirs, 'Notice', author='Arun (counsel)')
        data = out[0] if isinstance(out, tuple) else out
        # Add a Word comment on "Paid monthly."
        zin = zipfile.ZipFile(io.BytesIO(data))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as zout:
            for item in zin.infolist():
                body = zin.read(item.filename)
                if item.filename == 'word/document.xml':
                    xml = body.decode('utf-8')
                    i = xml.index('Paid monthly.')
                    start = xml.rindex('<w:r>', 0, i) if '<w:r>' in xml[:i] else xml.rindex('<w:r ', 0, i)
                    end = xml.index('</w:r>', i) + len('</w:r>')
                    xml = (xml[:start] + '<w:commentRangeStart w:id="0"/>' + xml[start:end]
                           + '<w:commentRangeEnd w:id="0"/>' + xml[end:])
                    body = xml.encode('utf-8')
                zout.writestr(item, body)
            zout.writestr('word/comments.xml',
                          '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                          '<w:comment w:id="0" w:author="Arun (counsel)"><w:p><w:r><w:t>Paid in advance?</w:t></w:r></w:p>'
                          '</w:comment></w:comments>')
        resp = self.upload(buf.getvalue(), from_name='')
        self.assertEqual(resp.status_code, 201, resp.content)
        body = resp.json()
        self.assertEqual((body['changes'], body['comments'], body['tracked'], body['from_name']),
                         (1, 1, True, 'Arun (counsel)'))
        rnd = DraftReviewRound.objects.get(id=body['round'])
        self.assertEqual((rnd.decider_id, rnd.external_from), (self.owner.id, 'Arun (counsel)'))
        comment = DraftComment.objects.get(session=self.session)
        self.assertEqual(comment.quote, 'Paid monthly.')
        self.assertIn('Paid in advance?', comment.body)
        self.assertEqual(comment.block_id, base[0]['block_id'])

    def test_not_word(self):
        self.assertEqual(self.upload(b'%PDF-1.4', name='back.pdf').status_code, 400)
        self.assertEqual(self.upload(b'not a zip').status_code, 400)


class TemplateLayoutTest(TestCase):
    """services/layout.py: a generated draft lays out like its Word template."""

    def make_docx(self):
        import tempfile
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        doc.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY   # inherited, not on the paragraph
        doc.add_paragraph('LEGAL NOTICE').alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph('From:').alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.add_paragraph('Under instructions from and on behalf of my client I serve this notice upon you.')
        doc.add_paragraph('The tenant has failed to pay the rent for two months despite repeated requests.')
        doc.add_paragraph('Yours faithfully, Advocate').alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        doc.save(f.name)
        f.close()
        return f.name

    def test_layout_and_generated_html(self):
        import os
        from .services.layout import docx_paragraphs, layout_for
        from .tasks import _styled_html
        path = self.make_docx()
        try:
            paras = docx_paragraphs(path)
        finally:
            os.unlink(path)
        self.assertEqual([a for _, a in paras], ['center', 'left', 'justify', 'justify', 'right'])
        text = '\n'.join(t for t, _ in paras)
        lay = layout_for(text, paras)
        self.assertEqual(lay['align'], 'justify')          # the clause's usual alignment
        # New wording: lines that match the template keep their own alignment, the rest follow the body.
        html = _styled_html('LEGAL NOTICE\nFrom:\nUnder instructions from and on behalf of my client Kannan.\n'
                            'A brand new sentence the model wrote.\nYours faithfully, Advocate',
                            {'align': lay['align']}, lay['lines'])
        self.assertEqual(html.count('text-align:center'), 1)
        self.assertEqual(html.count('text-align:justify'), 2)
        self.assertEqual(html.count('text-align:right'), 1)
        self.assertIn('<p>From:</p>', html)                  # left = the editor default


class ImportKeepsLayoutTest(TestCase):
    """Import changes keeps tables, alignment and lists of the received Word file."""

    def test_table_and_alignment(self):
        from .export.docx import render_session_docx
        from .incoming import _Reader, their_blocks
        from .models import DraftBlock
        owner = make_advocate(permissions=DRAFTER)
        session = DraftSession.objects.create(created_by_id=owner.id, facts={'document_title': 'Lease'}, status='ready')
        DraftBlock.objects.create(
            session=session, position=0, block_type='clause', heading='Schedule', text='',
            content_html='<p style="text-align:justify">The rent is payable as below.</p>'
                         '<table><tr><td><p>Month</p></td><td><p>Rent</p></td></tr>'
                         '<tr><td><p>August</p></td><td><p>40,000</p></td></tr></table>'
                         '<ul><li><p>Paid by transfer</p></li></ul>', source='generated')
        ours = render_session_docx(session)
        self.assertEqual(their_blocks(session, _Reader(ours))[1][0]['content_html'],
                         session.blocks.get().content_html)          # nothing changed: untouched
        # Their copy: the rent cell and the opening sentence changed.
        block = session.blocks.get()
        block.content_html = block.content_html.replace('40,000', '45,000').replace('as below', 'monthly as below')
        block.save()
        theirs = render_session_docx(session)
        block.content_html = block.content_html.replace('45,000', '40,000').replace('monthly as below', 'as below')
        block.save()
        html = their_blocks(session, _Reader(theirs))[1][0]['content_html']
        self.assertIn('<table>', html)
        self.assertIn('<td><p>45,000</p></td>', html)
        self.assertIn('text-align: justify', html)
        self.assertIn('<ul><li><p>Paid by transfer</p></li></ul>', html)


class PdfTableTextTest(TestCase):
    """A PDF's tables reach the draft as Markdown rows; they must come out as a grid."""

    def test_markdown_rows_become_a_table(self):
        from types import SimpleNamespace
        from .tasks import _styled_html, _tables_to_html
        text = ('Fees are payable as below.\n| Tier | Fee |\n|------|-----|\n| Up to 10,000 | Rs. 5 |\n'
                '| Above 10,000 | Rs. 4 |\nPayment within 30 days.')
        html = _styled_html(text, {'align': 'justify'}, [])
        self.assertIn('<table><tbody><tr><th><p>Tier</p></th><th><p>Fee</p></th></tr>', html)
        self.assertIn('<td><p>Rs. 4</p></td>', html)
        self.assertNotIn('---', html)
        self.assertEqual(html.count('text-align:justify'), 2)          # the sentences around it
        block = SimpleNamespace(text=text, content_html='')
        _tables_to_html([block])                                        # unstyled block: still a grid
        self.assertIn('<table>', block.content_html)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
class FirmSetupOwnershipTest(TestCase):
    """Templates and playbooks are firm set-up: never visible to another firm;
    "delete" archives, allowed only to the creator or a Super Admin; restorable."""
    databases = {'default'}

    def setUp(self):
        from core.models import AdvocateRole, Role
        self.rajesh = make_advocate(permissions=ALL_PERMISSIONS)                       # creator
        self.arjun = make_advocate(permissions=ALL_PERMISSIONS, parent_advocate_id=self.rajesh.id)
        self.admin = make_advocate(permissions=ALL_PERMISSIONS, parent_advocate_id=self.rajesh.id)
        role, _ = Role.objects.get_or_create(name='Super Admin')
        AdvocateRole.objects.create(advocate_id=self.admin.id, role_id=role.id)
        self.other_firm = make_advocate(permissions=ALL_PERMISSIONS)
        self.pb = Playbook.objects.create(name='NDA Playbook', category='NDA', method='scratch',
                                          status='ready', created_by_id=self.rajesh.id)
        self.tpl = Template.objects.create(name='Mutual NDA', document_type='nda', status='ready',
                                           created_by_id=self.rajesh.id)

    def _ids(self, url, who, **params):
        body = self.client.get(url, params, **auth(who)).json()
        return [r['id'] for r in body.get('results', body)]

    def test_another_firm_cannot_see_or_delete(self):
        for url, obj in (('/api/drafting/playbooks/', self.pb), ('/api/drafting/templates/', self.tpl)):
            self.assertNotIn(obj.id, self._ids(url, self.other_firm))
            self.assertEqual(self.client.delete(f'{url}{obj.id}/', **auth(self.other_firm)).status_code, 404)

    def test_a_colleague_sees_and_edits_but_cannot_delete(self):
        self.assertIn(self.pb.id, self._ids('/api/drafting/playbooks/', self.arjun))
        resp = self.client.delete(f'/api/drafting/playbooks/{self.pb.id}/', **auth(self.arjun))
        self.assertEqual(resp.status_code, 403)
        resp = self.client.patch(f'/api/drafting/playbooks/{self.pb.id}/', {'description': 'Our NDA rules'},
                                 content_type='application/json', **auth(self.arjun))
        self.assertEqual(resp.status_code, 200, resp.content)

    def test_the_creator_archives_and_restores(self):
        for url, obj in (('/api/drafting/playbooks/', self.pb), ('/api/drafting/templates/', self.tpl)):
            self.assertEqual(self.client.delete(f'{url}{obj.id}/', **auth(self.rajesh)).status_code, 204)
            obj.refresh_from_db()
            self.assertIsNotNone(obj.archived_at)                       # kept, not deleted
            self.assertNotIn(obj.id, self._ids(url, self.rajesh))
            self.assertIn(obj.id, self._ids(url, self.rajesh, archived='1'))
            self.assertEqual(self.client.post(f'{url}{obj.id}/restore/', **auth(self.arjun)).status_code, 403)
            self.assertEqual(self.client.post(f'{url}{obj.id}/restore/', **auth(self.rajesh)).status_code, 200)
            self.assertIn(obj.id, self._ids(url, self.rajesh))

    def test_a_super_admin_may_archive_anyones(self):
        self.assertEqual(self.client.delete(f'/api/drafting/playbooks/{self.pb.id}/', **auth(self.admin)).status_code, 204)

    def test_list_says_who_may_archive(self):
        rows = {r['id']: r for r in self.client.get('/api/drafting/playbooks/', **auth(self.arjun)).json()['results']}
        self.assertFalse(rows[self.pb.id]['can_archive'])
        rows = {r['id']: r for r in self.client.get('/api/drafting/playbooks/', **auth(self.rajesh)).json()['results']}
        self.assertTrue(rows[self.pb.id]['can_archive'])


class DurationFactsBriefTest(TestCase):
    """Duration facts reach the drafting prompt as plain sentences, whether saved as
    the form's phrase or as an older bare number."""

    def test_phrases_bare_numbers_and_open_ended(self):
        from drafting.tasks import _facts_brief
        brief = _facts_brief({'term_years': '11 months', 'termination_notice_days': '30',
                              'survival_years': '1'})
        self.assertIn('Agreement duration: 11 months from the Effective Date', brief)
        self.assertIn('Notice needed to end the agreement early: 30 days', brief)
        self.assertIn('Confidentiality continues after the agreement ends for: 1 year', brief)
        self.assertIn('no fixed end; it continues until terminated by notice',
                      _facts_brief({'term_years': 'until terminated'}))
        self.assertNotIn('Term years', _facts_brief({'term_years': '2'}))


def _letterhead_template_docx():
    """A Word template with a letterhead (logo + firm name in the header), a footer, a centred title,
    bold clause headings, an indented coloured body, a table and a right-aligned signature."""
    import io
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.left_margin = Cm(21), Cm(2.2)
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    logo = io.BytesIO()
    Image.new('RGB', (120, 40), (163, 35, 43)).save(logo, format='PNG')
    logo.seek(0)
    sec.header.paragraphs[0].add_run().add_picture(logo, width=Cm(3))
    sec.header.add_paragraph('KUMAR & ASSOCIATES, Advocates')
    sec.footer.paragraphs[0].text = 'No. 12, Anna Salai, Chennai 600002'
    t = doc.add_paragraph('LEASE AGREEMENT')
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.runs[0].bold, t.runs[0].font.size = True, Pt(16)
    for head, body in (('1. RENT', 'The Lessee shall pay a monthly rent of Rs. 40,000 on or before the 5th day.'),
                       ('2. TERM', 'This lease is for a term of eleven months from 1 April 2026.')):
        doc.add_paragraph(head).runs[0].bold = True
        b = doc.add_paragraph(body)
        b.paragraph_format.left_indent = Cm(1)
        b.runs[0].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.add_paragraph('3. SCHEDULE OF PAYMENTS').runs[0].bold = True
    tbl = doc.add_table(rows=2, cols=2)
    tbl.style = 'Table Grid'
    for i, (a, b) in enumerate((('Month', 'Rent'), ('April 2026', 'Rs. 40,000'))):
        tbl.cell(i, 0).text, tbl.cell(i, 1).text = a, b
    doc.add_paragraph('Rent includes maintenance.')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
class WordFormTest(TestCase):
    """drafting/wordform.py: a draft from a Word template looks like it in the editor, and Download
    writes into the original file (letterhead, header, footer, page, formatting kept)."""

    def setUp(self):
        from .models import DraftBlock
        from .services.parser import extract_template_slots
        from .tasks import _apply_word_form
        self.owner = make_advocate(permissions=DRAFTER)
        self.template = Template.objects.create(
            name='Lease', file=SimpleUploadedFile('lease.docx', _letterhead_template_docx()))
        self.slots = extract_template_slots(self.template.file.path)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready',
                                                   template=self.template)
        for n, slot in enumerate(self.slots):
            text = slot['description'].replace('Rs. 40,000 on', 'Rs. 45,000 on')
            if 'SCHEDULE' in text:
                text = ('3. SCHEDULE OF PAYMENTS\n| Month | Rent |\n|---|---|\n| April 2026 | Rs. 45,000 |\n'
                        '| May 2026 | Rs. 45,000 |\nRent includes maintenance.')
            block = DraftBlock(session=self.session, position=n, block_type='clause', heading='', text=text,
                               source='generated')
            _apply_word_form(block, slot.get('form'), f'template:{self.template.id}')
            block.save()

    def test_slots_know_their_place_and_look(self):
        self.assertEqual([s['label'] for s in self.slots], ['LEASE AGREEMENT', 'RENT', 'TERM', 'SCHEDULE OF PAYMENTS'])
        self.assertTrue(all(s.get('form') for s in self.slots))
        # Word tables reach generation as table rows (| a | b |), like a PDF's.
        self.assertIn('| Month | Rent |', self.slots[3]['description'])

    def test_editor_opens_in_the_template_look(self):
        rent = self.session.blocks.get(position=1)
        self.assertIn("font-family:'Times New Roman'", rent.content_html)
        self.assertIn('margin-left:28.4pt', rent.content_html)            # the 1 cm indent
        self.assertIn('color:#1F3A5F', rent.content_html)
        self.assertIn('<table>', self.session.blocks.get(position=3).content_html)
        self.assertIn('text-align:center', self.session.blocks.get(position=0).content_html)   # the title
        data = self.client.get(f'/api/drafting/draft-sessions/{self.session.id}/', **auth(self.owner)).json()
        form = data['word_form']
        self.assertEqual((form['source'], form['has_letterhead'], form['page']['width_mm']), ('template', True, 210.0))
        self.assertIn('<img src="data:image/png;base64,', form['header_html'])
        self.assertIn('Anna Salai', form['footer_html'])

    def test_download_writes_into_the_original(self):
        import io
        from docx import Document
        resp = self.client.get(f'/api/drafting/drafts/{self.session.id}/export/docx/?branding=1', **auth(self.owner))
        self.assertEqual(resp.status_code, 200)
        doc = Document(io.BytesIO(b''.join(resp.streaming_content) if hasattr(resp, 'streaming_content')
                                  else resp.content))
        body = [p for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(body[0].text, 'LEASE AGREEMENT')                  # outside the clauses: untouched
        self.assertTrue(body[0].runs[0].bold)
        rent = next(p for p in body if 'monthly rent' in p.text)
        self.assertIn('Rs. 45,000', rent.text)
        self.assertAlmostEqual(rent.paragraph_format.left_indent.cm, 1.0, places=1)
        self.assertEqual(str(rent.runs[0].font.color.rgb), '1F3A5F')
        rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        self.assertEqual(rows, [['Month', 'Rent'], ['April 2026', 'Rs. 45,000'], ['May 2026', 'Rs. 45,000']])
        self.assertEqual(doc.tables[0].style.name, 'Table Grid')
        sec = doc.sections[0]
        self.assertEqual(len(sec.header._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')), 1)
        self.assertEqual([p.text for p in sec.header.paragraphs if p.text], ['KUMAR & ASSOCIATES, Advocates'])
        self.assertEqual(sec.footer.paragraphs[0].text, 'No. 12, Anna Salai, Chennai 600002')
        # Its own letterhead: the firm letterhead asked for (?branding=1) is not added on top.
        self.assertEqual(len(sec.header.paragraphs), 2)

    def test_draft_without_word_source_keeps_house_style(self):
        from .wordform import render_draft_docx
        plain = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        from .models import DraftBlock
        DraftBlock.objects.create(session=plain, position=0, block_type='clause', heading='Rent', text='',
                                  content_html='<p>Rent is Rs. 1.</p>', source='generated')
        self.assertTrue(render_draft_docx(plain).startswith(b'PK'))
        self.assertIsNone(self.client.get(f'/api/drafting/draft-sessions/{plain.id}/', **auth(self.owner)).json()['word_form'])


class PdfToWordTest(TestCase):
    """services/pdf2word.py: a PDF becomes a Word file with its layout (needs LibreOffice to make the PDF)."""

    def test_round_trip(self):
        import io
        import os
        import tempfile as tf
        from unittest import skip  # noqa: F401
        from docx import Document
        from .export.pdf import PdfUnavailable, docx_to_pdf
        from .services.pdf2word import pdf_to_docx
        try:
            pdf = docx_to_pdf(_letterhead_template_docx())
        except PdfUnavailable:
            self.skipTest('LibreOffice is not installed')
        path = os.path.join(tf.mkdtemp(), 'lease.pdf')
        with open(path, 'wb') as f:
            f.write(pdf)
        doc = Document(io.BytesIO(pdf_to_docx(path)))
        body = [p for p in doc.paragraphs if p.text.strip()]
        title = next(p for p in body if p.text == 'LEASE AGREEMENT')
        self.assertEqual(int(title.alignment), 1)                          # centred
        self.assertTrue(any(p.style.name.startswith('Heading') and 'RENT' in p.text for p in body))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.sections[0].header._element.findall(
            './/{http://schemas.openxmlformats.org/drawingml/2006/main}blip')), 1)   # the logo
        self.assertIn('KUMAR & ASSOCIATES', ' '.join(p.text for p in doc.sections[0].header.paragraphs))

    def test_scan_gives_none(self):
        import os
        import tempfile as tf
        from PIL import Image
        from .services.pdf2word import pdf_to_docx
        path = os.path.join(tf.mkdtemp(), 'scan.pdf')
        Image.new('RGB', (600, 800), (255, 255, 255)).save(path, 'PDF')
        self.assertIsNone(pdf_to_docx(path))


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
class TemplatePrintViewTest(TestCase):
    """View on Firm Templates / Samples: the file as it prints (drafting/files.py preview)."""

    def setUp(self):
        self.user = make_advocate(permissions=ALL_PERMISSIONS)

    def get(self, url):
        resp = self.client.get(url, **auth(self.user))
        body = b''.join(resp.streaming_content) if hasattr(resp, 'streaming_content') else resp.content
        return resp, body

    def test_word_template_is_shown_as_a_pdf(self):
        t = Template.objects.create(name='Lease', created_by_id=self.user.id,
                                    file=SimpleUploadedFile('lease.docx', _letterhead_template_docx()))
        data = self.client.get(f'/api/drafting/templates/{t.id}/', **auth(self.user)).json()
        self.assertIn(f'/api/drafting/templates/{t.id}/preview/', data['preview'])
        resp, body = self.get(data['preview'])
        if resp.status_code == 503:                                   # no LibreOffice on this machine
            self.assertIn('LibreOffice', resp.json()['error'])
            return
        self.assertEqual((resp.status_code, resp['Content-Type']), (200, 'application/pdf'))
        self.assertTrue(body.startswith(b'%PDF'))
        again, body2 = self.get(data['preview'])                      # kept, not made again
        self.assertEqual(body, body2)

    def test_pdf_template_is_sent_as_it_is(self):
        pdf = b'%PDF-1.4\n%fake\n'
        t = Template.objects.create(name='NDA', created_by_id=self.user.id, file=SimpleUploadedFile('nda.pdf', pdf))
        resp, body = self.get(f'/api/drafting/templates/{t.id}/preview/')
        self.assertEqual((resp.status_code, body), (200, pdf))


class WordFormHeadingTest(TestCase):
    """Headings are kept apart from bodies (services/headings.py): the template's own heading, never the
    AI's slot name; fixed headings stay, variable ones are filled."""

    def test_slot_parts_use_the_templates_own_heading(self):
        from .tasks import _slot_parts
        slot = {'label': 'Lease Agreement Introduction', 'description': 'LEASE AGREEMENT\nThis Lease is made.',
                'heading': 'LEASE AGREEMENT', 'body': 'This Lease is made.', 'heading_kind': 'fixed'}
        self.assertEqual(_slot_parts(slot), ('LEASE AGREEMENT', 'This Lease is made.', 'fixed'))
        old = {'label': 'Premises', 'description': '1. PREMISES\nThe Lessor lets the flat.'}     # before headings
        self.assertEqual(_slot_parts(old), ('1. PREMISES', 'The Lessor lets the flat.', 'fixed'))

    def test_heading_kinds(self):
        from .services.headings import heading_kind
        self.assertEqual(heading_kind('VERSUS'), 'fixed')
        self.assertEqual(heading_kind('4. SCHEDULE OF PAYMENTS'), 'fixed')        # the clause number is not a detail
        self.assertEqual(heading_kind('IN THE COURT OF [[Court Name]], CHENNAI'), 'variable')
        self.assertEqual(heading_kind('IN THE COURT OF THE PRINCIPAL DISTRICT COURT, CHENNAI'), 'variable')
        self.assertEqual(heading_kind('O.S. No. 145 of 2024'), 'variable')
        self.assertEqual(heading_kind(''), '')

    def test_variable_heading_filled_fixed_left(self):
        from types import SimpleNamespace
        from .tasks import _fill_variable_headings
        llm = mock.Mock()
        llm.complete.return_value = 'IN THE COURT OF THE CITY CIVIL COURT, CHENNAI'
        blocks = [SimpleNamespace(position=0, heading='IN THE COURT OF [[Court Name]], CHENNAI'),
                  SimpleNamespace(position=1, heading='VERSUS')]
        specs = [{'position': 0, 'heading': blocks[0].heading, 'heading_kind': 'variable'},
                 {'position': 1, 'heading': 'VERSUS', 'heading_kind': 'fixed'}]
        _fill_variable_headings(llm, blocks, specs, 'Court: City Civil Court, Chennai')
        self.assertEqual(blocks[0].heading, 'IN THE COURT OF THE CITY CIVIL COURT, CHENNAI')
        self.assertEqual(blocks[1].heading, 'VERSUS')
        self.assertEqual(llm.complete.call_count, 1)                                 # only the variable one


class WordLinesTest(TestCase):
    """Word-form lines: unchanged lines keep their runs, layout lines and drawn shapes stay in the
    download, nothing is put back by rule."""

    def _docx(self):
        import base64
        import tempfile
        from docx import Document
        doc = Document()
        doc.add_paragraph('VAKALATNAMA')
        p = doc.add_paragraph('I, ')
        p.add_run('K. Kannan').bold = True
        p.add_run(' being the Plaintiff')
        doc.add_paragraph('))\t))')
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==')
        img = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
        img.write(png)
        img.close()
        boxed = doc.add_paragraph('Welfare stamp')
        boxed.add_run().add_picture(img.name)
        doc.add_paragraph('Dated: ____')
        f = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        f.close()
        doc.save(f.name)
        return f.name

    def _session(self, text, form):
        from types import SimpleNamespace

        from .tasks import _apply_word_form
        block = SimpleNamespace(text=text, heading='', content_html='', style_json={}, position=0)
        _apply_word_form(block, form, 'template:1')
        blocks = SimpleNamespace(all=lambda: [block], order_by=lambda *_: [block])
        return block, SimpleNamespace(blocks=blocks)

    def test_layout_lines_dropped_and_nothing_put_back(self):
        from .wordform import drop_layout_lines
        self.assertEqual(drop_layout_lines('A\n))\n-----\n))\tPLAINTIFF'), 'A\n))\tPLAINTIFF')
        path = self._docx()
        from .wordform import clause_forms
        form = clause_forms(path, ['VAKALATNAMA\nI, ____ being'])[0]
        self.assertEqual(form['shapes'], 1)
        block, _ = self._session('I, K. Kannan being the Plaintiff\nDated: 9 October 2026', form)
        self.assertNotIn('VAKALATNAMA', block.text)            # a template line the model left out is not forced back
        self.assertEqual(block.style_json['src']['shapes'], 1)

    def test_unchanged_line_keeps_runs_and_shapes_survive(self):
        import io

        from docx import Document
        from docx.oxml.ns import qn

        from .wordform import clause_forms, render_into_original
        path = self._docx()
        form = clause_forms(path, ['VAKALATNAMA\nI, ____ being'])[0]
        block, session = self._session(
            'VAKALATNAMA\nI, K. Kannan being the Plaintiff\n))\nWelfare stamp affixed as per law\nDated: 9 October 2026', form)
        self.assertIn('<strong>K. Kannan</strong>', block.content_html)   # editor: bold words of an unchanged line
        self.assertNotIn('<p>))</p>', block.content_html)
        out = Document(io.BytesIO(render_into_original(session, path)))
        texts = [p.text for p in out.paragraphs]
        self.assertEqual(texts.count('))\t))'), 1)                         # layout line from the file, once
        self.assertNotIn('))', [t.strip() for t in texts])
        kannan = next(p for p in out.paragraphs if 'Kannan' in p.text)
        self.assertTrue(any(r.bold and r.text == 'K. Kannan' for r in kannan.runs))
        self.assertEqual(len(out.element.body.findall('.//' + qn('w:drawing'))), 1)   # the box survives
        self.assertIn('Dated: 9 October 2026', texts)


class TemplateLookFixesTest(TestCase):
    """Fixes from testing a PDF sample: no repeated heading line, colours kept, signature lines."""

    def test_heading_shown_apart_is_not_repeated_in_the_text(self):
        from types import SimpleNamespace
        from .tasks import _apply_word_form
        form = {'range': [3, 4], 'heading': '1. PREMISES', 'paras': [
            {'i': 3, 'key': '1 premises', 'fmt': {'bold': True, 'color': '#7A1F2B'}, 'text': '1. PREMISES'},
            {'i': 4, 'key': 'the lessor lets to the lessee the', 'fmt': {'align': 'justify'},
             'text': 'The Lessor lets to the Lessee the premises at [[Premises Address]].'}]}
        block = SimpleNamespace(text='1. PREMISES\nThe Lessor lets to the Lessee the flat at T. Nagar.',
                                heading='1. PREMISES', content_html='', style_json={})
        _apply_word_form(block, form, 'sample:1')
        self.assertEqual(block.text, 'The Lessor lets to the Lessee the flat at T. Nagar.')
        self.assertNotIn('PREMISES', block.content_html)
        self.assertEqual(block.style_json['heading']['color'], '#7A1F2B')

    def test_pdf_colours(self):
        from .services.pdf2word import _hex
        self.assertEqual(_hex((0.478, 0.122, 0.169)), '7A1F2B')     # RGB
        self.assertIsNone(_hex((0,)))                                 # black
        self.assertIsNone(_hex((0, 0, 0, 1)))                         # CMYK black
        self.assertEqual(_hex((0, 0.75, 0.6, 0.52)), '7A1F31')        # CMYK maroon
        self.assertIsNone(_hex('pattern'))

    def test_signature_lines_are_not_headings(self):
        from .services.pdf2word import _SIGNATURE_LINE
        for line in ('LESSOR', 'LESSEE', 'For Kumar & Associates', 'Witness 1', 'Authorised Signatory'):
            self.assertTrue(_SIGNATURE_LINE.match(line), line)
        self.assertFalse(_SIGNATURE_LINE.match('1. PREMISES'))


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='drafting-test-'))
class ReferenceDocumentLinksTest(TestCase):
    """The draft page's Reference documents open: links go through the login, never /media."""

    def test_links_load_for_the_owner(self):
        owner = make_advocate(permissions=ALL_PERMISSIONS)
        sample = Sample.objects.create(name='Lease', uploaded_by_id=owner.id,
                                       file=SimpleUploadedFile('lease.pdf', b'%PDF-1.4\n%x\n'))
        session = DraftSession.objects.create(created_by_id=owner.id, facts={}, status='ready')
        session.samples.add(sample)
        doc = self.client.get(f'/api/drafting/draft-sessions/{session.id}/', **auth(owner)).json()['reference_documents'][0]
        self.assertNotIn('/media/', doc['url'])
        self.assertIn(f'/api/drafting/samples/{sample.id}/file/', doc['url'])
        self.assertIn(f'/api/drafting/samples/{sample.id}/preview/', doc['preview'])
        for url in (doc['url'], doc['preview']):
            resp = self.client.get(url, **auth(owner))
            body = b''.join(resp.streaming_content) if hasattr(resp, 'streaming_content') else resp.content
            self.assertEqual((resp.status_code, body[:5]), (200, b'%PDF-'))
            self.assertEqual(self.client.get(url).status_code, 401)        # never without the login


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class OnlyOfficeTest(TestCase):
    """ONLYOFFICE as the editor (drafting/onlyoffice.py): signed config, signed file link, save callback."""

    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.media_override = override_settings(MEDIA_ROOT=self.media)
        self.media_override.enable()
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.session.blocks.create(position=0, block_type='clause', heading='', text='First clause.', source='generated')

    def tearDown(self):
        self.media_override.disable()

    def _config(self, who=None):
        return self.client.get(f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/', **auth(who or self.owner))

    def test_desktop_editing_over_webdav(self):
        """Word opens the draft over WebDAV (documents/dav.py): its Save is a new version, and the PactPro
        editor is read-only while Word holds the lock."""
        links = self.client.get(f'/api/desktop-links?kind=draft&id={self.session.id}', **auth(self.owner)).json()
        self.assertTrue(links['editable'])
        self.assertTrue(links['office'].startswith('ms-word:ofe|u|'))
        path = '/dav/' + links['url'].split('/dav/', 1)[1]
        data = self.client.get(path).content
        self.assertEqual(data[:2], b'PK')
        self.assertEqual(self.client.generic('LOCK', path).status_code, 200)
        self.assertEqual(self._config().json()['role'], 'view')
        self.assertEqual(self.client.generic('PUT', path, data).status_code, 204)
        self.assertEqual(self.client.generic('PUT', path, data).status_code, 204)
        self.assertEqual(DraftFile.objects.get(session=self.session).version, 1)   # working copy only
        self.assertEqual(self.client.generic('UNLOCK', path).status_code, 204)
        self.assertEqual(self._config().json()['role'], 'decide')                 # own edits: go on with them
        resp = self.client.post(f'/api/desktop-work/draft/{self.session.id}', {'action': 'save'},
                                content_type='application/json', **auth(self.owner))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(DraftFile.objects.get(session=self.session).version, 2)
        self.assertEqual(self._config().json()['role'], 'decide')

    def test_redline_between_word_versions(self):
        """Download → Redline on an ONLYOFFICE draft: tracked changes between two Word versions, or to now."""
        import io as _io

        from docx import Document as _Doc
        from .docfile import apply_tracked
        from .onlyoffice import store_version
        self._config()
        df = DraftFile.objects.get(session=self.session)
        with df.file.open('rb') as f:
            data, ok = apply_tracked(f.read(), 'First clause.', 'First clause, as amended.')
        self.assertTrue(ok)
        store_version(df, data, self.owner.id)                               # v2: the clause reads amended
        url = f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/redline/'
        # Without LibreOffice: the standard-layout redline from each version's clause text.
        from .export.libre_redline import RedlineUnavailable
        with mock.patch('drafting.export.libre_redline.compare_docx', side_effect=RedlineUnavailable('none')):
            resp = self.client.get(url, {'from': 1, 'to': 2}, **auth(self.owner))
        self.assertEqual(resp.status_code, 200, resp.content[:300])
        self.assertEqual(resp['X-Redline-Layout'], 'standard')
        self.assertGreater(int(resp['X-Redline-Inserted']), 0)
        self.assertIn('_redline_v1-v2_', resp['Content-Disposition'])
        body = _Doc(_io.BytesIO(resp.content)).element.body.xml
        self.assertIn('w:ins', body)
        self.assertIn('amended', body)
        # With LibreOffice: its Compare Document on the two Word files, in the draft's own look.
        with mock.patch('drafting.export.libre_redline.compare_docx',
                        return_value=(b'PK compared', {'inserted': 3, 'deleted': 1})) as compare:
            now = self.client.get(url, **auth(self.owner))                  # default: the version before, to now
        self.assertEqual(now.status_code, 200)
        self.assertEqual((now.content, now['X-Redline-Layout'], now['X-Redline-Inserted']), (b'PK compared', 'template', '3'))
        self.assertIn('_redline_v1-current_', now['Content-Disposition'])
        older, newer, _author = compare.call_args[0]
        self.assertEqual(older[:2], b'PK')                                    # the two Word files themselves
        self.assertNotEqual(older, newer)
        self.assertEqual(self.client.get(url, {'from': 99}, **auth(self.owner)).status_code, 404)
        stranger = make_advocate(permissions=ALL_PERMISSIONS)
        self.assertEqual(self.client.get(url, **auth(stranger)).status_code, 404)

    def _callback(self, body, token=None, link=None):
        import json

        import jwt
        from django.conf import settings
        link = link or self._config().json()['config']['editorConfig']['callbackUrl'].split('?t=')[1]
        signed = jwt.encode({'payload': body}, token or settings.ONLYOFFICE_JWT_SECRET, algorithm='HS256')
        return self.client.post(f'/api/drafting/onlyoffice/callback/{self.session.id}/?t={link}',
                                data=json.dumps(body), content_type='application/json',
                                HTTP_AUTHORIZATION=f'Bearer {signed}')

    @override_settings(DRAFT_EDITOR='tiptap')
    def test_off_unless_switched_on(self):
        self.assertEqual(self._config().json(), {'enabled': False})

    def test_config_is_signed_and_the_file_link_works(self):
        import jwt
        from django.conf import settings
        data = self._config().json()
        self.assertTrue(data['enabled'])
        cfg = data['config']
        signed = jwt.decode(cfg['token'], settings.ONLYOFFICE_JWT_SECRET, algorithms=['HS256'])
        self.assertEqual(signed['document']['key'], cfg['document']['key'])
        self.assertEqual(cfg['editorConfig']['mode'], 'edit')
        path = cfg['document']['url'].replace('http://ams.test', '')
        resp = self.client.get(path)                                  # the document server: no AMS login
        self.assertEqual((resp.status_code, resp.content[:2]), (200, b'PK'))
        bad = path.split('?t=')[0] + '?t=forged'
        self.assertEqual(self.client.get(bad).status_code, 404)

    def test_other_users_get_404(self):
        stranger = make_advocate(permissions=ALL_PERMISSIONS)
        self.assertEqual(self._config(stranger).status_code, 404)

    def test_pactpro_save_callback_saves_a_new_version_and_changes_the_key(self):
        from .models import DraftFile
        self._config()
        df = DraftFile.objects.get(session=self.session)
        old_key = df.key
        with df.file.open('rb') as f:
            saved = f.read()                                          # a real Word file, as ONLYOFFICE sends
        with mock.patch('drafting.onlyoffice.requests.get') as get:
            get.return_value = mock.Mock(content=saved, raise_for_status=lambda: None)
            resp = self._callback({'key': old_key, 'status': 6, 'url': 'http://oo.test/cache/x.docx',
                                   'users': [str(self.owner.id)], 'userdata': 'pactpro'})
        self.assertEqual(resp.json(), {'error': 0})
        df.refresh_from_db()
        self.assertEqual(df.version, 2)
        self.assertNotEqual(df.key, old_key)
        self.assertEqual(df.versions.count(), 2)
        with df.file.open('rb') as f:
            self.assertEqual(f.read(), saved)

    def test_closing_or_the_editors_own_save_is_not_a_version(self):
        """Only PactPro's Save makes a version; closing the editor or its own Save keeps a working copy."""
        from documents.dav import work_of
        from .models import DraftFile
        self._config()
        df = DraftFile.objects.get(session=self.session)
        with df.file.open('rb') as f:
            saved = f.read()
        with mock.patch('drafting.onlyoffice.requests.get') as get:
            get.return_value = mock.Mock(content=saved + b'', raise_for_status=lambda: None)
            self._callback({'key': df.key, 'status': 6, 'url': 'http://oo.test/x', 'users': [str(self.owner.id)],
                            'forcesavetype': 1})
            self._callback({'key': df.key, 'status': 2, 'url': 'http://oo.test/x', 'users': [str(self.owner.id)]})
        df.refresh_from_db()
        self.assertEqual((df.version, df.editing), (1, False))
        self.assertIsNotNone(work_of('draft', self.session.id, self.owner.id))
        # Reopening: the owner goes on with their working copy, editable.
        cfg = self._config().json()
        self.assertEqual(cfg['role'], 'decide')
        self.assertTrue(cfg['pendingWork'])
        self.assertIn('&work=1', cfg['config']['document']['url'])
        # PactPro's Save (editor not open): the working copy becomes one version.
        resp = self.client.post(f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/save/', **auth(self.owner))
        self.assertEqual(resp.json()['version'], 2)
        self.assertIsNone(work_of('draft', self.session.id))

    def test_callback_rejects_a_wrong_signature(self):
        from .models import DraftFile
        self._config()
        df = DraftFile.objects.get(session=self.session)
        resp = self._callback({'key': df.key, 'status': 2, 'url': 'http://x'}, token='wrong-secret-wrong-secret-wrong!!')
        self.assertEqual(resp.status_code, 403)
        df.refresh_from_db()
        self.assertEqual(df.version, 1)

    def test_editing_status_and_download_uses_the_file(self):
        from .models import DraftFile
        self._config()
        df = DraftFile.objects.get(session=self.session)
        self._callback({'key': df.key, 'status': 1, 'users': [str(self.owner.id)]})
        df.refresh_from_db()
        self.assertTrue(df.editing)
        with mock.patch('drafting.onlyoffice._command', return_value={'error': 4}):
            resp = self.client.get(f'/api/drafting/drafts/{self.session.id}/export/docx/?watermark=0', **auth(self.owner))   # the stored file, no DRAFT mark
        with df.file.open('rb') as f:
            self.assertEqual(resp.content, f.read())


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class OnlyOfficePhase2Test(TestCase):
    """Clause markers in the file, text copied back on save, placeholders filled in the file."""

    def setUp(self):
        self.media_override = override_settings(MEDIA_ROOT=tempfile.mkdtemp())
        self.media_override.enable()
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.b1 = self.session.blocks.create(position=0, block_type='clause', heading='PARTIES', source='generated',
                                             text='This deed is made by [[Client Name]] of Chennai.')
        self.b2 = self.session.blocks.create(position=1, block_type='clause', heading='', source='generated',
                                             text='Rent is payable monthly.')

    def tearDown(self):
        self.media_override.disable()

    def _file(self):
        from .onlyoffice import ensure_file
        df = ensure_file(self.session, self.owner)
        with df.file.open('rb') as f:
            return df, f.read()

    def test_every_clause_has_a_marker_and_blanks_become_form_fields(self):
        from .docfile import clause_texts, fields, placeholders
        _, data = self._file()
        texts = clause_texts(data)
        self.assertEqual(set(texts), {self.b1.id, self.b2.id})
        # An empty field reads as its [[Label]] blank (for Lisa, search, compare)...
        self.assertIn('This deed is made by [[Client Name]] of Chennai.', texts[self.b1.id])
        # ...and is a form field in the file, not [[...]] text.
        self.assertEqual(fields(data), {'Client Name': {'count': 1, 'value': ''}})
        self.assertEqual(placeholders(data), [])

    def test_editor_placeholder_span_is_written_as_a_token(self):
        from .docfile import placeholders
        self.b1.content_html = '<p>Made by <span data-placeholder="Client Name" data-value="">[Client Name]</span> and ' \
                               '<span data-placeholder="Witness" data-value="R. Ravi">R. Ravi</span>.</p>'
        self.b1.save()
        _, data = self._file()
        from .docfile import clause_texts, fields
        self.assertEqual(fields(data), {'Client Name': {'count': 1, 'value': ''}})
        self.assertEqual(placeholders(data), [])
        self.assertIn('Made by [[Client Name]] and R. Ravi.', clause_texts(data)[self.b1.id])

    def test_saved_file_text_goes_back_into_the_clauses(self):
        from .docfile import fill_fields
        from .onlyoffice import store_version
        df, data = self._file()
        edited, _ = fill_fields(data, {'Client Name': 'K. Kannan'})
        store_version(df, edited, self.owner.id)
        self.b1.refresh_from_db()
        self.b2.refresh_from_db()
        self.assertEqual(self.b1.text, 'This deed is made by K. Kannan of Chennai.')   # heading kept apart
        self.assertTrue(self.b1.is_edited)
        self.assertEqual(self.b2.text, 'Rent is payable monthly.')
        self.assertFalse(self.b2.is_edited)

    def test_token_split_across_runs_is_replaced_keeping_the_first_runs_format(self):
        import io as _io

        from docx import Document
        from .docfile import fill_placeholder, placeholders
        doc = Document()
        p = doc.add_paragraph('Name: ')
        p.add_run('[[Cli').bold = True
        p.add_run('ent Name]]')
        p.add_run(' signs.')
        buf = _io.BytesIO()
        doc.save(buf)
        out, n = fill_placeholder(buf.getvalue(), 'Client Name', 'K. Kannan')
        self.assertEqual(n, 1)
        para = Document(_io.BytesIO(out)).paragraphs[0]
        self.assertEqual(para.text, 'Name: K. Kannan signs.')
        self.assertTrue(para.runs[1].bold and para.runs[1].text == 'K. Kannan')
        self.assertEqual(placeholders(out), [])

    def test_placeholder_endpoints(self):
        url = f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/placeholders/'
        self.assertEqual(self.client.get(url, **auth(self.owner)).json()['placeholders'], [])     # no file yet
        self._file()
        self.assertEqual(self.client.get(url, **auth(self.owner)).json()['placeholders'],
                         [{'label': 'Client Name', 'count': 1, 'value': '', 'kind': 'field'}])
        resp = self.client.post(url, {'label': 'Client Name', 'value': 'K. Kannan'}, content_type='application/json',
                                **auth(self.owner))
        self.assertEqual(resp.json()['filled'], 1)
        self.assertIn('key', resp.json()['config']['document'])
        # Filled, and still listed with its value so it can be changed later.
        self.assertEqual(self.client.get(url, **auth(self.owner)).json()['placeholders'],
                         [{'label': 'Client Name', 'count': 1, 'value': 'K. Kannan', 'kind': 'field'}])
        changed = self.client.post(url, {'values': {'Client Name': 'K. Kannan & Sons'}}, content_type='application/json',
                                   **auth(self.owner))
        self.assertEqual(changed.json()['filled'], 1)
        self.b1.refresh_from_db()
        self.assertIn('K. Kannan & Sons', self.b1.text)
        self.assertIn('K. Kannan', self.b1.text)
        stranger = make_advocate(permissions=ALL_PERMISSIONS)
        self.assertEqual(self.client.post(url, {'label': 'X', 'value': 'Y'}, content_type='application/json',
                                          **auth(stranger)).status_code, 404)

    def test_unedited_file_without_markers_is_made_again(self):
        from django.core.files.base import ContentFile
        from .docfile import clause_texts
        from .models import DraftFile
        df = DraftFile(session=self.session, key='old', version=1)
        df.file.save('old.docx', ContentFile(self._plain_docx()), save=True)
        new_df, data = self._file()
        self.assertNotEqual(new_df.pk, df.pk)
        self.assertTrue(clause_texts(data))

    @staticmethod
    def _plain_docx():
        import io as _io

        from docx import Document
        buf = _io.BytesIO()
        Document().save(buf)
        return buf.getvalue()


class HiddenMarkersTest(TestCase):
    """Clause markers stay invisible: new files, and files saved back with a frame / label."""

    def test_framed_marker_is_hidden_again(self):
        import io as _io

        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from .docfile import W15, clause_texts, hide_markers
        doc = Document()
        p = doc.add_paragraph('Rent is payable monthly.')
        sdt, pr, content = OxmlElement('w:sdt'), OxmlElement('w:sdtPr'), OxmlElement('w:sdtContent')
        alias, tag = OxmlElement('w:alias'), OxmlElement('w:tag')
        alias.set(qn('w:val'), 'Clause 7')
        tag.set(qn('w:val'), 'pactpro:block:7')
        pr.extend([alias, tag])
        sdt.extend([pr, content])
        p._p.addprevious(sdt)
        content.append(p._p)
        buf = _io.BytesIO()
        doc.save(buf)
        out = hide_markers(buf.getvalue())
        pr = Document(_io.BytesIO(out)).element.body.find(f'{qn("w:sdt")}/{qn("w:sdtPr")}')
        self.assertIsNone(pr.find(qn('w:alias')))
        self.assertEqual(pr.find(f'{{{W15}}}appearance').get(f'{{{W15}}}val'), 'hidden')
        self.assertEqual(clause_texts(out), {7: ['Rent is payable monthly.']})
        self.assertIs(hide_markers(out), out)                        # nothing left to change


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class OnlyOfficeReviewTest(TestCase):
    """Phase 3: editor roles, open tracked changes / comments in the Word file block approval."""

    def setUp(self):
        self.media_override = override_settings(MEDIA_ROOT=tempfile.mkdtemp())
        self.media_override.enable()
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.session.blocks.create(position=0, block_type='clause', heading='', source='generated', text='Rent is Rs. 40,000.')

    def tearDown(self):
        self.media_override.disable()

    def _role_of(self, who):
        resp = self.client.get(f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/', **auth(who))
        return resp.json() if resp.status_code == 200 else resp.status_code

    def _reviewer(self, authority):
        from .models import DraftReviewRequest
        who = make_advocate(permissions=ALL_PERMISSIONS)
        DraftReviewRequest.objects.create(session=self.session, requested_by_id=self.owner.id, reviewer_id=who.id,
                                          authority=authority)
        return who

    def test_roles(self):
        owner = self._role_of(self.owner)
        self.assertEqual(owner['role'], 'decide')
        perms = owner['config']['document']['permissions']
        self.assertEqual((perms['edit'], perms['review'], perms['comment']), (True, True, True))
        self.assertTrue(owner['config']['editorConfig']['customization']['layout']['toolbar']['collaboration'])

        suggester = self._role_of(self._reviewer('suggest'))
        self.assertEqual(suggester['role'], 'suggest')
        perms = suggester['config']['document']['permissions']
        self.assertEqual((perms['edit'], perms['review']), (False, True))           # review-only: edits are tracked
        self.assertEqual(suggester['config']['editorConfig']['mode'], 'edit')

        self.assertEqual(self._role_of(self._reviewer('binding'))['role'], 'decide')
        self.assertEqual(self._role_of(make_advocate(permissions=ALL_PERMISSIONS)), 404)

    def _file_with(self, tracked=0, comments=0, resolved=0):
        """The draft's Word file with tracked insertions and comment threads (some resolved)."""
        import io as _io
        import zipfile

        from .onlyoffice import ensure_file, store_version
        df = ensure_file(self.session, self.owner)
        with df.versions.order_by('version').first().file.open('rb') as f:     # the file as first made
            data = f.read()
        src = zipfile.ZipFile(_io.BytesIO(data))
        out = _io.BytesIO()
        w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        w15 = 'http://schemas.microsoft.com/office/word/2012/wordml'
        with zipfile.ZipFile(out, 'w') as z:
            for item in src.infolist():
                body = src.read(item.filename)
                if item.filename == 'word/document.xml' and tracked:
                    ins = ''.join(f'<w:ins w:id="{90 + n}" w:author="Lisa" w:date="2026-10-12T00:00:00Z"><w:r><w:t>x</w:t></w:r></w:ins>'
                                  for n in range(tracked))
                    body = body.replace(b'</w:body>', f'<w:p>{ins}</w:p></w:body>'.encode())
                z.writestr(item, body)
            if comments:
                z.writestr('word/comments.xml', f'<w:comments xmlns:w="{w}">' + ''.join(
                    f'<w:comment w:id="{n}" w:author="R"><w:p><w:r><w:t>q</w:t></w:r></w:p></w:comment>'
                    for n in range(comments)) + '</w:comments>')
                z.writestr('word/commentsExtended.xml', f'<w15:commentsEx xmlns:w15="{w15}">' + ''.join(
                    f'<w15:commentEx w15:paraId="{n:08X}" w15:done="{1 if n < resolved else 0}"/>'
                    for n in range(comments)) + '</w15:commentsEx>')
        store_version(df, out.getvalue(), self.owner.id)

    def test_review_state_counts_open_changes_and_threads(self):
        from .docfile import review_state
        from .models import DraftFile
        self._file_with(tracked=2, comments=3, resolved=1)
        with DraftFile.objects.get(session=self.session).file.open('rb') as f:
            self.assertEqual(review_state(f.read()), {'changes': 2, 'comments': 2})

    def test_open_changes_block_submit_and_approve(self):
        from .onlyoffice import file_review_blockers
        from .task_review import submit_blocker
        self.assertEqual(file_review_blockers(self.session), [])                  # no file yet
        self._file_with(tracked=1, comments=1)
        blockers = file_review_blockers(self.session)
        self.assertEqual(len(blockers), 2)
        self.assertIn('tracked change', submit_blocker(self.session))
        self._file_with(comments=1)                                                # changes decided, a question open
        self.assertIsNone(submit_blocker(self.session))                            # open comments don't block submitting
        self.assertEqual(len(file_review_blockers(self.session)), 1)               # ...but do block approving


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class ServerTrackedChangeTest(TestCase):
    """Lisa's Apply on the Community edition: written into the Word file as tracked changes by PactPro AI."""

    RENT = 'The Lessee shall pay a monthly rent of Rs. 40,000 on or before the 5th day of each month.'

    def setUp(self):
        self.media_override = override_settings(MEDIA_ROOT=tempfile.mkdtemp())
        self.media_override.enable()
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.block = self.session.blocks.create(position=0, block_type='clause', heading='', source='generated',
                                                text=self.RENT)
        self.session.blocks.create(position=1, block_type='clause', heading='', source='generated',
                                   text='A deposit of Rs. 2,40,000 is paid.')

    def tearDown(self):
        self.media_override.disable()

    def test_split_runs_keep_formatting_and_only_the_change_is_marked(self):
        import io as _io

        from docx import Document
        from docx.oxml.ns import qn
        from .docfile import AI_AUTHOR, apply_tracked, review_state
        doc = Document()
        p = doc.add_paragraph('Rent is Rs. ')
        p.add_run('40,').bold = True
        p.add_run('000 a month.')
        buf = _io.BytesIO()
        doc.save(buf)
        out, ok = apply_tracked(buf.getvalue(), 'Rent is Rs. 40,000 a month.', 'Rent is Rs. 45,000 a month.')
        self.assertTrue(ok)
        body = Document(_io.BytesIO(out)).element.body
        dele, ins = body.find('.//' + qn('w:del')), body.find('.//' + qn('w:ins'))
        self.assertEqual(''.join(t.text for t in dele.iter(qn('w:delText'))), '0')
        self.assertEqual(''.join(t.text for t in ins.iter(qn('w:t'))), '5')
        self.assertEqual(ins.get(qn('w:author')), AI_AUTHOR)
        self.assertEqual(review_state(out)['changes'], 2)
        # Accepted view (what the draft's clause text shows): the new amount.
        from .docfile import _para_text
        self.assertEqual(_para_text(body.find('.//' + qn('w:p'))), 'Rent is Rs. 45,000 a month.')

    def test_multi_paragraph_clause_with_a_blank_and_repeated_words_is_placed(self):
        """Lisa reads a clause as lines with [[Label]] blanks; the file has one paragraph per line, the blank
        as a form field showing its label, and the same signature line in two clauses (found by draft #114:
        each of these used to answer "the clause has changed since Lisa read it")."""
        import io as _io

        from docx import Document
        from docx.oxml.ns import qn
        from .docfile import apply_tracked, building, record, tokens_to_fields, wrap_blocks
        sign = 'Signed in My Presence. The Client.'
        doc = Document()
        token = building.set({})
        try:
            record(7, [doc.add_paragraph(t)._p for t in ('VERSUS', '[[Details]]Defendant(s)/Respondent(s)', sign)])
            record(8, [doc.add_paragraph(t)._p for t in ('Witness', sign)])
            wrap_blocks(doc)
        finally:
            building.reset(token)
        buf = _io.BytesIO()
        doc.save(buf)
        data = tokens_to_fields(buf.getvalue())
        before = f'VERSUS\n[[Details]]Defendant(s)/Respondent(s)\n{sign}'
        after = 'VERSUS\n[[Details]]Defendant(s) / Respondent(s)\nSigned in my presence. The Client.'
        out, ok = apply_tracked(data, before, after, block_id=7)
        self.assertTrue(ok)
        body = Document(_io.BytesIO(out)).element.body
        changed = [p for p in body.iter(qn('w:p')) if p.find('.//' + qn('w:ins')) is not None]
        self.assertEqual(len(changed), 2)                  # one line each, both inside clause 7
        clause8 = [s for s in body.iter(qn('w:sdt'))
                   if (s.find(f'{qn("w:sdtPr")}/{qn("w:tag")}').get(qn('w:val'))) == 'pactpro:block:8'][0]
        self.assertIsNone(clause8.find('.//' + qn('w:ins')))
        # Without the clause, the repeated line can't be pinned to one place.
        self.assertFalse(apply_tracked(data, sign, 'Signed in my presence. The Client.')[1])

    def test_lines_joined_by_the_model_still_place_the_word_change(self):
        """Lisa's text came back with two of the clause's lines joined by a space (draft #114, the enrolment
        number): the words go back onto the file's lines and only the changed number is marked."""
        import io as _io

        from docx import Document
        from docx.oxml.ns import qn
        from .docfile import apply_tracked
        doc = Document()
        doc.add_paragraph('Signed in My Presence. The Client.')
        doc.add_paragraph('Sd/S. Meenakshi, Advocate Enrolment No. MS/2210/2014 Place: Chennai')
        buf = _io.BytesIO()
        doc.save(buf)
        before = 'Signed in My Presence. The Client.\nSd/S. Meenakshi, Advocate Enrolment No. MS/2210/2014 Place: Chennai'
        after = 'Signed in My Presence. The Client. Sd/S. Meenakshi, Advocate Enrolment No. MS/2223/2026 Place: Chennai'
        out, ok = apply_tracked(buf.getvalue(), before, after)
        self.assertTrue(ok)
        body = Document(_io.BytesIO(out)).element.body
        self.assertEqual(len(list(body.iter(qn('w:p')))), 2)       # still two paragraphs
        self.assertEqual([''.join(t.text for t in d.iter(qn('w:delText'))) for d in body.iter(qn('w:del'))], ['10/2014'])
        self.assertEqual([''.join(t.text for t in i.iter(qn('w:t'))) for i in body.iter(qn('w:ins'))], ['23/2026'])

    def test_endpoint_writes_redline_and_refuses_when_it_cannot_place_it(self):
        from .docfile import clause_texts, review_state
        from .models import DraftFile
        url = f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/apply/'
        resp = self.client.post(url, {'before': self.RENT, 'after': self.RENT.replace('40,000', '45,000')},
                                content_type='application/json', **auth(self.owner))
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertIn('key', resp.json()['config']['document'])
        with DraftFile.objects.get(session=self.session).file.open('rb') as f:
            data = f.read()
        self.assertEqual(review_state(data)['changes'], 2)
        self.assertIn('Rs. 45,000', clause_texts(data)[self.block.id][0])
        self.block.refresh_from_db()
        self.assertIn('45,000', self.block.text)
        gone = self.client.post(url, {'before': 'Text that is not in the draft.', 'after': 'Other text.'},
                                content_type='application/json', **auth(self.owner))
        self.assertEqual(gone.status_code, 409)
        stranger = make_advocate(permissions=ALL_PERMISSIONS)
        self.assertEqual(self.client.post(url, {'before': 'a', 'after': 'b'}, content_type='application/json',
                                          **auth(stranger)).status_code, 404)

    def test_several_changes_in_one_request_make_one_version(self):
        """Lisa's "Apply all": every change in one pass, one version; ones that can't be placed are counted."""
        from .docfile import review_state
        from .models import DraftFile
        url = f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/apply/'
        self.client.get(f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/', **auth(self.owner))
        before = DraftFile.objects.get(session=self.session).version
        resp = self.client.post(url, {'changes': [
            {'before': self.RENT, 'after': self.RENT.replace('40,000', '45,000')},
            {'before': 'Text that is not in the draft.', 'after': 'Other text.'},
        ]}, content_type='application/json', **auth(self.owner))
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual((resp.json()['applied'], resp.json()['skipped']), (1, 1))
        df = DraftFile.objects.get(session=self.session)
        self.assertEqual(df.version, before + 1)
        with df.file.open('rb') as f:
            self.assertEqual(review_state(f.read())['changes'], 2)


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class OnlyOfficeHistoryTest(TestCase):
    """The editor's own version history, restore and compare, fed by PactPro."""

    def setUp(self):
        self.media_override = override_settings(MEDIA_ROOT=tempfile.mkdtemp())
        self.media_override.enable()
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.session.blocks.create(position=0, block_type='clause', heading='', source='generated', text='Rent is Rs. 40,000.')
        from .onlyoffice import ensure_file
        self.df = ensure_file(self.session, self.owner)

    def tearDown(self):
        self.media_override.disable()

    def _save_v2(self):
        from .docfile import apply_tracked
        from .onlyoffice import store_version
        with self.df.file.open('rb') as f:
            data, _ = apply_tracked(f.read(), 'Rent is Rs. 40,000.', 'Rent is Rs. 45,000.')
        self.df = store_version(self.df, data, self.owner.id, changes=b'PK-zip',
                                history={'serverVersion': '9.4.0', 'changes': [{'created': '2026-10-13 10:00:00'}]})

    def _url(self, tail=''):
        return f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/{tail}'

    def test_history_list_and_one_version_with_highlighting(self):
        import jwt
        from django.conf import settings
        self._save_v2()
        listing = self.client.get(self._url('history/'), **auth(self.owner)).json()
        self.assertEqual(listing['currentVersion'], 2)
        self.assertEqual([h['version'] for h in listing['history']], [1, 2])
        self.assertEqual(listing['history'][1]['serverVersion'], '9.4.0')
        self.assertNotEqual(listing['history'][0]['key'], listing['history'][1]['key'])
        one = self.client.get(self._url('history/2/'), **auth(self.owner)).json()
        self.assertEqual(one['previous']['key'], listing['history'][0]['key'])
        self.assertIn('changesUrl', one)
        self.assertEqual(jwt.decode(one['token'], settings.ONLYOFFICE_JWT_SECRET, algorithms=['HS256'])['version'], 2)
        # The document server fetches the files with the signed links, no login.
        for link in (one['url'], one['previous']['url'], one['changesUrl']):
            resp = self.client.get(link.replace('http://ams.test', ''))
            self.assertEqual(resp.status_code, 200, link)
        self.assertEqual(self.client.get(one['url'].replace('http://ams.test', '').split('?')[0] + '?t=x').status_code, 404)

    def test_restore_makes_a_new_version_and_only_who_decides(self):
        from .models import DraftReviewRequest
        self._save_v2()
        suggester = make_advocate(permissions=ALL_PERMISSIONS)
        DraftReviewRequest.objects.create(session=self.session, requested_by_id=self.owner.id,
                                          reviewer_id=suggester.id, authority='suggest')
        self.assertEqual(self.client.post(self._url('restore/'), {'version': 1}, content_type='application/json',
                                          **auth(suggester)).status_code, 403)
        resp = self.client.post(self._url('restore/'), {'version': 1}, content_type='application/json', **auth(self.owner))
        self.assertEqual(resp.status_code, 200)
        self.df.refresh_from_db()
        self.assertEqual(self.df.version, 3)                                   # nothing lost: v2 stays in history
        self.session.blocks.first().refresh_from_db()
        self.assertIn('40,000', self.session.blocks.first().text)

    def test_compare_with_a_version_or_an_uploaded_file(self):
        import jwt
        from django.conf import settings
        self._save_v2()
        by_version = self.client.post(self._url('compare/'), {'version': 1}, content_type='application/json',
                                      **auth(self.owner)).json()
        self.assertEqual(by_version['c'], 'compare')
        self.assertIn('/onlyoffice/version/', by_version['url'])
        with self.df.file.open('rb') as f:
            returned = SimpleUploadedFile('returned.docx', f.read())
        uploaded = self.client.post(self._url('compare/'), {'file': returned}, **auth(self.owner)).json()
        payload = jwt.decode(uploaded['token'], settings.ONLYOFFICE_JWT_SECRET, algorithms=['HS256'])
        self.assertEqual(payload['url'], uploaded['url'])
        self.assertEqual(self.client.get(uploaded['url'].replace('http://ams.test', '')).status_code, 200)
        bad = self.client.post(self._url('compare/'), {'file': SimpleUploadedFile('x.pdf', b'%PDF')}, **auth(self.owner))
        self.assertEqual(bad.status_code, 400)


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class OnlyOfficeFinaliseTest(TestCase):
    """Finalise / reopen, PDF/A, @mentions, and the DRAFT watermark on copies of an unfinished draft."""

    def setUp(self):
        self.media_override = override_settings(MEDIA_ROOT=tempfile.mkdtemp())
        self.media_override.enable()
        self.owner = make_advocate(permissions=ALL_PERMISSIONS)
        self.session = DraftSession.objects.create(created_by_id=self.owner.id, facts={}, status='ready')
        self.session.blocks.create(position=0, block_type='clause', heading='', source='generated', text='Rent is Rs. 40,000.')
        from .onlyoffice import ensure_file
        self.df = ensure_file(self.session, self.owner)

    def tearDown(self):
        self.media_override.disable()

    def _url(self, tail):
        return f'/api/drafting/draft-sessions/{self.session.id}/onlyoffice/{tail}'

    def test_finalise_locks_and_reopen_unlocks(self):
        from .docfile import apply_tracked
        from .onlyoffice import store_version
        # Open tracked changes refuse finalising.
        with self.df.file.open('rb') as f:
            data, _ = apply_tracked(f.read(), 'Rent is Rs. 40,000.', 'Rent is Rs. 45,000.')
        store_version(self.df, data, self.owner.id)
        refused = self.client.post(self._url('finalise/'), **auth(self.owner))
        self.assertEqual(refused.status_code, 409)
        self.assertIn('tracked change', refused.json()['error'])
        # Clean file: finalised, everyone read-only, saves no longer taken.
        with self.df.versions.order_by('version').first().file.open('rb') as f:
            store_version(self.df, f.read(), self.owner.id)
        done = self.client.post(self._url('finalise/'), **auth(self.owner))
        self.assertEqual(done.status_code, 200, done.content)
        self.assertEqual(len(done.json()['final']['sha256']), 64)
        state = self.client.get(self._url(''), **auth(self.owner)).json()
        self.assertEqual(state['role'], 'view')
        self.assertEqual(state['config']['editorConfig']['mode'], 'view')
        self.assertEqual(self.client.post(self._url('apply/'), {'before': 'Rent', 'after': 'Fee'},
                                          content_type='application/json', **auth(self.owner)).status_code, 403)
        # Reopen: editable again with a fresh key.
        old_key = state['config']['document']['key']
        reopened = self.client.delete(self._url('finalise/'), **auth(self.owner)).json()
        self.assertIsNone(reopened['final'])
        self.assertNotEqual(reopened['config']['document']['key'], old_key)
        self.assertEqual(self.client.get(self._url(''), **auth(self.owner)).json()['role'], 'decide')

    def test_pdfa_comes_from_the_document_server(self):
        with mock.patch('drafting.onlyoffice.convert', return_value=b'%PDF-1.4 pdfaid') as conv:
            resp = self.client.get(self._url('pdfa/'), **auth(self.owner))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(conv.call_args[0][1], 'pdfa')
        self.assertIn('PDF-A', resp['Content-Disposition'])

    def test_mentions_offer_and_notify_only_people_on_the_draft(self):
        from .models import DraftReviewRequest
        reviewer = make_advocate(permissions=ALL_PERMISSIONS)
        outsider = make_advocate(permissions=ALL_PERMISSIONS)
        DraftReviewRequest.objects.create(session=self.session, requested_by_id=self.owner.id,
                                          reviewer_id=reviewer.id, authority='suggest')
        users = self.client.get(self._url('users/'), **auth(self.owner)).json()['users']
        self.assertEqual([u['email'] for u in users], [reviewer.email])
        with mock.patch('notifications.service.notify') as notify:
            resp = self.client.post(self._url('notify/'), {'emails': [reviewer.email, outsider.email], 'message': 'Check this'},
                                    content_type='application/json', **auth(self.owner))
        self.assertEqual(resp.json()['notified'], 1)
        self.assertEqual(notify.call_args[0][0], reviewer.id)

    def test_download_has_draft_watermark_until_final(self):
        import io as _io
        from docx import Document

        def marks(resp):
            doc = Document(_io.BytesIO(resp.content))
            return sum(b'PactProWatermark' in __import__('lxml').etree.tostring(s.header._element) for s in doc.sections)

        url = f'/api/drafting/drafts/{self.session.id}/export/docx/'
        self.assertEqual(marks(self.client.get(url, **auth(self.owner))), 1)
        self.assertEqual(marks(self.client.get(url + '?watermark=0', **auth(self.owner))), 0)
        self.client.post(self._url('finalise/'), **auth(self.owner))
        self.assertEqual(marks(self.client.get(url, **auth(self.owner))), 0)


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class OnlyOfficeStateTest(TestCase):
    """The page's Save button follows saves made in the editor itself: it polls the stored version."""

    def test_version_goes_up_with_each_stored_save(self):
        from .onlyoffice import ensure_file, store_version
        with override_settings(MEDIA_ROOT=tempfile.mkdtemp()):
            owner = make_advocate(permissions=ALL_PERMISSIONS)
            session = DraftSession.objects.create(created_by_id=owner.id, facts={}, status='ready')
            session.blocks.create(position=0, block_type='clause', heading='', source='generated', text='Text.')
            url = f'/api/drafting/draft-sessions/{session.id}/onlyoffice/state/'
            self.assertEqual(self.client.get(url, **auth(owner)).json()['version'], 0)   # no file yet
            df = ensure_file(session, owner)
            self.assertEqual(self.client.get(url, **auth(owner)).json()['version'], 1)
            with df.file.open('rb') as f:
                store_version(df, f.read(), owner.id)                             # e.g. Ctrl+S in the editor
            self.assertEqual(self.client.get(url, **auth(owner)).json()['version'], 2)
            self.assertEqual(self.client.get(url, **auth(make_advocate(permissions=ALL_PERMISSIONS))).status_code, 404)


@override_settings(DRAFT_EDITOR='onlyoffice', ONLYOFFICE_URL='http://oo.test', ONLYOFFICE_INTERNAL_URL='http://oo.test',
                   ONLYOFFICE_JWT_SECRET='test-secret-test-secret-test-secret!', AMS_INTERNAL_URL='http://ams.test')
class PlaceholdersApplyAllTest(TestCase):
    """The panel's "Apply to document": every value typed so far, in one version (one editor reload)."""

    def test_several_values_one_version_empty_ones_stay_blank(self):
        from .models import DraftFile
        from .onlyoffice import ensure_file
        with override_settings(MEDIA_ROOT=tempfile.mkdtemp()):
            owner = make_advocate(permissions=ALL_PERMISSIONS)
            session = DraftSession.objects.create(created_by_id=owner.id, facts={}, status='ready')
            session.blocks.create(position=0, block_type='clause', heading='', source='generated',
                                  text='Made on [[Agreement Date]] by [[Lessor Name]] at [[Place]].')
            ensure_file(session, owner)
            url = f'/api/drafting/draft-sessions/{session.id}/onlyoffice/placeholders/'
            resp = self.client.post(url, {'values': {'Agreement Date': '1 November 2026', 'Lessor Name': 'K. Kannan',
                                                     'Place': '  ', 'Gone': 'x'}},
                                    content_type='application/json', **auth(owner))
            self.assertEqual(resp.status_code, 200, resp.content)
            self.assertEqual((resp.json()['filled'], resp.json()['missing']), (2, ['Gone']))
            df = DraftFile.objects.get(session=session)
            self.assertEqual(df.version, 2)                                       # one new version for all
            from .docfile import fields
            with df.file.open('rb') as f:
                state = fields(f.read())
            self.assertEqual((state['Agreement Date']['value'], state['Place']['value']), ('1 November 2026', ''))
            empty = self.client.post(url, {'values': {'Place': ''}}, content_type='application/json', **auth(owner))
            self.assertEqual(empty.status_code, 400)
