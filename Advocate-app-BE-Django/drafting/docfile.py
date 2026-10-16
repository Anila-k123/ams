"""The draft as a Word file (ONLYOFFICE, docs/ONLYOFFICE_PLAN.md phase 2): clause markers, the text copy
back into the draft's clauses, and placeholders filled in the file.

Each clause is a block-level content control tagged `pactpro:block:<id>` and locked against deletion
(its content stays editable), so per-clause features (Lisa, citations, review) still know which text
is which clause after people edit the file. Placeholders are written as `[[Label]]` text.
"""

import copy
import io
import re
import secrets
from contextvars import ContextVar
from html import escape

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from docx.oxml.ns import qn

TAG = 'pactpro:block:'
W15 = 'http://schemas.microsoft.com/office/word/2012/wordml'
PLACEHOLDER = re.compile(r'\[\[([^\[\]]{1,80})\]\]')

# Set while ensure_file builds the file: the renderers then record which body elements each clause
# became ({block id: [elements]}) and write empty placeholders as [[Label]].
building = ContextVar('drafting_docfile_building', default=None)


def record(block_id, elements):
    """Called by the renderers: these body elements are clause `block_id`."""
    captured = building.get()
    if captured is not None and block_id is not None and elements:
        captured.setdefault(block_id, []).extend(e for e in elements if e is not None)


def wrap_blocks(doc):
    """Wrap each recorded clause's elements in its content control (renderers call this before save)."""
    captured = building.get()
    if not captured:
        return
    body = doc.element.body
    for block_id, els in captured.items():
        els = [e for e in els if e.getparent() is body]
        if not els:
            continue
        els.sort(key=lambda e: body.index(e))
        sdt = OxmlElement('w:sdt')
        pr = OxmlElement('w:sdtPr')
        tag = OxmlElement('w:tag')
        tag.set(qn('w:val'), f'{TAG}{block_id}')
        lock = OxmlElement('w:lock')
        lock.set(qn('w:val'), 'sdtLocked')                 # the marker can't be deleted; its text can
        # Invisible to the user: no frame, no title label (Word 2013+ appearance "hidden"; no w:alias).
        appearance = etree.SubElement(pr, f'{{{W15}}}appearance')
        appearance.set(f'{{{W15}}}val', 'hidden')
        pr.insert(0, tag)
        pr.insert(1, lock)
        content = OxmlElement('w:sdtContent')
        sdt.extend([pr, content])
        els[0].addprevious(sdt)
        for e in els:
            content.append(e)                              # moves it into the control


def _hide(pr):
    """Make one clause control invisible: no title label, appearance "hidden". True when it changed."""
    changed = False
    for alias in pr.findall(qn('w:alias')):
        pr.remove(alias)
        changed = True
    found = pr.find(f'{{{W15}}}appearance')
    if found is None:
        found = etree.SubElement(pr, f'{{{W15}}}appearance')
        changed = True
    if found.get(f'{{{W15}}}val') != 'hidden':
        found.set(f'{{{W15}}}val', 'hidden')
        changed = True
    return changed


def hide_markers(data):
    """The file with every clause control invisible (files made before the markers were hidden, and
    files the editor saved back with a frame). Returns the bytes unchanged when nothing to do."""
    doc = Document(io.BytesIO(data))
    changed = False
    for sdt in doc.element.body.iter(qn('w:sdt')):
        pr = sdt.find(qn('w:sdtPr'))
        tag = pr.find(qn('w:tag')) if pr is not None else None
        if tag is not None and (tag.get(qn('w:val')) or '').startswith(TAG):
            changed = _hide(pr) or changed
    if not changed:
        return data
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- reading the file back ----------------------------------------------------------------------

def _para_text(p):
    """A paragraph's text; an empty form field reads as its `[[Label]]` blank (the hint text isn't text)."""
    out = []

    def walk(el):
        for node in el:
            if node.tag == qn('w:sdt') and _field_label(node) is not None:
                if _field_empty(node):
                    out.append(f'[[{_field_label(node)}]]')
                    continue
            if node.tag == qn('w:t'):
                out.append(node.text or '')
            elif node.tag == qn('w:tab'):
                out.append('\t')
            elif node.tag in (qn('w:br'), qn('w:cr')):
                out.append('\n')
            elif node.tag != qn('w:sdtPr'):
                walk(node)
    walk(p)
    return ''.join(out)


# --- form fields: the blanks of new drafts (docs/ONLYOFFICE_PLAN.md, "Form fields") ------------------
# A blank is a Word plain-text content control tagged `pactpro:field:<Label>`, showing its label as a grey
# hint until filled. Unlike `[[Label]]` text it can't be half-deleted by accident, and a value can be
# changed later from the Placeholders panel (every field with the label at once).

FIELD = 'pactpro:field:'


def _field_label(sdt):
    tag = sdt.find(f'{qn("w:sdtPr")}/{qn("w:tag")}')
    val = (tag.get(qn('w:val')) if tag is not None else '') or ''
    return val[len(FIELD):] if val.startswith(FIELD) else None


def _field_empty(sdt):
    return sdt.find(f'{qn("w:sdtPr")}/{qn("w:showingPlcHdr")}') is not None


def _field_runs(sdt):
    content = sdt.find(qn('w:sdtContent'))
    return content, ([r for r in content.iter(qn('w:r'))] if content is not None else [])


_HINT_COLOR, _HINT_FILL = '7F7F7F', 'E7E6E6'


def _own_rpr(sdt):
    """The field's own text formatting (kept in its properties, w:sdtPr/w:rPr), so a value never takes the
    grey hint look. Older fields have none yet: taken from their run, without the hint marks."""
    pr = sdt.find(qn('w:sdtPr'))
    own = pr.find(qn('w:rPr'))
    if own is None:
        _, runs = _field_runs(sdt)
        first = runs[0].find(qn('w:rPr')) if runs else None
        own = copy.deepcopy(first) if first is not None else OxmlElement('w:rPr')
        for tag in ('w:color', 'w:shd'):
            for el in own.findall(qn(tag)):
                if el.get(qn('w:val')) == _HINT_COLOR or el.get(qn('w:fill')) == _HINT_FILL:
                    own.remove(el)
        pr.append(own)
    return own


def _hint_rpr(own):
    """An empty field's look: its own formatting in grey on light grey, so every blank stands out."""
    rpr = copy.deepcopy(own)
    rpr.tag = qn('w:rPr')
    for tag in ('w:color', 'w:shd'):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), _HINT_COLOR)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), _HINT_FILL)
    rpr.extend([color, shd])
    return rpr


def _set_field(sdt, value):
    """Put `value` in the field in its own formatting; an empty value shows the grey hint again."""
    label = _field_label(sdt)
    own = _own_rpr(sdt)
    content, _ = _field_runs(sdt)
    for child in list(content):
        content.remove(child)
    r = OxmlElement('w:r')
    pr = sdt.find(qn('w:sdtPr'))
    hint = pr.find(qn('w:showingPlcHdr'))
    if value:
        if hint is not None:
            pr.remove(hint)
        if len(own):
            r.append(copy.deepcopy(own))
        _put_text(r, value)
    else:
        if hint is None:
            pr.append(OxmlElement('w:showingPlcHdr'))
        r.append(_hint_rpr(own))
        _put_text(r, label)
    content.append(r)


def style_fields(data):
    """Give fields their look (empty: grey hint; filled: own formatting). For fields made before the hint
    look existed, and after the editor saves (a value typed into a field keeps the hint's grey otherwise).
    Returns the bytes unchanged when nothing to do."""
    doc = Document(io.BytesIO(data))
    changed = False
    for sdt in list(doc.element.body.iter(qn('w:sdt'))):
        if _field_label(sdt) is None:
            continue
        had_own = sdt.find(f'{qn("w:sdtPr")}/{qn("w:rPr")}') is not None
        own = _own_rpr(sdt)
        content, runs = _field_runs(sdt)
        want_hint = _field_empty(sdt)
        for r in runs:
            rpr = r.find(qn('w:rPr'))
            grey = rpr is not None and any(c.get(qn('w:val')) == _HINT_COLOR for c in rpr.findall(qn('w:color')))
            if want_hint and not grey:
                if rpr is not None:
                    r.remove(rpr)
                r.insert(0, _hint_rpr(own))
                changed = True
            elif not want_hint and grey:
                r.remove(rpr)
                if len(own):
                    r.insert(0, copy.deepcopy(own))
                changed = True
        changed = changed or not had_own
    if not changed:
        return data
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _make_field(label, rpr):
    sdt = OxmlElement('w:sdt')
    pr = OxmlElement('w:sdtPr')
    alias = OxmlElement('w:alias')
    alias.set(qn('w:val'), label)                          # the field's name when the user clicks it
    tag = OxmlElement('w:tag')
    tag.set(qn('w:val'), f'{FIELD}{label}')
    lock = OxmlElement('w:lock')
    lock.set(qn('w:val'), 'sdtLocked')                     # the field can't be deleted; its text can change
    own = copy.deepcopy(rpr) if rpr is not None else OxmlElement('w:rPr')
    pr.extend([alias, tag, lock, own, OxmlElement('w:showingPlcHdr'), OxmlElement('w:text')])
    content = OxmlElement('w:sdtContent')
    r = OxmlElement('w:r')
    r.append(_hint_rpr(own))
    _put_text(r, label)
    content.append(r)
    sdt.extend([pr, content])
    return sdt


def tokens_to_fields(data):
    """New drafts: every `[[Label]]` blank in the file becomes a form field (when the file is made)."""
    doc = Document(io.BytesIO(data))
    changed = False
    for p in [p for p in doc.element.body.iter(qn('w:p'))]:
        while True:
            runs = _text_runs(p)
            text = ''.join(_run_text(r) for r in runs)
            m = PLACEHOLDER.search(text)
            if not m:
                break
            start, end = m.span()
            for edge in (start, end):                     # the token on run edges
                pos = 0
                for r in _text_runs(p):
                    n = len(_run_text(r))
                    if pos < edge < pos + n:
                        _split(r, edge - pos)
                        break
                    pos += n
            pos, inside = 0, []
            for r in _text_runs(p):
                n = len(_run_text(r))
                if start <= pos and pos + n <= end:
                    inside.append(r)
                pos += n
            if not inside:
                break
            field = _make_field(m.group(1).strip(), inside[0].find(qn('w:rPr')))
            inside[0].addprevious(field)
            for r in inside:
                r.getparent().remove(r)
            changed = True
    if not changed:
        return data
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def fields(data):
    """{label: {'count', 'value'}} of the file's form fields ('' while a field shows its hint)."""
    doc = Document(io.BytesIO(data))
    out = {}
    for sdt in doc.element.body.iter(qn('w:sdt')):
        label = _field_label(sdt)
        if label is None:
            continue
        item = out.setdefault(label, {'count': 0, 'value': ''})
        item['count'] += 1
        if not _field_empty(sdt) and not item['value']:
            content, _ = _field_runs(sdt)
            item['value'] = ''.join(_run_text(r) for r in content.iter(qn('w:r'))) if content is not None else ''
    return out


def fill_fields(data, values):
    """The file with every field of each label in `values` set to its value ('' empties it again).
    Returns (bytes, {label: how many fields})."""
    doc = Document(io.BytesIO(data))
    done = {}
    for sdt in list(doc.element.body.iter(qn('w:sdt'))):
        label = _field_label(sdt)
        if label is not None and label in values:
            _set_field(sdt, values[label])
            done[label] = done.get(label, 0) + 1
    if not done:
        return data, done
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), done


def _lines(content):
    """Paragraph texts of a control's content, in order; a table row is one `a | b` line."""
    lines = []
    for el in content.iterchildren():
        if el.tag == qn('w:p'):
            lines.append(_para_text(el))
        elif el.tag == qn('w:tbl'):
            for tr in el.iter(qn('w:tr')):
                cells = [' '.join(_para_text(p) for p in tc.iter(qn('w:p'))).strip() for tc in tr.findall(qn('w:tc'))]
                lines.append('| ' + ' | '.join(cells) + ' |')
        elif el.tag == qn('w:sdt'):                         # a control inside (form field): its paragraphs
            inner = el.find(qn('w:sdtContent'))
            if inner is not None:
                lines.extend(_lines(inner))
    return lines


def clause_texts(data):
    """{block id: [lines]} from a Word file's clause controls."""
    doc = Document(io.BytesIO(data))
    out = {}
    for sdt in doc.element.body.iter(qn('w:sdt')):
        tag = sdt.find(f'{qn("w:sdtPr")}/{qn("w:tag")}')
        val = tag.get(qn('w:val')) if tag is not None else ''
        if not (val or '').startswith(TAG) or not val[len(TAG):].isdigit():
            continue
        content = sdt.find(qn('w:sdtContent'))
        if content is not None:
            out[int(val[len(TAG):])] = _lines(content)
    return out


def _norm(text):
    return re.sub(r'[^a-z0-9]+', ' ', (text or '').lower()).strip()


def _clause_body(heading, lines):
    """A clause's body lines from its control's lines: the heading (the first paragraph) is kept apart."""
    lines = [x.rstrip() for x in lines]
    first = next((n for n, x in enumerate(lines) if x.strip()), None)
    if heading and first is not None and _norm(lines[first]) == _norm(heading):
        lines = lines[first + 1:]
    return [x for x in lines if x.strip()]


def sync_blocks(session, data):
    """Copy the clauses' text from the saved file into the draft's blocks, so search, Lisa, compare and
    review read what was typed in the editor. A clause whose marker is gone keeps its last text."""
    texts = clause_texts(data)
    for block in session.blocks.all():
        lines = texts.get(block.id)
        if lines is None:
            continue
        body = _clause_body(block.heading, lines)
        text = '\n'.join(body)
        if text == (block.text or '').strip():
            continue
        block.text = text
        block.content_html = ''.join(f'<p>{escape(x, quote=False)}</p>' for x in body) or '<p></p>'
        block.is_edited = True
        block.save(update_fields=['text', 'content_html', 'is_edited', 'updated_at'])


def file_blocks(session, data):
    """The draft's clauses as they read in one version of its Word file, in the shape of a version snapshot
    (drafting/versions.snapshot_blocks), for the redline / compare engine. A clause missing from that
    file (added later, or its marker removed) is left out, so it shows as added / removed."""
    texts = clause_texts(data)
    out = []
    for b in session.blocks.order_by('position'):
        lines = texts.get(b.id)
        if lines is None:
            continue
        body = _clause_body(b.heading, lines)
        out.append({'block_id': b.id, 'position': b.position, 'block_type': b.block_type,
                    'heading': b.heading or '', 'text': '\n'.join(body),
                    'content_html': ''.join(f'<p>{escape(x, quote=False)}</p>' for x in body) or '<p></p>',
                    'style_json': b.style_json or {}})
    return out


# --- review state (phase 3) -----------------------------------------------------------------------

_CHANGE_TAGS = ('w:ins', 'w:del', 'w:moveFrom', 'w:moveTo', 'w:rPrChange', 'w:pPrChange')


def review_state(data):
    """{'changes': tracked changes not yet accepted / rejected, 'comments': open comment threads}
    in a Word file. A thread is open until its top comment is marked done (resolved)."""
    import zipfile
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = set(zf.namelist())
    body = etree.fromstring(zf.read('word/document.xml'))
    changes = sum(len(body.findall('.//' + qn(t))) for t in _CHANGE_TAGS)
    comments = 0
    if 'word/comments.xml' in names:
        all_comments = etree.fromstring(zf.read('word/comments.xml')).findall(qn('w:comment'))
        if 'word/commentsExtended.xml' in names:
            ext = etree.fromstring(zf.read('word/commentsExtended.xml'))
            threads = [e for e in ext.iter(f'{{{W15}}}commentEx') if not e.get(f'{{{W15}}}paraIdParent')]
            comments = sum(1 for e in threads if e.get(f'{{{W15}}}done') not in ('1', 'true'))
        else:
            comments = len(all_comments)                   # no resolved state recorded: all open
    return {'changes': changes, 'comments': comments}


# --- tracked changes written on the server (Lisa's Apply on the Community edition) ----------------

AI_AUTHOR = 'PactPro AI'
_MAX_CONTEXT_WORDS = 12
_MERGE_GAP = 40


def _run_text(r):
    return ''.join((n.text or '') if n.tag == qn('w:t') else '\t' for n in r
                   if n.tag in (qn('w:t'), qn('w:tab')))


def _text_runs(p):
    """The paragraph's runs that hold text (not deleted text), in order."""
    return [r for r in p.iter(qn('w:r')) if r.getparent().tag != qn('w:del') and _run_text(r)]


def _doc_paragraphs(doc):
    return [p for p in doc.element.body.iter(qn('w:p')) if not any(a.tag == qn('w:p') for a in p.iterancestors())]


def edits_for(before, after, document_text):
    """The smallest find / replace pairs turning `before` into `after`, each widened word by word until it
    occurs exactly once in `document_text`. None when a change can't be pinned to one place.
    (The same rule as the page's utils/textEdits.ts.)"""
    import difflib
    if before == after:
        return []
    a, b = re.split(r'(\s+)', before), re.split(r'(\s+)', after)
    starts, pos = [], 0
    for tok in a:
        starts.append(pos)
        pos += len(tok)
    starts.append(pos)
    hunks = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op != 'equal':
            hunks.append([starts[i1], starts[i2], ''.join(b[j1:j2])])
    merged = []
    for h in hunks:
        if merged and h[0] - merged[-1][1] < _MERGE_GAP:
            last = merged[-1]
            last[2] += before[last[1]:h[0]] + h[2]
            last[1] = h[1]
        else:
            merged.append(h)

    def widen(s, e):
        while s > 0 and before[s - 1].isspace():
            s -= 1
        while s > 0 and not before[s - 1].isspace():
            s -= 1
        while e < len(before) and before[e].isspace():
            e += 1
        while e < len(before) and not before[e].isspace():
            e += 1
        return s, e

    out = []
    for hs, he, new in merged:
        s, e = (hs, he) if hs != he else widen(hs, he)
        for _ in range(_MAX_CONTEXT_WORDS):
            if document_text.count(before[s:e]) == 1:
                break
            ns, ne = widen(s, e)
            if (ns, ne) == (s, e):
                break
            s, e = ns, ne
        find = before[s:e]
        if not find.strip() or document_text.count(find) != 1:
            return None
        out.append((find, before[s:hs] + new + before[he:e]))
    return out


def _split(r, at):
    """Split run `r` at character `at` of its text; returns (left, right), both with r's formatting."""
    text = _run_text(r)
    rpr = r.find(qn('w:rPr'))
    right = OxmlElement('w:r')
    if rpr is not None:
        right.append(copy.deepcopy(rpr))
    for child in [c for c in r if c.tag != qn('w:rPr')]:
        r.remove(child)
    _put_text(r, text[:at])
    _put_text(right, text[at:])
    r.addnext(right)
    return r, right


def _put_text(r, text):
    for k, piece in enumerate(text.split('\t')):
        if k:
            r.append(OxmlElement('w:tab'))
        if piece:
            t = OxmlElement('w:t')
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            t.text = piece
            r.append(t)


def _tracked(tag, author, when):
    el = OxmlElement(tag)
    el.set(qn('w:id'), str(secrets.randbelow(10 ** 8) + 10 ** 6))
    el.set(qn('w:author'), author)
    el.set(qn('w:date'), when)
    return el


def _track_in_paragraph(p, start, end, insert, author, when):
    """Mark [start, end) of the paragraph's text deleted and `insert` inserted there, as tracked changes."""
    runs, pos = _text_runs(p), 0
    # Split the runs so the deleted span starts and ends on run edges.
    for edge in (start, end):
        pos = 0
        for r in _text_runs(p):
            n = len(_run_text(r))
            if pos < edge < pos + n:
                _split(r, edge - pos)
                break
            pos += n
    runs, pos, inside, anchor, fmt = _text_runs(p), 0, [], None, None
    for r in runs:
        n = len(_run_text(r))
        if start <= pos and pos + n <= end and n:
            inside.append(r)
        elif pos + n <= start:
            anchor = r
        pos += n
    if inside:
        fmt = inside[0].find(qn('w:rPr'))
        dele = _tracked('w:del', author, when)
        inside[0].addprevious(dele)
        for r in inside:
            for t in r.findall(qn('w:t')):
                t.tag = qn('w:delText')
            dele.append(r)
        anchor = dele
    elif anchor is not None:
        fmt = anchor.find(qn('w:rPr'))
    if insert:
        ins = _tracked('w:ins', author, when)
        new = OxmlElement('w:r')
        if fmt is not None:
            new.append(copy.deepcopy(fmt))
        _put_text(new, insert)
        ins.append(new)
        if anchor is not None:
            anchor.addnext(ins)
        else:
            first = p.find(qn('w:pPr'))
            (first.addnext(ins) if first is not None else p.insert(0, ins))


def _clause_paragraphs(doc, block_id):
    """The paragraphs inside clause `block_id`'s marker, or [] when the file has no such marker."""
    for sdt in doc.element.body.iter(qn('w:sdt')):
        tag = sdt.find(f'{qn("w:sdtPr")}/{qn("w:tag")}')
        if tag is not None and tag.get(qn('w:val')) == f'{TAG}{block_id}':
            content = sdt.find(qn('w:sdtContent'))
            return [p for p in content.iter(qn('w:p'))] if content is not None else []
    return []


def _onto_lines(b_lines, after):
    """`after` cut into as many lines as `b_lines`, each new word going to the line of the old word it
    matches or follows. The model often joins two lines with a space (or splits one) while changing a
    word; a tracked change can't merge or split paragraphs, so the words go back onto the file's lines.
    None when `b_lines` has no words."""
    import difflib
    old = [(w, n) for n, line in enumerate(b_lines) for w in line.split()]
    if not old:
        return None
    new = after.split()
    out = [[] for _ in b_lines]
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=[w for w, _ in old], b=new, autojunk=False).get_opcodes():
        if tag == 'equal':
            for k in range(j2 - j1):
                out[old[i1 + k][1]].append(new[j1 + k])
        elif tag in ('replace', 'insert'):
            line = old[i1][1] if i1 < len(old) else old[-1][1]
            out[line].extend(new[j1:j2])
    # Lines whose words didn't change keep their own spacing.
    return [b if b.split() == words else ' '.join(words) for b, words in zip(b_lines, out)]


def apply_tracked(data, before, after, author=AI_AUTHOR, block_id=None):
    """The Word file with the change from `before` to `after` (a clause's text) written as tracked
    changes by `author`, for the person who decides to accept or reject in the editor.
    Returns (bytes, True), or (data, False) when a change can't be pinned to exactly one place.

    Lisa reads a clause as its block text: one line per Word paragraph, blanks as [[Label]]. The file
    shows a blank as its form field's text (the value, or the label as a grey hint), so blanks are put
    back the way the file shows them, and the change is placed line by line: a clause spans several
    paragraphs, a tracked change can't."""
    import difflib
    from datetime import datetime, timezone
    shown = {label: f['value'] or label for label, f in fields(data).items()}
    as_in_file = lambda s: PLACEHOLDER.sub(lambda m: shown.get(m.group(1).strip(), m.group(0)), s)
    before, after = as_in_file(before), as_in_file(after)
    doc = Document(io.BytesIO(data))
    paras = _doc_paragraphs(doc)
    if block_id is not None:
        # Only inside that clause's marker: the same words elsewhere (a repeated signature line) don't count.
        own = _clause_paragraphs(doc, block_id)
        paras = own or paras
    texts = [''.join(_run_text(r) for r in _text_runs(p)) for p in paras]
    b_lines, a_lines = before.split('\n'), after.split('\n')
    pairs = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=b_lines, b=a_lines, autojunk=False).get_opcodes():
        if tag == 'equal':
            continue
        if tag == 'replace' and i2 - i1 == j2 - j1:       # reworded lines: each stays in its paragraph
            pairs += list(zip(b_lines[i1:i2], a_lines[j1:j2]))
        elif i2 > i1 and j2 > j1 and (lines := _onto_lines(b_lines[i1:i2], ' '.join(a_lines[j1:j2]))):
            # Lines joined or split by the model: its words back onto the file's lines, one change each.
            pairs += [(b, a) for b, a in zip(b_lines[i1:i2], lines) if b != a]
        else:                                              # lines added / removed: only if it fits one paragraph
            pairs.append(('\n'.join(b_lines[i1:i2]), '\n'.join(a_lines[j1:j2])))
    if not pairs:
        return data, False
    when = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    for old, new in pairs:
        edits = edits_for(old, new, '\n'.join(texts))
        if not edits:
            return data, False
        for find, replace in edits:
            hit = [(n, t.find(find)) for n, t in enumerate(texts) if find in t]
            if len(hit) != 1:                              # spans paragraphs, or no longer unique
                return data, False
            n, at = hit[0]
            # Only the part that differs is marked: the context around it stays as it is.
            pre = 0
            while pre < min(len(find), len(replace)) and find[pre] == replace[pre]:
                pre += 1
            suf = 0
            while suf < min(len(find), len(replace)) - pre and find[-1 - suf] == replace[-1 - suf]:
                suf += 1
            _track_in_paragraph(paras[n], at + pre, at + len(find) - suf, replace[pre:len(replace) - suf], author, when)
            texts[n] = ''.join(_run_text(r) for r in _text_runs(paras[n]))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), True


# --- Lisa's edits as find / replace on numbered paragraphs (services/doc_edits.py) ----------------
# The model never rewrites text: it names a paragraph and the exact words to replace there, and every
# edit is checked against the file before anyone sees it and again when it is written.

def _block_of(p):
    """The clause id of the marker a paragraph sits in, or None (title, letterhead, signatures)."""
    for anc in p.iterancestors(qn('w:sdt')):
        tag = anc.find(f'{qn("w:sdtPr")}/{qn("w:tag")}')
        val = (tag.get(qn('w:val')) if tag is not None else '') or ''
        if val.startswith(TAG) and val[len(TAG):].isdigit():
            return int(val[len(TAG):])
    return None


def paragraph_listing(data):
    """Every paragraph of the file in order: [{n, text, block_id}]. `text` is what a reader sees, with an
    empty form field as its [[Label]] blank; `n` is the paragraph's index, the edits' address."""
    doc = Document(io.BytesIO(data))
    return [{'n': n, 'text': _para_text(p), 'block_id': _block_of(p)} for n, p in enumerate(_doc_paragraphs(doc))]


def _in_file(data):
    """Turns text as the listing shows it into the paragraph's run text: an empty field reads as its label
    (the grey hint), a filled one as its value."""
    shown = {label: f['value'] or label for label, f in fields(data).items()}
    return lambda s: PLACEHOLDER.sub(lambda m: shown.get(m.group(1).strip(), m.group(0)), s)


def apply_edits(data, edits, author=AI_AUTHOR):
    """Write [{id, n, text, find, replace}] into the file as tracked changes by `author`. Each edit is found
    by its paragraph: index `n` when that paragraph still reads `text`, else the one paragraph that does,
    else the one paragraph containing `find`; every occurrence of `find` there is replaced. Returns
    (bytes, applied ids, skipped ids); an edit that can't be pinned down is skipped, never guessed."""
    from datetime import datetime, timezone
    doc = Document(io.BytesIO(data))
    paras = _doc_paragraphs(doc)
    shown = [_para_text(p) for p in paras]                # as the edits saw them (kept while editing)
    in_file = _in_file(data)
    when = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    applied, skipped = [], []
    for e in edits:
        find, replace = in_file(str(e.get('find') or '')), in_file(str(e.get('replace') or ''))
        if not find or find == replace or '\n' in find:
            skipped.append(e.get('id'))
            continue
        n = e.get('n')
        if not (isinstance(n, int) and 0 <= n < len(paras) and shown[n] == e.get('text')):
            same = [i for i, t in enumerate(shown) if t == e.get('text')]
            if len(same) == 1:
                n = same[0]
            else:
                holding = [i for i, p in enumerate(paras) if find in ''.join(_run_text(r) for r in _text_runs(p))]
                n = holding[0] if len(holding) == 1 else None
        if n is None:
            skipped.append(e.get('id'))
            continue
        text = ''.join(_run_text(r) for r in _text_runs(paras[n]))
        spots = [m.start() for m in re.finditer(re.escape(find), text)]
        if not spots:
            skipped.append(e.get('id'))
            continue
        pre = 0
        while pre < min(len(find), len(replace)) and find[pre] == replace[pre]:
            pre += 1
        suf = 0
        while suf < min(len(find), len(replace)) - pre and find[-1 - suf] == replace[-1 - suf]:
            suf += 1
        for at in reversed(spots):                         # right to left: earlier offsets stay valid
            _track_in_paragraph(paras[n], at + pre, at + len(find) - suf, replace[pre:len(replace) - suf], author, when)
        applied.append(e.get('id'))
    if not applied:
        return data, applied, skipped
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), applied, skipped


# --- DRAFT watermark on copies of a draft that isn't final --------------------------------------

_WATERMARK = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
    '<w:pPr><w:pStyle w:val="Header"/></w:pPr><w:r><w:pict>'
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e"><v:path textpathok="t" o:connecttype="custom"/>'
    '<v:textpath on="t" fitshape="t"/><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>'
    '<v:shape id="PactProWatermark" o:spid="_x0000_s2049" type="#_x0000_t136" '
    'style="position:absolute;margin-left:0;margin-top:0;width:470pt;height:120pt;rotation:315;'
    'z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
    'mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" '
    'fillcolor="#c8c8c8" stroked="f"><v:fill opacity=".45"/>'
    '<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{text}"/>'
    '</v:shape></w:pict></w:r></w:p>')


def add_watermark(data, text='DRAFT'):
    """The Word file with a diagonal watermark on every page (in each section's header), the way Word
    draws one: on copies of a draft that isn't final yet, so a draft is never mistaken for the final."""
    from html import escape as _esc
    doc = Document(io.BytesIO(data))
    seen = set()
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header.is_linked_to_previous and section is not doc.sections[0]:
                continue
            el = header._element
            if id(el) in seen:
                continue
            seen.add(id(el))
            el.append(etree.fromstring(_WATERMARK.replace('{text}', _esc(text))))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- placeholders ---------------------------------------------------------------------------------

def _paragraphs(doc):
    for p in doc.element.body.iter(qn('w:p')):
        yield p
    for part in {id(s.header.part): s.header for s in doc.sections}.values():
        yield from part._element.iter(qn('w:p'))


def placeholders(data):
    """[{label, count}] of the [[Label]] text blanks still in the file, in document order (form fields,
    which read as [[Label]] while empty, are listed by `fields`)."""
    doc = Document(io.BytesIO(data))
    seen = {}
    for p in _paragraphs(doc):
        plain = ''.join(t.text or '' for t in p.iter(qn('w:t'))
                        if not any(a.tag == qn('w:sdt') and _field_label(a) is not None for a in t.iterancestors()))
        for m in PLACEHOLDER.finditer(plain):
            label = m.group(1).strip()
            seen[label] = seen.get(label, 0) + 1
    return [{'label': k, 'count': v} for k, v in seen.items()]


def _replace_in_paragraph(p, token, value):
    """Replace `token` in a paragraph whose text may be split across runs: the value goes into the run
    where the token starts (keeping that run's formatting), the rest of the token is cut from the runs
    after it. Returns how many were replaced."""
    count = 0
    while True:
        ts = [t for t in p.iter(qn('w:t'))]
        full = ''.join(t.text or '' for t in ts)
        at = full.find(token)
        if at < 0:
            return count
        end, pos = at + len(token), 0
        placed = False
        for t in ts:
            text = t.text or ''
            a, b = pos, pos + len(text)
            pos = b
            if b <= at or a >= end:
                continue
            lo, hi = max(at, a) - a, min(end, b) - a
            t.text = text[:lo] + ('' if placed else value) + text[hi:]
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            placed = True
        count += 1


def fill_placeholder(data, label, value):
    """The file with every [[label]] replaced by `value`. Returns (bytes, how many)."""
    doc = Document(io.BytesIO(data))
    token, n = f'[[{label}]]', 0
    for p in _paragraphs(doc):
        if token in _para_text(p):
            n += _replace_in_paragraph(p, token, value)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), n
