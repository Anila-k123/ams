"""Word form: a draft that looks like the Word template / sample it came from (docs/DRAFT_EXPORT.md,
"Template look").

The AI still writes the text. What this module adds:

1. Reading (`clause_forms`): where each clause of a .docx sits (its paragraphs and tables, in body
   order) and how each of those paragraphs is formatted (font, size, colour, alignment, indents,
   spacing, emphasis), plus the page size, margins, header and footer (`page_form`).
2. Laying text into it (`styled_lines_html`): each generated line takes the formatting of the
   template paragraph it matches (or the one at its position), so the editor opens in the
   template's look.
3. Writing (`render_into_original`): Download / Save to case open the ORIGINAL file and write each
   clause's final text into that clause's own paragraphs, reusing their formatting; everything
   outside the clauses (title, logo, header, footer, signature images, page set-up) is untouched.

A clause's place is kept on its block as style_json['src'] = {'file': 'template:<id>' | 'sample:<id>',
'range': [first, last] element index, 'heading': original heading text or ''}.
"""

import base64
import copy
import io
import re
from functools import lru_cache

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

PREFIX = 32
_ALIGN = {0: 'left', 1: 'center', 2: 'right', 3: 'justify'}
_JC = {'left': 'left', 'center': 'center', 'right': 'right', 'justify': 'both'}
_NUM_LEAD = re.compile(r'^\s*(?:\(?[0-9]+(?:\.[0-9]+)*[.)]?|\(?[a-z]\)|\(?[ivxlc]+\))\s+', re.I)


def _norm(text):
    return re.sub(r'[^a-z0-9]+', ' ', (text or '').lower()).strip()


# --- reading --------------------------------------------------------------------------

def _chain(style):
    while style is not None:
        yield style
        style = getattr(style, 'base_style', None)


def _pf(para, attr):
    """A paragraph-format value as Word resolves it: the paragraph's own, else its style chain's."""
    val = getattr(para.paragraph_format, attr)
    if val is not None:
        return val
    for st in _chain(para.style):
        fmt = getattr(st, 'paragraph_format', None)
        if fmt is not None and getattr(fmt, attr) is not None:
            return getattr(fmt, attr)
    return None


def _font(para, attr):
    """A font value of the paragraph's first visible run, else its style chain's."""
    runs = [r for r in para.runs if (r.text or '').strip()]
    if runs:
        try:
            val = getattr(runs[0].font, attr)
            if val is not None:
                return val
        except Exception:                                         # noqa: BLE001
            pass
    for st in _chain(para.style):
        font = getattr(st, 'font', None)
        try:
            if font is not None and getattr(font, attr) is not None:
                return getattr(font, attr)
        except Exception:                                         # noqa: BLE001
            continue
    return None


def _pt(length):
    return round(length.pt, 1) if length is not None else None


def para_format(para):
    """The paragraph's look as a small dict (keys the editor CSS and the heading style read)."""
    fmt = {}
    align = _pf(para, 'alignment')
    if align is not None:
        fmt['align'] = _ALIGN.get(int(align), 'left')
    name = _font(para, 'name')
    if name:
        fmt['font'] = name
    size = _font(para, 'size')
    if size is not None:
        fmt['size'] = round(size.pt, 1)
    try:
        color = _font(para, 'color')
        if color is not None and color.type is not None and color.rgb is not None and str(color.rgb) != '000000':
            fmt['color'] = f'#{color.rgb}'
    except Exception:                                             # noqa: BLE001
        pass
    runs = [r for r in para.runs if (r.text or '').strip()]
    if runs and all(r.bold or (r.bold is None and _font(para, 'bold')) for r in runs):
        fmt['bold'] = True
    if runs and all(r.italic for r in runs):
        fmt['italic'] = True
    if runs and all(r.underline for r in runs):
        fmt['underline'] = True
    for key, attr in (('indent_left', 'left_indent'), ('indent_first', 'first_line_indent'),
                      ('space_before', 'space_before'), ('space_after', 'space_after')):
        val = _pt(_pf(para, attr))
        if val:
            fmt[key] = val
    spacing = _pf(para, 'line_spacing')
    if isinstance(spacing, float) and spacing and abs(spacing - 1.0) > 0.01:
        fmt['line_height'] = round(spacing, 2)
    if para.paragraph_format.element.pPr is not None and para.paragraph_format.element.pPr.numPr is not None:
        fmt['numbered'] = True
    return fmt


def _body_elements(doc):
    """The body's paragraphs and tables in order (section properties left out)."""
    return [el for el in doc.element.body.iterchildren() if el.tag in (qn('w:p'), qn('w:tbl'))]


def _el_text(el, doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    if el.tag == qn('w:p'):
        return Paragraph(el, doc).text or ''
    return '\n'.join(' — '.join(c.text.strip() for c in row.cells if c.text.strip()) for row in Table(el, doc).rows)


def clause_forms(path, clause_texts, headings=None):
    """For each clause (its source text, in document order) of the .docx at `path`:
    {'range': [first, last], 'heading': heading text or '', 'paras': [{'i', 'key', 'fmt'}...]} or None
    when the clause can't be found. Element indices count the body's paragraphs and tables."""
    from docx.text.paragraph import Paragraph
    try:
        doc = Document(path)
    except Exception:                                             # noqa: BLE001
        return [None] * len(clause_texts)
    els = _body_elements(doc)
    keys = [_norm(_el_text(el, doc).split('\n')[0]) for el in els]
    starts, cursor = [], 0
    for text in clause_texts:
        first = next((_norm(x) for x in (text or '').split('\n') if x.strip()), '')
        hit = next((i for i in range(cursor, len(els)) if first and keys[i] and keys[i][:PREFIX] == first[:PREFIX]), None)
        starts.append(hit)
        if hit is not None:
            cursor = hit + 1
    out = []
    for n, start in enumerate(starts):
        if start is None:
            out.append(None)
            continue
        nxt = next((s for s in starts[n + 1:] if s is not None), len(els))
        end = nxt - 1
        while end > start and not _el_text(els[end], doc).strip():
            end -= 1                                              # trailing empty paragraphs belong to nobody
        paras = []
        for i in range(start, end + 1):
            el = els[i]
            if el.tag == qn('w:tbl'):
                paras.append({'i': i, 'table': True})
                continue
            p = Paragraph(el, doc)
            if not (p.text or '').strip():
                continue
            paras.append({'i': i, 'key': _norm(p.text)[:PREFIX], 'fmt': para_format(p), 'text': p.text.strip(),
                          'html': runs_html(p)})
        shapes = sum(_drawings(els[i]) for i in range(start, end + 1))
        # The clause's heading is the one the parser detected (services/headings.py), when it is the
        # clause's first paragraph here; never a guess from a short first line.
        detected = ((headings or [None] * len(clause_texts))[n] or '').strip()
        is_heading = bool(detected) and bool(paras) and not paras[0].get('table') \
            and paras[0].get('key') == _norm(detected)[:PREFIX]
        out.append({'range': [start, end], 'heading': detected if is_heading else '', 'paras': paras,
                    'shapes': shapes})
    return out


def _drawings(el):
    """Drawn shapes / pictures in a body element (stamp boxes, lines, images)."""
    # A shape is often stored twice (a drawing plus its older VML fallback): count one of them.
    return len(el.findall('.//' + qn('w:drawing'))) or len(el.findall('.//' + qn('w:pict')))


_TAB = '  '                                             # a tab stop, as the editor can show it


def same_line(a, b):
    """Two lines with the same words (tabs / spacing don't count)."""
    return ' '.join((a or '').split()) == ' '.join((b or '').split())


def is_layout_line(text):
    """A line with no letters or digits ("))", "-----"): layout from tab stops, not content."""
    return bool((text or '').strip()) and not re.search(r'[^\W_]', text)


def drop_layout_lines(text):
    return '\n'.join(x for x in (text or '').split('\n') if not is_layout_line(x))


def runs_html(para):
    """The paragraph's text as HTML keeping its runs' own bold / italic / underline (a few bold words
    inside a line) and its tabs, for lines the draft keeps word for word."""
    from html import escape
    out = []
    for r in para.runs:
        bits = []
        for node in r._element.iterchildren():
            if node.tag == qn('w:t'):
                bits.append(escape(node.text or '', quote=False))
            elif node.tag == qn('w:tab'):
                bits.append(_TAB)
            elif node.tag in (qn('w:br'), qn('w:cr')):
                bits.append('<br>')
        text = ''.join(bits)
        if not text:
            continue
        if r.bold:
            text = f'<strong>{text}</strong>'
        if r.italic:
            text = f'<em>{text}</em>'
        if r.underline:
            text = f'<u>{text}</u>'
        out.append(text)
    return ''.join(out)


def _emu_mm(v):
    return round(Emu(v).mm, 1) if v is not None else None


def _part_html(part_el, part):
    """A header / footer as simple HTML: paragraphs (alignment, bold, size), tables, and images as
    data: URIs. Shown above / below the draft in the editor; never edited there."""
    from html import escape
    out = []

    def para_html(p):
        jc = p.find(f'{qn("w:pPr")}/{qn("w:jc")}')
        align = {'center': 'center', 'right': 'right', 'end': 'right', 'both': 'justify'}.get(
            jc.get(qn('w:val')) if jc is not None else '', '')
        bits = []
        for node in p.iter():
            if node.tag == qn('w:t'):
                bits.append(escape(node.text or ''))
            elif node.tag == qn('w:br'):
                bits.append('<br>')
            elif node.tag == qn('a:blip'):
                rid = node.get(qn('r:embed'))
                try:
                    img = part.related_parts[rid]
                    data = base64.b64encode(img.blob).decode()
                    bits.append(f'<img src="data:{img.content_type};base64,{data}" alt="">')
                except Exception:                                 # noqa: BLE001
                    pass
        style = f' style="text-align:{align}"' if align else ''
        return f'<p{style}>{"".join(bits) or "&nbsp;"}</p>'

    for el in part_el.iterchildren():
        if el.tag == qn('w:p'):
            out.append(para_html(el))
        elif el.tag == qn('w:tbl'):
            rows = []
            for tr in el.iter(qn('w:tr')):
                cells = ''.join('<td>' + ''.join(para_html(p) for p in tc.iter(qn('w:p'))) + '</td>'
                                for tc in tr.findall(qn('w:tc')))
                rows.append(f'<tr>{cells}</tr>')
            out.append(f'<table>{"".join(rows)}</table>')
    html = ''.join(out)
    return html if re.search(r'<img|[^\s<>&;]', re.sub(r'<[^>]+>|&nbsp;', '', html)) or '<img' in html else ''


@lru_cache(maxsize=64)
def _page_form_cached(path, _mtime):
    try:
        doc = Document(path)
    except Exception:                                             # noqa: BLE001
        return None
    sec = doc.sections[0]
    header = _part_html(sec.header._element, sec.header.part)
    footer = _part_html(sec.footer._element, sec.footer.part)
    return {
        'page': {'width_mm': _emu_mm(sec.page_width), 'margin_left_mm': _emu_mm(sec.left_margin),
                 'margin_right_mm': _emu_mm(sec.right_margin), 'margin_top_mm': _emu_mm(sec.top_margin),
                 'margin_bottom_mm': _emu_mm(sec.bottom_margin),
                 # Where Word prints the header / footer: inside the top / bottom margin, this far from
                 # the page edge (the editor draws them there, not after the margin).
                 'header_mm': _emu_mm(sec.header_distance), 'footer_mm': _emu_mm(sec.footer_distance)},
        'header_html': header, 'footer_html': footer,
        'has_letterhead': bool(header),
    }


def page_form(path):
    """Page size / margins, header and footer HTML, and whether the file has its own letterhead."""
    import os
    try:
        return _page_form_cached(str(path), os.path.getmtime(path))
    except OSError:
        return None


# --- which file a draft follows ---------------------------------------------------------

def _word_path(obj):
    """The Word file of a template / sample: its converted Word version, else its own .docx."""
    for f in (getattr(obj, 'word_file', None), getattr(obj, 'file', None)):
        try:
            if f and f.name and f.name.lower().endswith('.docx') and f.path:
                import os
                if os.path.exists(f.path):
                    return f.path
        except Exception:                                         # noqa: BLE001
            continue
    return None


def source_for(ref):
    """'template:<id>' / 'sample:<id>' -> the Word file path, or None."""
    from .models import Sample, Template
    kind, _, pk = (ref or '').partition(':')
    model = {'template': Template, 'sample': Sample}.get(kind)
    obj = model.objects.filter(id=pk).first() if model and pk.isdigit() else None
    return _word_path(obj) if obj else None


def session_source(session):
    """The Word file this draft follows (from its blocks' src), or None."""
    for b in session.blocks.all():
        ref = ((b.style_json or {}).get('src') or {}).get('file')
        if ref:
            return ref, source_for(ref)
    return None, None


# --- laying generated text into the form -------------------------------------------------

def fmt_css(fmt):
    """A paragraph format dict as inline CSS for the editor."""
    if not fmt:
        return ''
    css = []
    if fmt.get('align') and fmt['align'] != 'left':
        css.append(f"text-align:{fmt['align']}")
    if fmt.get('font'):
        css.append(f"font-family:'{fmt['font']}'")
    if fmt.get('size'):
        css.append(f"font-size:{fmt['size']}pt")
    if fmt.get('color'):
        css.append(f"color:{fmt['color']}")
    if fmt.get('bold'):
        css.append('font-weight:bold')
    if fmt.get('italic'):
        css.append('font-style:italic')
    if fmt.get('underline'):
        css.append('text-decoration:underline')
    if fmt.get('indent_left'):
        css.append(f"margin-left:{fmt['indent_left']}pt")
    if fmt.get('indent_first'):
        css.append(f"text-indent:{fmt['indent_first']}pt")
    if fmt.get('space_before'):
        css.append(f"margin-top:{fmt['space_before']}pt")
    if fmt.get('space_after'):
        css.append(f"margin-bottom:{fmt['space_after']}pt")
    if fmt.get('line_height'):
        css.append(f"line-height:{fmt['line_height']}")
    return ';'.join(css)


def body_paras(form):
    """The clause's paragraph formats without its heading line (the editor shows that apart)."""
    paras = [p for p in form.get('paras', []) if not p.get('table')]
    if form.get('heading') and paras and paras[0].get('key', '') == _norm(form['heading'])[:PREFIX]:
        paras = paras[1:]
    return paras


def pick_format(line, paras, index, every=None):
    """The template paragraph a generated line takes its look from: the one it starts like (among
    `every` paragraph of the clause, heading included), else the body paragraph at its position
    (the last one for extra lines)."""
    key = _norm(line)[:PREFIX]
    for p in (every if every is not None else paras):
        if p.get('table'):
            continue
        k = p.get('key', '')
        n = min(len(k), len(key), 20)
        if k and (k == key or (n >= 8 and k[:n] == key[:n])):
            return p['fmt']
    if not paras:
        return {}
    return paras[min(index, len(paras) - 1)]['fmt']


def drop_leading_heading(text, heading):
    """The text without a first line that only repeats `heading` (the clause shows it apart)."""
    lines = (text or '').split('\n')
    first = next((n for n, line in enumerate(lines) if line.strip()), None)
    if first is not None and _norm(heading) and _norm(lines[first]) == _norm(heading):
        lines = lines[first + 1:]
    return '\n'.join(lines)


def styled_lines_html(text, form, table_html):
    """content_html for a generated clause in the form's look. `table_html(rows)` turns a run of
    `| a | b |` lines into a table (tasks._table_html)."""
    from html import escape
    paras, every = body_paras(form), form.get('paras', [])
    lines = (text or '').split('\n')
    out, i, n = [], 0, 0
    row_re = re.compile(r'^\s*\|.*\|\s*$')
    while i < len(lines):
        rows = []
        while i < len(lines) and row_re.match(lines[i]):
            rows.append(lines[i])
            i += 1
        if rows:
            out.append(table_html(rows))
            continue
        line = lines[i]
        i += 1
        if not line.strip():
            continue
        css = fmt_css(pick_format(line, paras, n, every))
        n += 1
        # A line kept word for word shows the template's own runs (bold words, tabs).
        same = next((p for p in every if not p.get('table') and p.get('html') and same_line(p.get('text'), line)), None)
        inner = same['html'] if same else escape(line, quote=False).replace('\t', _TAB)
        out.append(f'<p style="{css}">{inner}</p>' if css else f'<p>{inner}</p>')
    return ''.join(out) or '<p></p>'


# --- writing into the original ----------------------------------------------------------

def _clear_runs(p_el):
    for child in list(p_el):
        if child.tag != qn('w:pPr'):
            p_el.remove(child)


def _first_rpr(p_el):
    r = p_el.find(qn('w:r'))
    rpr = r.find(qn('w:rPr')) if r is not None else None
    return copy.deepcopy(rpr) if rpr is not None else None


def _set_jc(p_el, align):
    if not align:
        return
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        from docx.oxml import OxmlElement
        ppr = OxmlElement('w:pPr')
        p_el.insert(0, ppr)
    jc = ppr.find(qn('w:jc'))
    if jc is None:
        from docx.oxml import OxmlElement
        jc = OxmlElement('w:jc')
        ppr.append(jc)
    jc.set(qn('w:val'), _JC.get(align, 'left'))


def _write_runs(p_el, runs, rpr):
    """Put (text, formats) runs into a cleared paragraph, each run with the template's first-run
    properties plus the editor's bold / italic / underline / strike."""
    from docx.oxml import OxmlElement
    for text, fmt in runs:
        r = OxmlElement('w:r')
        props = copy.deepcopy(rpr) if rpr is not None else OxmlElement('w:rPr')
        for name, tag in (('bold', 'w:b'), ('italic', 'w:i'), ('strike', 'w:strike')):
            if fmt and name in fmt and props.find(qn(tag)) is None:
                props.append(OxmlElement(tag))
        if fmt and 'underline' in fmt and props.find(qn('w:u')) is None:
            u = OxmlElement('w:u')
            u.set(qn('w:val'), 'single')
            props.append(u)
        if len(props):
            r.append(props)
        if text == '\n':
            r.append(OxmlElement('w:br'))
        else:
            # Tabs (typed, or a template tab shown as wide spaces) print as Word tab stops again.
            for k, piece in enumerate(re.split(r'\t|' + _TAB, text)):
                if k:
                    r.append(OxmlElement('w:tab'))
                if piece:
                    t = OxmlElement('w:t')
                    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    t.text = piece
                    r.append(t)
        p_el.append(r)


def _para_from_template(tpl_el, para):
    """A new paragraph: the template paragraph's properties with the draft paragraph's text."""
    new = copy.deepcopy(tpl_el)
    rpr = _first_rpr(tpl_el)
    _clear_runs(new)
    runs = list(para.runs)
    numbered = new.find(f'{qn("w:pPr")}/{qn("w:numPr")}') is not None
    if numbered and runs and runs[0][0] != '\n':                 # Word numbers it: drop a typed number
        runs[0] = (_NUM_LEAD.sub('', runs[0][0], count=1), runs[0][1])
    _set_jc(new, para.align)
    _write_runs(new, runs, rpr)
    return new


def _table_from_template(tpl_tbl, table, para_tpl):
    """A table with the template table's look and the draft table's rows / cells."""
    from docx.oxml import OxmlElement
    if tpl_tbl is None:
        tbl = OxmlElement('w:tbl')
        tblpr = OxmlElement('w:tblPr')
        style = OxmlElement('w:tblStyle')
        style.set(qn('w:val'), 'TableGrid')
        tblpr.append(style)
        tbl.append(tblpr)
        tpl_rows = []
    else:
        tbl = copy.deepcopy(tpl_tbl)
        tpl_rows = tbl.findall(qn('w:tr'))
        for tr in tpl_rows:
            tbl.remove(tr)
    from .export.docx import Para
    for n, row in enumerate(table.rows):
        if tpl_rows:
            tr = copy.deepcopy(tpl_rows[min(n, len(tpl_rows) - 1)])
            tcs = tr.findall(qn('w:tc'))
        else:
            tr, tcs = OxmlElement('w:tr'), []
        for j, cell in enumerate(row):
            if tcs and j < len(tcs):
                tc = tcs[j]
            else:
                tc = copy.deepcopy(tcs[-1]) if tcs else OxmlElement('w:tc')
                tr.append(tc)
            ps = tc.findall(qn('w:p'))
            tpl_p = ps[0] if ps else (para_tpl if para_tpl is not None else OxmlElement('w:p'))
            for p in ps:
                tc.remove(p)
            cell_paras = [x for x in cell if isinstance(x, Para)] or [Para(runs=[('', frozenset())])]
            for cp in cell_paras:
                tc.append(_para_from_template(tpl_p, cp))
        for extra in tcs[len(row):]:
            tr.remove(extra)
        tbl.append(tr)
    return tbl


def _draft_items(block, heading_text):
    """The block's body as docx Para / Table items, without a first line that repeats its heading."""
    from .export.docx import _text_to_blocks, html_to_blocks
    items = html_to_blocks(block.content_html) if (block.content_html or '').strip() else _text_to_blocks(block.text)
    first = next((x for x in items if hasattr(x, 'runs')), None)
    if first is not None and items and items[0] is first and heading_text:
        t = _norm(''.join(r[0] for r in first.runs))
        if t and (t == _norm(heading_text) or t == _norm(block.heading)):
            items = items[1:]
    return items


def render_into_original(session, path, branding=None):
    """The draft written into its Word source: bytes of a .docx."""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    from .docfile import record, wrap_blocks
    from .export.docx import Para, _branding, _signature_block
    doc = Document(path)
    els = _body_elements(doc)
    blocks = list(session.blocks.order_by('position'))
    last_el = None
    for block in blocks:
        src = (block.style_json or {}).get('src') or {}
        rng = src.get('range')
        if not rng or rng[0] >= len(els):
            # A clause with no place in the file: after the previous clause, in its body look.
            items = _draft_items(block, '')
            anchor = last_el if last_el is not None else (els[-1] if els else None)
            tpl = anchor if anchor is not None and anchor.tag == qn('w:p') else OxmlElement('w:p')
            new = []
            if (block.heading or '').strip():
                new.append(_para_from_template(tpl, Para(runs=[(block.heading, frozenset(['bold']))])))
            new += [_para_from_template(tpl, it) if isinstance(it, Para) else _table_from_template(None, it, tpl)
                    for it in items]
            for el in new:
                if anchor is not None:
                    anchor.addnext(el)
                    anchor = el
            record(getattr(block, 'id', None), new)
            last_el = anchor
            continue
        start, end = rng[0], min(rng[1], len(els) - 1)
        part = els[start:end + 1]
        heading_el = None
        original_heading = src.get('heading') or ''
        # The template's heading paragraph is kept apart only when the draft shows the heading apart
        # (block.heading). Otherwise the heading is the first line of the draft's own text and is
        # written from there like every other line, so a changed first line can't print twice.
        if not (block.heading or '').strip():
            original_heading = ''
        if original_heading and part and part[0].tag == qn('w:p'):
            heading_el = part[0]
            part = part[1:]
            # Keep the template's heading paragraph (its numbering, look) unless the heading was renamed.
            if _norm(block.heading) and _norm(block.heading) not in _norm(original_heading):
                tpl = copy.deepcopy(heading_el)
                new_heading = _para_from_template(tpl, Para(runs=[(block.heading, frozenset())]))
                heading_el.addprevious(new_heading)
                heading_el.getparent().remove(heading_el)
                heading_el = new_heading
        all_ps = [el for el in part if el.tag == qn('w:p')]
        texts = {id(el): Paragraph(el, doc).text or '' for el in all_ps}
        tpl_paras = [el for el in all_ps if texts[id(el)].strip() and not is_layout_line(texts[id(el)])]
        tpl_tables = [el for el in part if el.tag == qn('w:tbl')]
        items = [it for it in _draft_items(block, original_heading)
                 if not (isinstance(it, Para) and is_layout_line(''.join(r[0] for r in it.runs)))]
        keys = [_norm(texts[id(el)])[:PREFIX] for el in tpl_paras]
        new, came_from, reused, n, t = [], [], set(), 0, 0
        for it in items:
            if isinstance(it, Para):
                line = ''.join(r[0] for r in it.runs)
                same = next((el for el in tpl_paras if id(el) not in reused and same_line(texts[id(el)], line)), None)
                if same is not None:
                    # Kept word for word: the template's own paragraph, every run (bold words, tabs,
                    # inline shapes) as the file has it.
                    reused.add(id(same))
                    new.append(copy.deepcopy(same))
                    came_from.append(all_ps.index(same))
                    n = tpl_paras.index(same) + 1
                    continue
                key = _norm(line)[:PREFIX]
                hit = next((k for k, kk in enumerate(keys) if kk and key and (kk == key or kk[:20] == key[:20])), None)
                # A rewritten line takes the template paragraph after the one used last (the model may
                # merge or drop lines, so a count of lines drifts).
                tpl = tpl_paras[hit] if hit is not None else (
                    tpl_paras[min(n, len(tpl_paras) - 1)] if tpl_paras else (heading_el if heading_el is not None
                                                                             else OxmlElement('w:p')))
                new.append(_para_from_template(tpl, it))
                came_from.append(all_ps.index(tpl) if tpl in all_ps else -1)
                n = (tpl_paras.index(tpl) + 1) if tpl in tpl_paras else n + 1
            else:
                tpl_tbl = tpl_tables[min(t, len(tpl_tables) - 1)] if tpl_tables else None
                new.append(_table_from_template(tpl_tbl, it, tpl_paras[0] if tpl_paras else None))
                came_from.append(came_from[-1] if came_from else -1)
                t += 1
        # Layout the editor can't show stays where the file has it: bracket / dash lines of a tab-stop
        # layout, and drawn shapes (stamp boxes) of paragraphs not kept above (their text goes).
        for j, el in enumerate(all_ps):
            text = texts[id(el)]
            if id(el) in reused or not (is_layout_line(text) or _drawings(el)):
                continue
            keep = copy.deepcopy(el)
            if not is_layout_line(text):
                for r in keep.findall(qn('w:r')):
                    if not _drawings(r):
                        keep.remove(r)
            at = next((k for k, src_j in enumerate(came_from) if src_j > j), len(new))
            new.insert(at, keep)
            came_from.insert(at, j)
        anchor = heading_el if heading_el is not None else (part[0].getprevious() if part else None)
        parent = doc.element.body
        if anchor is None and part:
            for el in new:
                part[0].addprevious(el)
        else:
            for el in new:
                if anchor is not None:
                    anchor.addnext(el)
                    anchor = el
                else:
                    parent.insert(0, el)
        for el in part:
            el.getparent().remove(el)
        record(getattr(block, 'id', None), ([heading_el] if heading_el is not None else []) + new)
        last_el = new[-1] if new else heading_el
    # Clauses of the file the draft no longer has (deleted in the editor) leave with them.
    used = {i for b in blocks for i in range(*(lambda r: (r[0], r[1] + 1))(((b.style_json or {}).get('src') or {})
                                                                          .get('range') or [0, -1]))}
    ranges = [((b.style_json or {}).get('src') or {}).get('range') for b in blocks]
    ranges = [r for r in ranges if r]
    if ranges:
        lo, hi = min(r[0] for r in ranges), max(r[1] for r in ranges)
        for i in range(lo, min(hi, len(els) - 1) + 1):
            if i not in used and els[i].getparent() is not None and _el_text(els[i], doc).strip():
                els[i].getparent().remove(els[i])
    if branding and not (page_form(path) or {}).get('has_letterhead'):
        _branding(doc, branding)
        _signature_block(doc, branding)
    wrap_blocks(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def render_draft_docx(session, branding=None):
    """Download / Save to case: in the Word source's own form when the draft has one, else the
    house-style document (export/docx.py)."""
    from .export.docx import render_session_docx
    # Edited in ONLYOFFICE (docs/ONLYOFFICE_PLAN.md): that Word file is the draft.
    from .onlyoffice import current_docx
    edited = current_docx(session)
    if edited is not None:
        return edited
    ref, path = session_source(session)
    if path:
        try:
            return render_into_original(session, path, branding)
        except Exception:                                         # noqa: BLE001 - never lose a download
            import logging
            logging.getLogger(__name__).exception('render_into_original failed for draft %s', session.id)
    return render_session_docx(session, branding)
