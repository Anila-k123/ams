"""Redline: two versions of a draft as one .docx with real Word tracked changes.

render_redline_docx(before, after, title, author, branding=None) -> (bytes, stats)

`before` / `after` are block snapshots (drafting/versions.py::snapshot_blocks). Word
shows the result in Review mode, so opposite counsel or the client can Accept or
Reject each change, the same as a file marked up in Word itself.

What counts as a change is decided by the comparison engine (export/compare.py), which
the on-screen compare view uses too; this module only writes it as Word XML.
"""

import datetime
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .compare import DEL, INS, SAME, compare_blocks
from .docx import ALIGN, _branding, _setup, _style_or_normal


class _Writer:
    def __init__(self, doc, author):
        self.doc = doc
        self.author = author
        self.date = datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
        self.next_id = 1

    def _mark(self, tag):
        el = OxmlElement(tag)
        el.set(qn('w:id'), str(self.next_id))
        el.set(qn('w:author'), self.author)
        el.set(qn('w:date'), self.date)
        self.next_id += 1
        return el

    def _run(self, text, fmt, deleted=False):
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        for name, tag in (('bold', 'w:b'), ('italic', 'w:i'), ('strike', 'w:strike')):
            if name in fmt:
                rpr.append(OxmlElement(tag))
        if 'underline' in fmt:
            u = OxmlElement('w:u')
            u.set(qn('w:val'), 'single')
            rpr.append(u)
        if len(rpr):
            r.append(rpr)
        for i, part in enumerate(text.split('\n')):
            if i:
                r.append(OxmlElement('w:br'))
            if part:
                t = OxmlElement('w:delText' if deleted else 'w:t')
                t.text = part
                t.set(qn('xml:space'), 'preserve')
                r.append(t)
        return r

    def paragraph(self, pd):
        """One ParaDiff as a Word paragraph. A wholly inserted / deleted paragraph marks its
        paragraph mark too, so accepting a deleted paragraph removes it entirely."""
        p = self.doc.add_paragraph(style=_style_or_normal(self.doc, pd.style))
        if pd.align:
            p.alignment = ALIGN[pd.align]
        if pd.kind in (INS, DEL):
            ppr = p._p.get_or_add_pPr()
            rpr = OxmlElement('w:rPr')
            rpr.append(self._mark('w:ins' if pd.kind == INS else 'w:del'))
            ppr.append(rpr)
        holder, holder_op = None, None
        for seg in pd.segs:
            if seg.op == SAME:
                holder, holder_op = p._p, SAME
            elif seg.op != holder_op:   # consecutive changed runs share one <w:ins> / <w:del>
                holder, holder_op = self._mark('w:ins' if seg.op == INS else 'w:del'), seg.op
                p._p.append(holder)
            holder.append(self._run(seg.text, seg.fmt, deleted=seg.op == DEL))


def render_redline_docx(before, after, title='', author='PactPro', branding=None):
    doc = Document()
    _setup(doc)
    if branding:
        _branding(doc, branding)
    if title:
        doc.add_paragraph(title, style='Title').alignment = WD_ALIGN_PARAGRAPH.CENTER

    comparison = compare_blocks(before, after, title)
    w = _Writer(doc, author)
    for clause in comparison.clauses:
        for pd in clause.heading + clause.body:
            w.paragraph(pd)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), {'inserted': comparison.inserted, 'deleted': comparison.deleted}
