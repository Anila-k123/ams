"""Server-side .docx export of a draft session.

render_session_docx(session, branding=None) -> bytes

Blocks are written in `position` order: `heading` as a Heading 2, the body from
`content_html` (TipTap's HTML) converted to Word paragraphs/runs, falling back
to the plain `text`. The converter handles what the editor produces: headings,
paragraphs, bold/italic/underline/strike, ordered/unordered (nested) lists,
tables, line breaks and text-align. Anything else is kept as plain text.

`branding` is the firm's AMS profile (drafting/branding.py::firm_branding), optionally with
image bytes under 'logo', 'signature' and 'seal'. Missing pieces are skipped.
"""

import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

FONT = 'Times New Roman'
BLOCK_TAGS = {'p', 'div', 'blockquote', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li', 'pre'}
HEADING_LEVEL = {'h1': 1, 'h2': 2, 'h3': 3, 'h4': 3, 'h5': 3, 'h6': 3}
ALIGN = {'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
         'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT}
FORMAT_TAGS = {'strong': 'bold', 'b': 'bold', 'em': 'italic', 'i': 'italic', 'u': 'underline',
               's': 'strike', 'strike': 'strike', 'del': 'strike'}


# --- HTML -> intermediate blocks ---------------------------------------------

@dataclass
class Para:
    style: str = 'Normal'
    align: str | None = None
    runs: list = field(default_factory=list)  # (text, frozenset(formats)) or ('\n', None) for <br>

    def is_empty(self):
        return not any(t.strip() for t, _ in self.runs if t != '\n')


@dataclass
class Table:
    rows: list = field(default_factory=list)  # rows -> cells -> [Para]


class _HtmlToBlocks(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []            # top-level output
        self.targets = [self.blocks]  # where new paragraphs go (top level or a table cell)
        self.formats = []           # stack of active inline formats
        self.lists = []             # stack of 'ol' / 'ul'
        self.tables = []            # stack of Table being built
        self.para = None
        self.pending = ('Normal', None)
        self.li_fresh = False       # just opened an <li>, no content yet
        self.skip = 0               # inside <style>/<script>

    # paragraphs
    def _close(self):
        if self.para is not None and not self.para.is_empty():
            self.targets[-1].append(self.para)
        self.para = None

    def _open(self, style, align):
        self._close()
        self.pending = (style, align)

    def _current(self):
        if self.para is None:
            self.para = Para(*self.pending)
        return self.para

    @staticmethod
    def _align(attrs):
        m = re.search(r'text-align\s*:\s*(\w+)', dict(attrs).get('style') or '')
        return m.group(1).lower() if m and m.group(1).lower() in ALIGN else None

    def _list_style(self):
        kind = 'List Number' if self.lists[-1] == 'ol' else 'List Bullet'
        depth = min(len(self.lists), 3)
        return kind if depth == 1 else f'{kind} {depth}'

    def handle_starttag(self, tag, attrs):
        if tag in ('style', 'script'):
            self.skip += 1
        elif tag in FORMAT_TAGS:
            self.formats.append(FORMAT_TAGS[tag])
        elif tag in ('ul', 'ol'):
            self._close()
            self.lists.append(tag)
        elif tag == 'li':
            self._open(self._list_style() if self.lists else 'Normal', self._align(attrs))
            self.li_fresh = True
        elif tag in HEADING_LEVEL:
            self.li_fresh = False
            self._open(f'Heading {HEADING_LEVEL[tag]}', self._align(attrs))
        elif tag in BLOCK_TAGS:
            if self.li_fresh:
                # TipTap writes <li><p>…</p></li>: the first <p> IS the list item's paragraph.
                self.li_fresh = False
                return
            self._open('Normal', self._align(attrs))
        elif tag == 'br':
            self._current().runs.append(('\n', None))
        elif tag == 'table':
            self._close()
            self.tables.append(Table())
        elif tag == 'tr' and self.tables:
            self.tables[-1].rows.append([])
        elif tag in ('td', 'th') and self.tables:
            self._close()
            cell = []
            if not self.tables[-1].rows:
                self.tables[-1].rows.append([])
            self.tables[-1].rows[-1].append(cell)
            self.targets.append(cell)
            self.pending = ('Normal', None)
            if tag == 'th':
                self.formats.append('bold')

    def handle_endtag(self, tag):
        if tag in ('style', 'script'):
            self.skip = max(0, self.skip - 1)
        elif tag in FORMAT_TAGS:
            if FORMAT_TAGS[tag] in self.formats:
                self.formats.reverse()
                self.formats.remove(FORMAT_TAGS[tag])
                self.formats.reverse()
        elif tag in ('ul', 'ol'):
            self._close()
            if self.lists:
                self.lists.pop()
            self.pending = ('Normal', None)
        elif tag in HEADING_LEVEL or tag in BLOCK_TAGS:
            self._close()
            self.li_fresh = False
            self.pending = ('Normal', None)
        elif tag in ('td', 'th') and len(self.targets) > 1:
            self._close()
            self.targets.pop()
            if tag == 'th' and 'bold' in self.formats:
                self.formats.remove('bold')
        elif tag == 'table' and self.tables:
            self._close()
            table = self.tables.pop()
            if any(table.rows):
                self.targets[-1].append(table)

    def handle_data(self, data):
        if self.skip:
            return
        text = re.sub(r'\s+', ' ', data)
        if not text.strip() and self.para is None:
            return  # whitespace between block tags
        self.li_fresh = False
        para = self._current()
        if not para.runs or para.runs[-1][0] == '\n':
            text = text.lstrip()
        if text:
            para.runs.append((text, frozenset(self.formats)))

    def close(self):
        super().close()
        self._close()
        while self.tables:  # unclosed table
            table = self.tables.pop()
            self.targets = [self.blocks]
            if any(table.rows):
                self.blocks.append(table)
        return self.blocks


def html_to_blocks(html):
    parser = _HtmlToBlocks()
    parser.feed(html or '')
    return parser.close()


# --- blocks -> python-docx ---------------------------------------------------

def _style_or_normal(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles['Normal']


def _write_para(doc, container, para):
    p = container.add_paragraph(style=_style_or_normal(doc, para.style))
    if para.align:
        p.alignment = ALIGN[para.align]
    runs = list(para.runs)
    while runs and runs[-1][0] == '\n':
        runs.pop()
    for text, fmt in runs:
        if text == '\n':
            (p.runs[-1] if p.runs else p.add_run()).add_break()
            continue
        run = p.add_run(text)
        run.bold = 'bold' in fmt or None
        run.italic = 'italic' in fmt or None
        run.underline = 'underline' in fmt or None
        run.font.strike = 'strike' in fmt or None
    return p


def _write_blocks(doc, container, blocks):
    for block in blocks:
        if isinstance(block, Para):
            _write_para(doc, container, block)
        else:
            cols = max(len(r) for r in block.rows)
            table = container.add_table(rows=0, cols=cols)
            table.style = _style_or_normal(doc, 'Table Grid')
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for row in block.rows:
                cells = table.add_row().cells
                for i, cell_blocks in enumerate(row):
                    cell = cells[i]
                    _write_blocks(doc, cell, cell_blocks)
                    # add_row() gives each cell an empty first paragraph; drop it if we added our own.
                    if len(cell.paragraphs) > 1 and not cell.paragraphs[0].text:
                        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)


def _text_to_blocks(text):
    return [Para(runs=[(line.strip(), frozenset())]) for line in (text or '').split('\n') if line.strip()]


def _setup(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (('Title', 16), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 12)):
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.bold = FONT, Pt(size), True
        style.font.color.rgb = RGBColor(0, 0, 0)
        # East-Asian font slot too, or Word keeps the theme font for some text.
        rpr = style.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', FONT)


def _add_image(paragraph, data, height_cm):
    if not data:
        return False
    try:
        paragraph.add_run().add_picture(io.BytesIO(data), height=Cm(height_cm))
        return True
    except Exception:  # not an image python-docx understands
        return False


def _branding(doc, branding):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if _add_image(header, branding.get('logo'), 1.5):
        header.add_run().add_break()
    if branding.get('officeName'):
        header.add_run(branding['officeName']).bold = True
    address = branding.get('officeAddress') or ', '.join(
        x for x in (branding.get('address'), branding.get('city'), branding.get('state'),
                    branding.get('pinCode')) if x)
    if address:
        line = section.header.add_paragraph(address)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.runs[0].font.size = Pt(9)

    ids = ' · '.join(f'{label}: {branding[key]}' for label, key in
                     (('GSTIN', 'gstNumber'), ('PAN', 'panNumber')) if branding.get(key))
    if ids:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(ids).font.size = Pt(9)


def _signature_block(doc, branding):
    if not any(branding.get(k) for k in ('signature', 'seal', 'fullName')):
        return
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if branding.get('officeName'):
        p.add_run(f"For {branding['officeName']}").bold = True
    images = doc.add_paragraph()
    images.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_image(images, branding.get('signature'), 1.8)
    if branding.get('seal'):
        images.add_run('   ')
        _add_image(images, branding.get('seal'), 2.2)
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if branding.get('fullName'):
        name.add_run(branding['fullName']).bold = True
    if branding.get('barCouncilId'):
        name.add_run().add_break()
        name.add_run(f"Enrolment No. {branding['barCouncilId']}")


DOC_TITLE_RE = re.compile(r'\b(agreement|deed|nda|mou|memorandum|affidavit|vakalatnama|vakalathnama|'
                          r'contract|lease|power of attorney|undertaking|petition|plaint|notice)\b', re.I)


def _title_from_blocks(session):
    """The draft's own title: the first line of the first block when it reads like
    one (short, and mostly upper-case or naming a legal document). Mirrors the
    editor's extractDocTitle (DraftPage.tsx)."""
    first = session.blocks.order_by('position').first()
    if first is None:
        return ''
    # The title is either the first block's heading or the first line of its text.
    first_line = next((ln.strip() for ln in re.sub(r'\*\*|__|^#+\s*', '', first.text or '', flags=re.M)
                       .split('\n') if ln.strip()), '')
    # A heading only counts if it names a document ("DEFINITIONS" is a clause, not a title).
    for candidate, need_keyword in ((first.heading or '', True), (first_line, False)):
        title = _as_title(re.sub(r'\s+', ' ', candidate).strip(), need_keyword)
        if title:
            return title
    return ''


def _as_title(line, need_keyword=False):
    if not line or len(line) > 90 or len(line.split()) > 10:
        return ''
    letters = [c for c in line if c.isalpha()]
    upper = letters and sum(c.isupper() for c in letters) / len(letters) > 0.8
    # A keyword only makes a title of a line that doesn't read as a sentence.
    sentence = ',' in line or line.rstrip().endswith(('.', '…', ';'))
    has_keyword = bool(DOC_TITLE_RE.search(line)) and not sentence
    if not (has_keyword if need_keyword else (upper or has_keyword)):
        return ''
    line = re.sub(r'\s+', ' ', line).strip(' .:')
    return line.title() if upper else line


def session_title(session):
    """Same precedence as the editor: typed title, else the draft's own title line,
    else template name (hidden for from-scratch drafts), else the document type."""
    title = ((session.facts or {}).get('document_title') or '').strip() or _title_from_blocks(session)
    if not title and session.template_id:
        if session.mode != 'library':
            title = session.template.name
        title = title or (session.template.document_type or '').upper()
    return title


def render_session_docx(session, branding=None):
    doc = Document()
    _setup(doc)
    if branding:
        _branding(doc, branding)

    title = session_title(session)
    if title:
        doc.add_paragraph(title, style='Title').alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in session.blocks.order_by('position'):
        if block.heading:
            doc.add_paragraph(block.heading, style='Heading 2')
        body = html_to_blocks(block.content_html) if (block.content_html or '').strip() \
            else _text_to_blocks(block.text)
        _write_blocks(doc, doc, body)

    if branding:
        _signature_block(doc, branding)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
